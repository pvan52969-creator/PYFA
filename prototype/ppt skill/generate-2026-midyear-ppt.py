#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""厦大马来项目 2026 年中汇报 PPT。只写产品设计已能量化的进展，不留待填空。"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_MID = RGBColor(0x14, 0x32, 0x58)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
BG = RGBColor(0xF6, 0xF8, 0xFB)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0xEE, 0xF2, 0xF7)

OUT = Path(__file__).resolve().parent / "厦大马来项目2026年中汇报-教务与考勤.pptx"
TOTAL = 5


def set_run(run, text, size, color, bold=False, name="微软雅黑"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    return sh


def add_round(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def add_tb(slide, l, t, w, h, lines, size=14, color=INK, bold=False,
           align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, sz, col, bd = item
        else:
            text, sz, col, bd = item, size, color, bold
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(3)
        set_run(p.add_run(), text, sz, col, bd)
    return tb


def footer(slide, page):
    add_box(slide, Inches(0), Inches(7.28), W, Inches(0.22), NAVY)
    add_tb(
        slide, Inches(0.5), Inches(7.28), Inches(8), Inches(0.22),
        [("厦大马来分校项目  ·  2026 年中汇报  ·  产品设计进展", 9, WHITE, False)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_tb(
        slide, Inches(11.2), Inches(7.28), Inches(1.6), Inches(0.22),
        [(f"{page}  /  {TOTAL}", 9, WHITE, False)],
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def accent_bar(slide):
    add_box(slide, Inches(0), Inches(0), Inches(0.12), H, GOLD)


def title_block(slide, title, subtitle):
    add_tb(slide, Inches(0.5), Inches(0.28), Inches(12.2), Inches(0.42),
           [(title, 24, NAVY, True)])
    add_box(slide, Inches(0.5), Inches(0.74), Inches(1.1), Inches(0.05), GOLD)
    add_tb(slide, Inches(0.5), Inches(0.84), Inches(12.2), Inches(0.32),
           [(subtitle, 12, MUTED, False)])


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, NAVY)
    add_box(s, 0, 0, Inches(0.16), H, GOLD)
    add_box(s, Inches(0.5), Inches(1.85), Inches(1.4), Inches(0.06), GOLD)
    add_tb(s, Inches(0.5), Inches(1.15), Inches(12), Inches(0.36),
           [("XMUM  ·  马来项目", 14, GOLD, True)])
    add_tb(s, Inches(0.5), Inches(2.05), Inches(12), Inches(0.7),
           [("2026 年中工作汇报", 40, WHITE, True)])
    add_tb(s, Inches(0.5), Inches(2.85), Inches(12), Inches(0.4),
           [("产品设计进展  ·  六个模块原型 + 考勤已上线", 18, RGBColor(0xCB, 0xD5, 0xE1), False)])

    row1 = ["基础数据 18 页", "学籍 9 页", "选课 12 页"]
    row2 = ["培养方案 11 页", "开课 24 页", "排课 23 页", "考勤 4 月已上线"]
    x = Inches(0.5)
    for chip in row1:
        w = Inches(2.35)
        add_round(s, x, Inches(3.7), w, Inches(0.42), NAVY_MID)
        add_tb(s, x, Inches(3.7), w, Inches(0.42),
               [(chip, 13, WHITE, False)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += w + Inches(0.14)
    x = Inches(0.5)
    widths = [Inches(2.35), Inches(2.2), Inches(2.2), Inches(2.7)]
    for chip, w in zip(row2, widths):
        add_round(s, x, Inches(4.28), w, Inches(0.42), NAVY_MID)
        add_tb(s, x, Inches(4.28), w, Inches(0.42),
               [(chip, 13, WHITE, False)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += w + Inches(0.14)

    add_tb(s, Inches(0.5), Inches(5.05), Inches(12), Inches(0.7),
           [
               ("合计 97 张独立原型页  ·  侧栏入口 86 个  ·  考勤不按原型页计", 14, RGBColor(0xCB, 0xD5, 0xE1), False),
               ("来源：基础数据复刻（基础数据 / 学籍 / 选课）+ 本教务原型（培养方案 / 开课 / 排课）", 12, RGBColor(0x94, 0xA3, 0xB8), False),
           ])
    add_tb(s, Inches(0.5), Inches(6.55), Inches(12), Inches(0.3),
           [("2026 年 8 月  ·  开发进度由开发负责人另行汇报", 12, RGBColor(0x94, 0xA3, 0xB8), False)])


def slide_scope(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "产品设计覆盖：六个模块 + 考勤已上线",
                "只报菜单入口和独立页实数。弹窗、向导、页内 Tab 不另计页。")

    kpis = [
        ("6", "设计中的模块", "基础数据 · 学籍 · 选课\n培养方案 · 开课 · 排课"),
        ("86", "侧栏可进入口", "基础三件套 39  ·  教务三件套 47"),
        ("97", "独立原型页面", "基础三件套 39  ·  教务三件套 58"),
        ("1", "已上线系统", "考勤，4 月正式使用"),
    ]
    for i, (n, label, sub) in enumerate(kpis):
        x = Inches(0.45) + i * Inches(3.18)
        add_round(s, x, Inches(1.28), Inches(3.05), Inches(1.95), CARD, LINE)
        add_tb(s, x, Inches(1.36), Inches(3.05), Inches(0.58),
               [(n, 28, NAVY, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x + Inches(0.12), Inches(1.94), Inches(2.81), Inches(0.32),
               [(label, 13, INK, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x + Inches(0.12), Inches(2.28), Inches(2.81), Inches(0.78),
               [(sub, 11, MUTED, False)], align=PP_ALIGN.CENTER)

    modules = [
        ("基础数据", "18 入口 / 18 页", "学校 · 专业 · 课程 · 讲师 · 校历"),
        ("学籍", "9 入口 / 9 页", "档案 + 异动申请到查询"),
        ("选课", "12 入口 / 12 页", "学生端 3  ·  管理端 8  ·  说明 1"),
        ("培养方案", "9 入口 / 11 页", "版本全周期 + 2 个编辑子页"),
        ("开课", "16 入口 / 24 页", "计划 / 安排 / 名单已闭环"),
        ("排课", "22 入口 / 23 页", "时间教室双轨 + 调课"),
    ]
    for i, (name, num, desc) in enumerate(modules):
        x = Inches(0.45) + i * Inches(2.12)
        add_round(s, x, Inches(3.42), Inches(2.0), Inches(2.05), CARD, LINE)
        add_box(s, x, Inches(3.42), Inches(2.0), Inches(0.07), BLUE)
        add_tb(s, x + Inches(0.08), Inches(3.56), Inches(1.84), Inches(0.36),
               [(name, 13, NAVY, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x + Inches(0.08), Inches(3.92), Inches(1.84), Inches(0.4),
               [(num, 12, GOLD, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x + Inches(0.08), Inches(4.32), Inches(1.84), Inches(0.95),
               [(desc, 11, MUTED, False)], align=PP_ALIGN.CENTER)

    add_round(s, Inches(0.45), Inches(5.62), Inches(12.4), Inches(1.4), SLATE)
    add_tb(s, Inches(0.7), Inches(5.72), Inches(12), Inches(1.2),
           [
               ("链路：基础数据 → 学籍 → 培养方案 → 开课 → 排课 → 选课 → 考勤", 14, NAVY, True),
               ("基础数据 / 学籍 / 选课：基础数据复刻现挂菜单。培养方案 / 开课 / 排课：本原型 index.html 现存独立页。", 12, MUTED, False),
               ("已下线页、空壳页不计入。开发进度由开发负责人汇报，产品侧不报完成百分比。", 12, MUTED, False),
           ])
    footer(s, 2)


def slide_academic_chain(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "教务三件套：培养方案 / 开课 / 排课",
                "来源：本教务原型。培养方案 11 页，开课 24 页，排课 23 页。")

    cols = [
        ("培养方案", "11 页", [
            ("侧栏 9 个入口", "流程图、版本、审批、查询、变更申请 / 审核、执行计划、两类统计"),
            ("另 2 个编辑子页", "方案版本编辑、执行计划编辑"),
            ("当前设计状态", "版本全周期 + 执行计划成套，开课从执行计划取数"),
        ]),
        ("开课管理", "24 页", [
            ("侧栏 16 个入口", "开课设置 4、专业开课 3、选修开课 3、课程班管理 6"),
            ("另 8 页", "分组、安排样式一、特殊开课 3 页、共同授课与变更日志"),
            ("当前设计状态", "计划 → 安排 → 名单已闭环；特殊开课从专业开课入口添加"),
        ]),
        ("排课管理", "23 页", [
            ("侧栏 22 个入口", "设置 4、排时间 3、排教室 3、课表查询 6、调课 6"),
            ("另 1 个详情子页", "排课表时间 / 教室共用的排课详情"),
            ("当前设计状态", "时间与教室双轨；多维课表查询；调课申请到批量调课"),
        ]),
    ]
    for i, (title, pages, items) in enumerate(cols):
        x = Inches(0.4) + i * Inches(4.28)
        add_round(s, x, Inches(1.28), Inches(4.12), Inches(5.72), CARD, LINE)
        add_box(s, x, Inches(1.28), Inches(4.12), Inches(0.56), NAVY)
        add_tb(s, x + Inches(0.16), Inches(1.28), Inches(2.4), Inches(0.56),
               [(title, 16, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        add_tb(s, x + Inches(2.3), Inches(1.28), Inches(1.65), Inches(0.56),
               [(pages, 14, GOLD, True)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y = Inches(2.04)
        for h, d in items:
            add_tb(s, x + Inches(0.18), y, Inches(3.76), Inches(0.32), [(h, 13, NAVY, True)])
            add_tb(s, x + Inches(0.18), y + Inches(0.32), Inches(3.76), Inches(0.95), [(d, 12, MUTED, False)])
            y += Inches(1.48)
    footer(s, 3)


def slide_foundation(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "基础三件套：基础数据 / 学籍 / 选课",
                "来源：基础数据复刻现挂菜单。基础数据 18 页，学籍 9 页，选课 12 页。")

    cols = [
        ("基础数据", "18 页", [
            ("侧栏 18 个入口", "仪表盘、学校 / 院系 / 代码集、专业版本 / 批次、楼栋教室、课程全周期、讲师、学期校历"),
            ("口径", "菜单即页面；创建 / 审批向导在页内打开，不另计页"),
            ("当前设计状态", "院校基础配置齐；课程申请、审批、变更已成套"),
        ]),
        ("学籍管理", "9 页", [
            ("侧栏 9 个入口", "学生基本信息 1；异动类别 / 同意书 / 规则 / 申请双端 / 审批 / 维护 / 查询 8"),
            ("口径", "转专业、休学、复学、退学是申请页内 Tab，不另计页"),
            ("当前设计状态", "档案 + 异动主流程已挂出；异动统计有页面未挂菜单，不计"),
        ]),
        ("选课管理", "12 页", [
            ("侧栏 12 个入口", "学生端 3（在线选课 / 历史 / 加退课）；管理端 8（批次规则、监控审批、缴费补注册、结果日志）；流程说明 1"),
            ("口径", "流程说明菜单标注正式上线去掉，按现存页计入"),
            ("当前设计状态", "批次轮次、学生选课、监控审批、结果日志已齐"),
        ]),
    ]
    for i, (title, pages, items) in enumerate(cols):
        x = Inches(0.4) + i * Inches(4.28)
        add_round(s, x, Inches(1.28), Inches(4.12), Inches(5.72), CARD, LINE)
        add_box(s, x, Inches(1.28), Inches(4.12), Inches(0.56), NAVY)
        add_tb(s, x + Inches(0.16), Inches(1.28), Inches(2.4), Inches(0.56),
               [(title, 16, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        add_tb(s, x + Inches(2.3), Inches(1.28), Inches(1.65), Inches(0.56),
               [(pages, 14, GOLD, True)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y = Inches(2.04)
        for h, d in items:
            add_tb(s, x + Inches(0.18), y, Inches(3.76), Inches(0.32), [(h, 13, NAVY, True)])
            add_tb(s, x + Inches(0.18), y + Inches(0.32), Inches(3.76), Inches(0.95), [(d, 12, MUTED, False)])
            y += Inches(1.48)
    footer(s, 4)


def slide_attendance(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "考勤：已上线，上半年以运维和接需求为主",
                "只写已发生的事实。缺陷条数、使用量由开发 / 运维口径另报。")

    facts = [
        ("年初", "完成系统设计"),
        ("4 月", "正式上线使用"),
        ("数据", "对接学校老教务系统"),
        ("4 月学期", "持续修缺陷、接新需求"),
        ("过程", "上线后问题较多，边用边改"),
    ]
    for i, (k, v) in enumerate(facts):
        x = Inches(0.45) + i * Inches(2.52)
        add_round(s, x, Inches(1.38), Inches(2.38), Inches(2.15), CARD, LINE)
        add_tb(s, x + Inches(0.14), Inches(1.55), Inches(2.1), Inches(0.4),
               [(k, 13, GOLD, True)])
        add_tb(s, x + Inches(0.14), Inches(2.05), Inches(2.1), Inches(1.15),
               [(v, 16, NAVY, True)])

    add_round(s, Inches(0.45), Inches(3.75), Inches(12.4), Inches(3.2), CARD, LINE)
    add_tb(s, Inches(0.7), Inches(3.92), Inches(12), Inches(0.36),
           [("和教务设计的关系", 16, NAVY, True)])
    notes = [
        "考勤与教务并行：考勤已先上线，其余六个模块仍在产品设计与确认阶段。",
        "上半年产品精力两边并行——教务与基础模块出原型，考勤跟现场问题和新增需求。",
        "下半年继续收口六个模块的设计确认；考勤继续消化运行中的缺陷与新需求。",
        "开发进度、排期与缺陷量化，由开发负责人补充。",
    ]
    y = Inches(4.4)
    for t in notes:
        add_box(s, Inches(0.75), y + Inches(0.08), Inches(0.1), Inches(0.1), BLUE)
        add_tb(s, Inches(1.05), y, Inches(11.5), Inches(0.5), [(t, 14, INK, False)])
        y += Inches(0.52)
    footer(s, 5)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide_cover(prs)
    slide_scope(prs)
    slide_academic_chain(prs)
    slide_foundation(prs)
    slide_attendance(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
