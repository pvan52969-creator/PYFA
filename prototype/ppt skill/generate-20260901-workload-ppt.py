#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""六个模块原型工作量统计 PPT。数字来自 2026-09-01 代码复核，不编未落档功能点。"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W = Inches(13.333)
H = Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
NAVY_MID = RGBColor(0x14, 0x32, 0x58)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
BG = RGBColor(0xF6, 0xF8, 0xFB)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0xEE, 0xF2, 0xF7)
ROW_ALT = RGBColor(0xF8, 0xFA, 0xFC)

OUT = Path(__file__).resolve().parent / "六个模块原型工作量统计-20260901.pptx"
TOTAL = 5


def set_run(run, text, size, color, bold=False, name="微软雅黑"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", name)


def add_box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    return sh


def add_round(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def add_tb(slide, l, t, w, h, lines, size=14, color=INK, bold=False,
           align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    for i, item in enumerate(lines):
        if isinstance(item, tuple):
            text, sz, col, bd = item
        else:
            text, sz, col, bd = item, size, color, bold
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        set_run(p.add_run(), text, sz, col, bd)
    return tb


def footer(slide, page):
    add_box(slide, Inches(0), Inches(7.28), W, Inches(0.22), NAVY)
    add_tb(
        slide, Inches(0.5), Inches(7.28), Inches(9.2), Inches(0.22),
        [("厦大马来分校项目  ·  产品设计工作量  ·  2026-09-01 代码复核", 9, WHITE, False)],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_tb(
        slide, Inches(11.2), Inches(7.28), Inches(1.6), Inches(0.22),
        [(f"{page}  /  {TOTAL}", 9, WHITE, False)],
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def accent_bar(slide):
    add_box(slide, Inches(0), Inches(0), Inches(0.12), H, GOLD)


def title_block(slide, title, subtitle):
    add_tb(slide, Inches(0.5), Inches(0.24), Inches(12.2), Inches(0.40),
           [(title, 22, NAVY, True)])
    add_box(slide, Inches(0.5), Inches(0.68), Inches(1.1), Inches(0.05), GOLD)
    add_tb(slide, Inches(0.5), Inches(0.76), Inches(12.2), Inches(0.28),
           [(subtitle, 12, MUTED, False)])


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, NAVY)
    add_box(s, 0, 0, Inches(0.16), H, GOLD)
    add_box(s, Inches(0.5), Inches(1.7), Inches(1.4), Inches(0.06), GOLD)
    add_tb(s, Inches(0.5), Inches(1.05), Inches(12), Inches(0.36),
           [("XMUM  ·  马来项目  ·  产品设计", 14, GOLD, True)])
    add_tb(s, Inches(0.5), Inches(1.9), Inches(12.2), Inches(0.7),
           [("六个模块原型工作量", 36, WHITE, True)])
    add_tb(s, Inches(0.5), Inches(2.7), Inches(12.2), Inches(0.4),
           [("菜单入口  ·  独立原型页  ·  已落档功能点", 18, RGBColor(0xCB, 0xD5, 0xE1), False)])

    kpis = [
        ("87", "叶子菜单"),
        ("98", "独立原型页"),
        ("189", "已落档功能点"),
        ("6", "设计模块"),
    ]
    for i, (n, label) in enumerate(kpis):
        x = Inches(0.5) + i * Inches(3.1)
        add_round(s, x, Inches(3.5), Inches(2.9), Inches(1.35), NAVY_MID)
        add_tb(s, x, Inches(3.58), Inches(2.9), Inches(0.7),
               [(n, 28, GOLD, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x, Inches(4.28), Inches(2.9), Inches(0.4),
               [(label, 13, WHITE, False)], align=PP_ALIGN.CENTER)

    add_tb(s, Inches(0.5), Inches(5.15), Inches(12.2), Inches(0.7),
           [
               ("基础数据 18 页  ·  学籍 9 页 / 67 功能  ·  选课 13 页", 14, RGBColor(0xCB, 0xD5, 0xE1), False),
               ("培养方案 11 页 / 54 功能  ·  开课 24 页 / 68 功能  ·  排课 23 页", 14, RGBColor(0xCB, 0xD5, 0xE1), False),
           ])
    add_tb(s, Inches(0.5), Inches(6.5), Inches(12.2), Inches(0.3),
           [("2026 年 9 月 1 日  ·  功能点只报已出 PRD 的清单，未落档不估算", 12, RGBColor(0x94, 0xA3, 0xB8), False)])


def slide_summary(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "工作量总表：87 个菜单，98 张页面",
                "功能点只加已落档三项：培养方案 54 + 开课 68 + 学籍 67 = 189。基础数据 / 选课 / 排课不硬加总。")

    kpis = [
        ("40", "基础三件套菜单", "基础 18  ·  学籍 9  ·  选课 13"),
        ("47", "教务三件套菜单", "方案 9  ·  开课 16  ·  排课 22"),
        ("40", "基础三件套页面", "基础 18  ·  学籍 9  ·  选课 13"),
        ("58", "教务三件套页面", "方案 11  ·  开课 24  ·  排课 23"),
    ]
    for i, (n, label, sub) in enumerate(kpis):
        x = Inches(0.45) + i * Inches(3.18)
        add_round(s, x, Inches(1.16), Inches(3.05), Inches(1.28), CARD, LINE)
        add_tb(s, x, Inches(1.22), Inches(3.05), Inches(0.48),
               [(n, 24, NAVY, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x + Inches(0.1), Inches(1.68), Inches(2.85), Inches(0.28),
               [(label, 12, INK, True)], align=PP_ALIGN.CENTER)
        add_tb(s, x + Inches(0.1), Inches(1.96), Inches(2.85), Inches(0.36),
               [(sub, 10, MUTED, False)], align=PP_ALIGN.CENTER)

    headers = ["模块", "叶子菜单", "独立页", "功能点", "说明"]
    rows = [
        ["基础数据", "18", "18", "—", "菜单即页面；功能清单未落档"],
        ["学籍管理", "9", "9", "67", "PRD 按钮清单已落档"],
        ["选课管理", "13", "13", "—", "较年中 PPT +1 个管理端入口"],
        ["培养方案", "9", "11", "54", "含 2 个编辑子页；PRD V2.2"],
        ["开课管理", "16", "24", "68", "10 / 16 个菜单已出功能清单"],
        ["排课管理", "22", "23", "—", "本仓库无排课功能清单"],
        ["合计", "87", "98", "189", "功能点仅已落档三项可加总"],
    ]
    cols_w = [Inches(2.1), Inches(1.5), Inches(1.4), Inches(1.4), Inches(5.7)]
    x0, y0, rh = Inches(0.45), Inches(2.6), Inches(0.48)
    # header
    x = x0
    for w, htxt in zip(cols_w, headers):
        add_box(s, x, y0, w, rh, NAVY)
        add_tb(s, x, y0, w, rh, [(htxt, 12, WHITE, True)],
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += w
    for r, row in enumerate(rows):
        y = y0 + rh * (r + 1)
        fill = SLATE if r == len(rows) - 1 else (CARD if r % 2 == 0 else ROW_ALT)
        x = x0
        aligns = [PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.CENTER, PP_ALIGN.LEFT]
        pads = [Inches(0.12), 0, 0, 0, Inches(0.12)]
        bold = r == len(rows) - 1
        for c, (w, cell, al, pad) in enumerate(zip(cols_w, row, aligns, pads)):
            add_box(s, x, y, w, rh, fill, LINE)
            color = NAVY if (c in (1, 2) or r == len(rows) - 1) else INK
            if c == 3:
                color = GOLD if cell not in ("—",) else MUTED
            add_tb(s, x + pad, y, w - pad, rh,
                   [(cell, 12, color, bold or c == 0)],
                   align=al, anchor=MSO_ANCHOR.MIDDLE)
            x += w

    add_tb(s, Inches(0.45), Inches(6.9), Inches(12.4), Inches(0.28),
           [("口径：侧栏可点入口计菜单；独立页含未挂侧栏的下钻页；弹窗 / Tab / 空壳不计页。", 11, MUTED, False)])
    footer(s, 2)


def slide_academic(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "教务三件套：培养方案 / 开课 / 排课",
                "来源：本教务原型。47 个菜单，58 张页面；已落档功能点 54 + 68 = 122。")

    cols = [
        ("培养方案", "11 页 · 54 功能", [
            ("侧栏 9 个入口", "流程图、版本、审批、查询、变更申请 / 审核、执行计划、两类统计"),
            ("另 2 个编辑子页", "方案版本编辑、执行计划编辑"),
            ("功能点 54", "版本管理含编辑页 34；变更 9；执行计划 5；审批 / 查询 / 统计 / 流程图 6"),
        ]),
        ("开课管理", "24 页 · 68 功能", [
            ("侧栏 16 个入口", "开课设置 4、专业开课 3、选修开课 3、课程班管理 6"),
            ("另 8 页未挂侧栏", "分组工作台、安排样式一、特殊开课 3 页、共同授课与两条变更日志"),
            ("功能点 68 / 部分", "10 个菜单已出 PRD；选修 3 页、教师替换、排课计划、流程图未出清单"),
        ]),
        ("排课管理", "23 页 · 未落档", [
            ("侧栏 22 个入口", "流程图 1、基础设置 3、排时间 3、排教室 3、课表查询 6、调课 6"),
            ("另 1 个详情子页", "排课表时间 / 教室共用的排课详情"),
            ("功能点", "本仓库无排课功能清单，会上只报菜单和页数"),
        ]),
    ]
    for i, (title, pages, items) in enumerate(cols):
        x = Inches(0.4) + i * Inches(4.28)
        add_round(s, x, Inches(1.16), Inches(4.12), Inches(5.84), CARD, LINE)
        add_box(s, x, Inches(1.16), Inches(4.12), Inches(0.56), NAVY)
        add_tb(s, x + Inches(0.16), Inches(1.16), Inches(1.9), Inches(0.56),
               [(title, 15, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        add_tb(s, x + Inches(1.95), Inches(1.16), Inches(2.05), Inches(0.56),
               [(pages, 12, GOLD, True)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y = Inches(1.88)
        for htxt, d in items:
            add_tb(s, x + Inches(0.18), y, Inches(3.76), Inches(0.32), [(htxt, 13, NAVY, True)])
            add_tb(s, x + Inches(0.18), y + Inches(0.32), Inches(3.76), Inches(1.05),
                   [(d, 12, MUTED, False)])
            y += Inches(1.52)
    footer(s, 3)


def slide_foundation(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "基础三件套：基础数据 / 学籍 / 选课",
                "来源：基础数据复刻现挂菜单。40 个菜单，40 张页面；仅学籍功能点可报 67。")

    cols = [
        ("基础数据", "18 页 · 未落档", [
            ("侧栏 18 个入口", "仪表盘、学校 / 院系 / 代码集、专业版本 / 批次、楼栋教室、课程全周期、讲师、学期校历"),
            ("页面口径", "菜单即页面；创建 / 审批向导在页内打开，不另计页"),
            ("功能点", "无同口径功能清单，不上会估算"),
        ]),
        ("学籍管理", "9 页 · 67 功能", [
            ("侧栏 9 个入口", "学生基本信息 1；异动类别 / 同意书 / 规则 / 申请双端 / 审批 / 维护 / 查询 8"),
            ("页面口径", "转专业、休学、复学、退学是申请页内 Tab，不另计页"),
            ("功能点 67", "PRD 按钮清单；异动统计有页未挂菜单，不计页也不计功能"),
        ]),
        ("选课管理", "13 页 · 未落档", [
            ("侧栏 13 个入口", "学生端 3；配置 3；过程 4；结果 2；流程说明 1"),
            ("较年中变化", "管理端现 9 个入口，年中按 8 个计，合计从 12 增至 13"),
            ("功能点", "无同口径功能清单，不上会估算"),
        ]),
    ]
    for i, (title, pages, items) in enumerate(cols):
        x = Inches(0.4) + i * Inches(4.28)
        add_round(s, x, Inches(1.16), Inches(4.12), Inches(5.84), CARD, LINE)
        add_box(s, x, Inches(1.16), Inches(4.12), Inches(0.56), NAVY)
        add_tb(s, x + Inches(0.16), Inches(1.16), Inches(1.9), Inches(0.56),
               [(title, 15, WHITE, True)], anchor=MSO_ANCHOR.MIDDLE)
        add_tb(s, x + Inches(1.95), Inches(1.16), Inches(2.05), Inches(0.56),
               [(pages, 12, GOLD, True)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y = Inches(1.88)
        for htxt, d in items:
            add_tb(s, x + Inches(0.18), y, Inches(3.76), Inches(0.32), [(htxt, 13, NAVY, True)])
            add_tb(s, x + Inches(0.18), y + Inches(0.32), Inches(3.76), Inches(1.05),
                   [(d, 12, MUTED, False)])
            y += Inches(1.52)
    footer(s, 4)


def slide_caliber(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_box(s, 0, 0, W, H, BG)
    accent_bar(s)
    title_block(s, "统计口径：只报能对账的数",
                "会上被问「功能点怎么加出来的」时，按本页回答。")

    cards = [
        ("计菜单", "侧栏当前可点击的叶子入口。", "不计分组标题、已注释下线菜单。"),
        ("计页面", "独立原型页，含未挂侧栏的下钻页。", "不计弹窗、抽屉、页内 Tab、空壳页。"),
        ("计功能点", "PRD 功能清单 / 功能按钮，一条算一项。", "未出清单的模块写「—」，不估算。"),
    ]
    for i, (title, a, b) in enumerate(cards):
        x = Inches(0.45) + i * Inches(4.2)
        add_round(s, x, Inches(1.18), Inches(4.0), Inches(2.15), CARD, LINE)
        add_box(s, x, Inches(1.18), Inches(0.1), Inches(2.15), BLUE)
        add_tb(s, x + Inches(0.28), Inches(1.32), Inches(3.55), Inches(0.4),
               [(title, 16, NAVY, True)])
        add_tb(s, x + Inches(0.28), Inches(1.78), Inches(3.55), Inches(1.35),
               [(a, 13, INK, False), (b, 13, MUTED, False)])

    add_round(s, Inches(0.45), Inches(3.52), Inches(12.4), Inches(3.48), CARD, LINE)
    add_tb(s, Inches(0.7), Inches(3.66), Inches(12), Inches(0.36),
           [("汇报时建议这样说", 16, NAVY, True)])
    notes = [
        "覆盖规模：六个模块 87 个菜单入口、98 张独立原型页。",
        "已文档化功能：培养方案 54 + 开课 68 + 学籍 67，共 189 条。",
        "开课 68 只覆盖 16 个菜单中的 10 个；选修三页、教师替换、排课计划、流程图尚未出清单。",
        "基础数据、选课、排课目前没有同口径功能清单，所以功能点列为空，不凑整数。",
        "选课比 8 月年中稿多 1 个管理端入口（12→13）；其余菜单和页数与当时一致。",
    ]
    y = Inches(4.1)
    for t in notes:
        add_box(s, Inches(0.75), y + Inches(0.1), Inches(0.1), Inches(0.1), BLUE)
        add_tb(s, Inches(1.05), y, Inches(11.5), Inches(0.42), [(t, 13, INK, False)])
        y += Inches(0.5)
    footer(s, 5)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide_cover(prs)
    slide_summary(prs)
    slide_academic(prs)
    slide_foundation(prs)
    slide_caliber(prs)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
