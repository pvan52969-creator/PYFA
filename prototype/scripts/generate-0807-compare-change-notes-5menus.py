#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""为 5 个已有 20260902V2 菜单生成「对照开课0807」变更说明 md+docx。"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "参考文档" / "0、模板" / "需求调整变更说明模板（简版）20260901V3.docx"
BASE = ROOT / "参考文档" / "2、开课管理"
CHANGE_DATE = "2026-09-04"

# menu -> (folder relative to BASE, 0807 from date+ver, to date+ver, path label, prototype, rows)
JOBS = {
    "开课时间设置": {
        "dir": BASE / "01_开课设置/01_开课时间设置/开课时间设置20260902V2",
        "from_tag": "20260804V2",
        "to_tag": "20260902V2",
        "path": "开课管理 → 开课设置 → 开课时间设置",
        "prototype": "prototype/index.html → page-course-time-setting",
        "from_file": "开课0807/开课设置/开课时间设置/开课时间设置20260804V2",
        "to_file": "开课时间设置20260902V2.docx",
        "rows": [
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "其他", "obj_id": "—", "where": "文档结构 / PRD 模板", "type": "修改", "content": "开课0807（20260804V2 旧结构）→ 现行按 PRD 模板 V3.1 完整重写，落档界面树 P01-*、FD/F/BR 对象ID", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "字段", "obj_id": "FD001", "where": "P01-F01 查询区", "type": "删除", "content": "筛选「默认展示学期」：0807 有全部/是/否 → 现行仅保留学年学期（FD001），不再列默认展示筛选项", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "字段", "obj_id": "FD010–FD014", "where": "P01-L01 学期开课时间列表", "type": "修改", "content": "列表字段矩阵重排：0807 含教学周数、窗口状态、开课时间范围合并展示 → 现行拆为开课开始/截止（FD011–012）等，教学周数/窗口状态未再入现行 FD 表", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "字段", "obj_id": "FD024 / BR001 / UC001", "where": "默认展示学期", "type": "澄清", "content": "0807 写死「全局最多一个是」且列表开关排他 → 现行默认改表单勾选（FD024），BR001 写「是否全局唯一以原型为准」，UC001 未确认", "status": "未确认"},
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "规则", "obj_id": "BR003", "where": "备注列 / 表单备注", "type": "修改", "content": "0807 修改弹窗含备注（≤200）→ 现行 BR003：列表备注 HTML 已注释，接口不得强制备注必填；不可见则不作为必填", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "界面", "obj_id": "P01-M02 / FD030–FD032", "where": "新增所属部门覆盖", "type": "修改", "content": "单位覆盖：0807 为展开子表行内编辑 + 多选开课单位弹窗 → 现行独立弹窗 P01-M02（所属部门下拉 + 部门开课起止 FD030–032）；用语「开课单位」→「所属部门」", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 开课时间设置", "obj_type": "规则", "obj_id": "—", "where": "教学周数派生", "type": "澄清", "content": "0807 明确列表展示派生教学周（2月5周/4·9月14周）且弹窗文案与算法不一致待统一 → 现行 FD 表未单列教学周；校历周数口径仍需与校历决议对齐 UC", "status": "未确认"},
        ],
    },
    "特殊课程设置": {
        "dir": BASE / "01_开课设置/03_特殊课程设置/特殊课程设置20260902V2",
        "from_tag": "20260804V2",
        "to_tag": "20260902V2",
        "path": "开课管理 → 开课设置 → 特殊课程设置",
        "prototype": "prototype/index.html → page-special-course-settings",
        "from_file": "开课0807/开课设置/特殊课程设置/特殊课程设置20260804V2",
        "to_file": "特殊课程设置20260902V2.docx",
        "rows": [
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "其他", "obj_id": "—", "where": "文档结构 / PRD 模板", "type": "修改", "content": "开课0807（20260804V2）→ 现行 V3.1 完整重写，落档 P01 / P01-D01 分区与 FD/F/BR", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "字段", "obj_id": "FD023 / BR013", "where": "P01-D01 学时类型设置", "type": "新增", "content": "抽屉新增学时类型设置（全部合并/部分合并/分别安排，默认 separate）；0807 抽屉无此控件，仅四学时教室行", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "字段", "obj_id": "FD028 / FD018", "where": "共同授课课号", "type": "新增", "content": "抽屉/列表新增共同授课课号维护与展示；0807 需求稿无共同授课配置", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "规则", "obj_id": "BR004 / FD026", "where": "Support历史", "type": "修改", "content": "0807 Support历史可手改 Y/N → 现行只读，由默认 Support 专业勾选自动推导（勾选→Y，未勾选→N，不可手改）", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "字段", "obj_id": "FD017 / FD027", "where": "默认 Support 对象", "type": "修改", "content": "用语与对象：0807「默认 Support 学院」→ 现行「默认 Support 专业 / Support Programmes」；列表同步改称", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "字段", "obj_id": "FD010–FD018 / FD003", "where": "列表与查询", "type": "修改", "content": "列表：0807 不展示教室摘要 → 现行增加学时类型设置/教室信息/所属专业/共同授课等列；查询新增课名（FD003）", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "其他", "obj_id": "—", "where": "§2.3 非目标", "type": "澄清", "content": "现行明确：学分/总学时/计划人数/预留/选课方式/是否排课排场地等不在本抽屉，禁止从开课计划抄入；0807 本页本就不含这些字段", "status": "已明确"},
            {"module": "开课管理 → 开课设置 → 特殊课程设置", "obj_type": "规则", "obj_id": "BR001 / BR002 / UC001–UC002", "where": "删除约束 / 带出覆盖", "type": "澄清", "content": "0807 对删除不回滚已落班、新建/onlyIfEmpty 带出有较完整叙述 → 现行 BR001/BR002 标未确认（已被开课引用能否删、是否覆盖已改字段）", "status": "未确认"},
        ],
    },
    "开课计划": {
        "dir": BASE / "02_专业开课/01_开课计划/开课计划20260902V2",
        "from_tag": "20260804V4",
        "to_tag": "20260902V2",
        "path": "开课管理 → 专业开课 → 开课计划",
        "prototype": "prototype/index.html → page-course-offering-major",
        "from_file": "开课0807/专业开课/开课计划/开课计划20260804V4",
        "to_file": "开课计划20260902V2.docx",
        "rows": [
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "其他", "obj_id": "—", "where": "文档结构 / PRD 模板", "type": "修改", "content": "开课0807（20260804V4）→ 现行 V3.1 完整重写，界面树含 P01-M01/M02、P01-D01/D02 及 FD/F/BR", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "界面", "obj_id": "P01-M02 / F004", "where": "特殊开课弹窗", "type": "新增", "content": "工具栏新增「特殊开课」入口（P01-M02）；0807 非目标写「不含通识选修/特殊开课计划」", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "字段", "obj_id": "FD015 / FD017 / FD040", "where": "列表标记 / 选课方式", "type": "修改", "content": "列表新增「标记」专业开课/特殊开课（FD015）；选课类型展示口径由开放/不开放 → 现行必修/选修（选课方式 FD017/FD040）", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "字段", "obj_id": "FD031 / FD050 / FD053–FD054 / BR010–BR013", "where": "P01-D01 修改教学任务抽屉", "type": "新增", "content": "抽屉相对 0807 补：课程班名称（FD031）、学时类型设置 hoursMode（FD050/BR013）、选修「可选专业批次」与必修「上课专业批次」分流（BR010）", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "规则", "obj_id": "BR012 / FD036–FD038", "where": "计划总人数 / 预留 / 名额上限", "type": "修改", "content": "0807：开放选课计划总人数可手填；预留≥0 → 现行：计划总人数一律只读（BR012）；预留默认5且可为负；名额上限=计划+预留且须>0", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "其他", "obj_id": "—", "where": "工具栏隐藏项", "type": "澄清", "content": "现行非目标/界面说明：ME开课、撤回、导出、授课确认、学院筛选在原型已隐藏，禁止当本页生效入口；0807 筛选学院亦暂隐", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "规则", "obj_id": "BR004 / UC002", "where": "删除约束", "type": "澄清", "content": "0807 硬规则：仅非已生效可删，已生效须先安排退回 → 现行 BR004「已生效或已被安排引用是否可删」未确认", "status": "未确认"},
            {"module": "开课管理 → 专业开课 → 开课计划", "obj_type": "规则", "obj_id": "BR002 / UC001", "where": "合班人数累加", "type": "澄清", "content": "0807 §7.10 写明合班后人数/预留累加等副作用 → 现行合班一致性已明确，但合班后总人数/预置累加口径终裁仍 UC001", "status": "未确认"},
        ],
    },
    "开课安排": {
        "dir": BASE / "02_专业开课/02_开课安排/开课安排20260902V2",
        "from_tag": "20260804V4",
        "to_tag": "20260902V2",
        "path": "开课管理 → 专业开课 → 开课安排",
        "prototype": "prototype/index.html → page-course-major-offering-task-style2",
        "from_file": "开课0807/专业开课/开课安排/开课安排20260804V4",
        "to_file": "开课安排20260902V2.docx",
        "rows": [
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "界面", "obj_id": "P01-S01 / P01-S02 / BR004", "where": "向导步骤条", "type": "修改", "content": "0807 三步（开课前期设置 / 课程班·合班Support / 师资安排）→ 现行两步 Course Setting + 师资安排；旧三步作废（BR004）", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "字段", "obj_id": "FD048 / FD077 / BR011", "where": "联动安排", "type": "新增", "content": "展开面板与安排详情新增「联动安排」列及共同授课组联动规则（BR011）；0807 展开仅有同时多组授课，无联动列", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "规则", "obj_id": "BR010 / FD076", "where": "同时授课", "type": "修改", "content": "0807 同时授课为是/否条件必选 → 现行须手动选是/否、新增默认空、禁止自动默认；多组选否则拆行（BR010 tip）", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "规则", "obj_id": "BR001", "where": "安排完整度 / 生效校验", "type": "修改", "content": "0807 §7.1 概括硬拦/警示 → 现行 BR001 分点清单（硬阻断1–5、二次确认6–7、副作用8–10、保存vs生效11–12），含停课/选修人数/授课确认分教师等", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "界面", "obj_id": "P02-A01 / P02-W01 / P02-M02", "where": "分组工作台子区块", "type": "修改", "content": "工作台侧栏动作/Teaching Load/课时详情（含 No. of Groups、Latest Evaluation 占位 BR014–015）按现行可见 UI 落档；相对 0807 字段表更强调嵌入区与 tip", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "字段", "obj_id": "FD012 / FD013 / FD020", "where": "P01-L01 安排列表", "type": "修改", "content": "列表：Course Name→课程班名称；新增标记；小组数 0807 写「列表暂隐」→ 现行 FD020 展示；Support 管理不再作为独立步骤列集主路径", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "其他", "obj_id": "—", "where": "非目标 / 名单 Tab", "type": "澄清", "content": "现行非目标：停课主入口在课程班；学生名单 Tab 默认隐藏（从名单管理进入另见开课名单）；授课确认教师端深度见独立菜单", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课安排", "obj_type": "规则", "obj_id": "BR003 / UC002", "where": "撤回 vs 退回", "type": "澄清", "content": "0807 §7.2 对撤回/退回副作用有完整对照表 → 现行 BR003 写「撤回/退回差异及备注未完全确认」，UC002 仍开", "status": "未确认"},
        ],
    },
    "开课名单": {
        "dir": BASE / "02_专业开课/03_开课名单/开课名单20260902V2",
        "from_tag": "20260804V2",
        "to_tag": "20260902V2",
        "path": "开课管理 → 专业开课 → 开课名单",
        "prototype": "prototype/index.html → page-course-major-offering-roster",
        "from_file": "开课0807/专业开课/开课名单/开课名单20260804V2",
        "to_file": "开课名单20260902V2.docx",
        "rows": [
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "其他", "obj_id": "—", "where": "文档结构 / PRD 模板", "type": "修改", "content": "开课0807（20260804V2）→ 现行 V3.1 完整重写；明确 P01 列表 + P02 名单工作台（共用 page-offering-grouping 学生名单模式）", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "字段", "obj_id": "FD010–FD012 / FD015–FD016 / FD027–FD028", "where": "P01-L01 开课任务列表", "type": "新增", "content": "列表相对 0807 增补：生效状态、名单状态、分配情况、课程班名称、标记、共同授课状态/课号、所属专业等", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "功能", "obj_id": "F015 / P02-M02", "where": "预重修名单", "type": "新增", "content": "工作台按钮新增「预重修名单」；0807 功能清单无此入口", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "规则", "obj_id": "BR010 / P02-M01", "where": "一键分配", "type": "修改", "content": "0807 §7.2 简述维度/模式 → 现行 BR010 按 tip 全条：维度勾选+优先级连续正整数、跳过手动添加、遵守组上限、重置分组清空规则等", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "规则", "obj_id": "BR001 / FD048", "where": "可入名单学籍", "type": "澄清", "content": "现行明确可入名单学籍为 Active+Normal（BR001）；0807 字段表有学籍状态列但未写此硬口径", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "规则", "obj_id": "BR002 / UC001", "where": "退回后名单保留", "type": "澄清", "content": "0807 写退回安排后「已维护名单保留」→ 现行 BR002「退回后名单保留策略：未确认」（UC001）；与安排页撤回重置小组分配的对照待终裁", "status": "未确认"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "其他", "obj_id": "—", "where": "查询区隐藏项", "type": "澄清", "content": "授课确认状态筛选：0807 与现行均标明原型已注释隐藏，禁止当生效筛选项实现", "status": "已明确"},
            {"module": "开课管理 → 专业开课 → 开课名单", "obj_type": "其他", "obj_id": "—", "where": "§12 名单/分配状态枚举", "type": "澄清", "content": "现行列表有名单状态/分配情况展示，但枚举码「未单独文档化处标未确认」；0807 写本页无独立名单提交工作流", "status": "未确认"},
        ],
    },
}


def set_cell_text(cell, text: str):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text or ""))
    run.font.size = Pt(10.5)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def fill_kv(table, mapping: dict[str, str]):
    for row in table.rows:
        label = (row.cells[0].text or "").strip().replace("\n", "")
        for key, val in mapping.items():
            if key in label:
                set_cell_text(row.cells[1], val)
                break


def write_md(menu: str, job: dict, out: Path):
    if out.exists():
        raise SystemExit(f"拒绝覆盖：{out}")
    rows = job["rows"]
    lines = [
        f"# {menu}——需求调整变更说明",
        "",
        f"> 模板：`参考文档/0、模板/需求调整变更说明模板（简版）20260901V3.docx`",
        f"> **对比基线（开课0807）：** `{job['from_file']}`",
        f"> **本版：** `{job['to_file']}`（正式链现行有效版）",
        "",
        "## 1 版本信息",
        "",
        "| 项目 | 内容 |",
        "|------|------|",
        f"| 需求文档名称 | 厦大马来分校本科教务系统产品需求文档 — {menu} |",
        f"| 菜单路径 | {job['path']} |",
        f"| 上一有效版 → 本版 | 开课0807·{job['from_tag']} → {job['to_tag']} |",
        f"| 对比基线文件 | {job['from_file']} |",
        f"| 本版文件 | {job['to_file']} |",
        "| 本版 PRD 模板 | V3.1 |",
        f"| 原型地址 | {job['prototype']} |",
        f"| 变更日期 | {CHANGE_DATE} |",
        "",
        "## 2 本次改了什么（总览）",
        "",
        "| 序号 | 模块（菜单路径） | 对象类型 | 对象ID | 功能 / 位置（可读名） | 变更类型 | 调整内容（从什么变成什么） | 状态 |",
        "|------|------------------|----------|--------|---------------------|----------|------------------------------|------|",
    ]
    for i, r in enumerate(rows, 1):
        lines.append(
            f"| {i} | {r['module']} | {r['obj_type']} | {r['obj_id']} | {r['where']} | "
            f"{r['type']} | {r['content']} | {r['status']} |"
        )
    lines += [
        "",
        "## 3 前后对照（可选）",
        "",
        "详见上表「调整内容」；复杂字段以现行 PRD 对象ID 为准。",
        "",
        "## 4 连带影响（可选）",
        "",
        "| 序号 | 影响到哪里 | 影响说明 | 需同步 |",
        "|------|------------|----------|--------|",
        "| 1 | 开课0807 历史稿 | 仅作对比基线，不覆盖开课0807 文件夹内文件 | — |",
        "| 2 | 原型 / 开发 / 测试 | 以现行 20260902V2 与可见 UI 为准 | ☑原型 ☑开发 ☑测试 |",
        "",
    ]
    out.write_text("\n".join(lines), encoding="utf-8")


def write_docx(menu: str, job: dict, out: Path):
    if out.exists():
        raise SystemExit(f"拒绝覆盖：{out}")
    if not TEMPLATE.exists():
        raise SystemExit(f"缺少模板：{TEMPLATE}")
    doc = Document(str(TEMPLATE))
    fill_kv(
        doc.tables[0],
        {
            "需求文档名称": f"厦大马来分校本科教务系统产品需求文档 — {menu}",
            "菜单路径": job["path"],
            "上一有效版": f"开课0807·{job['from_tag']} → {job['to_tag']}",
            "上一版 → 本版": f"开课0807·{job['from_tag']} → {job['to_tag']}",
            "上一版文件": job["from_file"],
            "本版文件": job["to_file"],
            "本版 PRD 模板": "V3.1",
            "原型地址": job["prototype"],
            "变更日期": CHANGE_DATE,
        },
    )
    t1 = doc.tables[1]
    rows = job["rows"]
    while len(t1.rows) - 1 < len(rows):
        t1.add_row()
    data_rows = t1.rows[1:]
    for i, row in enumerate(data_rows):
        if i >= len(rows):
            for c in row.cells:
                set_cell_text(c, "")
            continue
        r = rows[i]
        vals = [
            str(i + 1),
            r["module"],
            r["obj_type"],
            r["obj_id"],
            r["where"],
            r["type"],
            r["content"],
            r["status"],
        ]
        for cell, val in zip(row.cells, vals):
            set_cell_text(cell, val)
    doc.save(out)


def main():
    for menu, job in JOBS.items():
        stem = f"{menu}{job['from_tag']}→{job['to_tag']}变更说明"
        # 文件名与规范一致：含日期，如 开课时间设置20260804V2→20260902V2变更说明
        out_dir = job["dir"]
        if not out_dir.is_dir():
            raise SystemExit(f"目录不存在：{out_dir}")
        md_out = out_dir / f"{stem}.md"
        docx_out = out_dir / f"{stem}.docx"
        write_md(menu, job, md_out)
        write_docx(menu, job, docx_out)
        print(f"OK {menu}: {md_out.name} + docx ({len(job['rows'])} rows)")


if __name__ == "__main__":
    main()
