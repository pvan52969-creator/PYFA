#!/usr/bin/env python3
"""按 XMUM 产品需求文档模板结构，为各二级菜单生成独立 PRD docx。"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from prd_folder_paths import BASE as BASE_OUT, menu_dir, version_dir
from prd_submenu_fields import FIELD_COL_WIDTHS, FIELD_HEADERS, MENU_FIELD_PACKS

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent

# 二级菜单 → 一级菜单（母文件夹键名，路径见 prd_folder_paths）
MENU_PARENT = {
    "开课时间设置": "开课设置",
    "特殊课程设置": "开课设置",
    "校选课程管理": "开课设置",
    "开课计划": "专业开课",
    "开课安排": "专业开课",
    "开课名单": "专业开课",
    "课程班": "课程班管理",
}

LANDSCAPE_SCALE = 1.35
FIELD_TABLE_COL_WIDTHS = FIELD_COL_WIDTHS


def fn(num, cn, en, desc, interaction, notes, drill):
    return (num, cn, en, desc, interaction, notes, drill)


def f(*cells):
    """兼容 DOCS 内遗留摘要字段表写法（生成时会被 MENU_FIELD_PACKS 全量表覆盖）。"""
    return list(cells)


def scale_widths(widths):
    if isinstance(widths, tuple):
        return tuple(round(w * LANDSCAPE_SCALE, 1) for w in widths)
    return [round(w * LANDSCAPE_SCALE, 1) for w in widths]


def set_document_landscape(doc):
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return h


def add_kv_table(doc, rows, col_widths=(5, 18.5)):
    col_widths = scale_widths(col_widths)
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_col_widths(table, col_widths)
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = k
        c1.text = v
        for cell in (c0, c1):
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if c0.paragraphs[0].runs:
            c0.paragraphs[0].runs[0].bold = True
    doc.add_paragraph()
    return table


def add_grid_table(doc, headers, rows, col_widths=None):
    if col_widths is None:
        col_widths = None
    else:
        col_widths = scale_widths(col_widths) if max(col_widths) < 10 else col_widths
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            cells[ci].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[ci].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if col_widths:
        set_col_widths(table, col_widths)
    doc.add_paragraph()
    return table


def add_field_table(doc, title, rows):
    add_para(doc, title, bold=True)
    add_grid_table(doc, FIELD_HEADERS, rows, col_widths=FIELD_TABLE_COL_WIDTHS)


def add_function_item(doc, num, cn_name, en_name, desc, interaction, notes, drill_down):
    title = f"{num}、功能按钮——{cn_name}（英文名称：{en_name}）"
    add_para(doc, title, bold=True)
    add_para(doc, "a. 功能说明（描述、业务的事件交互、备注信息等）：")
    add_para(doc, f"· 描述：{desc}")
    add_para(doc, f"· 业务的事件交互：{interaction}")
    add_para(doc, f"· 备注信息（校验规则补充、其他说明等）：{notes}")
    add_para(doc, f"b. 下钻页面说明（备注说明）：{drill_down}")
    doc.add_paragraph()


def add_menu_block(
    doc,
    section_no,
    title,
    status,
    intro,
    list_fields,
    search_fields,
    flow_rel,
    flow_pre,
    flow_out,
    biz_flow,
    proto_link="",
    field_tables=None,
    functions=None,
    non_goals="",
    pending="",
    field_shared_note="",
    field_hidden_note="",
):
    add_heading(doc, f"{section_no}{title}（需求确认状态：{status}）", level=4)
    add_para(doc, "菜单介绍", bold=True)
    add_kv_table(doc, [("菜单介绍", ""), ("菜单内容简介", intro)])
    add_para(doc, "页面展示字段信息", bold=True)
    add_kv_table(doc, [("页面展示字段信息", list_fields)])
    add_para(doc, "支持查询检索的字段信息", bold=True)
    add_kv_table(doc, [("支持查询检索的字段信息", search_fields)])
    add_para(doc, "数据前后流转关系（说明、前置条件、下游输出等）", bold=True)
    add_kv_table(
        doc,
        [
            ("1）关系说明", flow_rel),
            ("2）前置条件", flow_pre),
            ("3）下游输出", flow_out),
        ],
    )
    add_para(doc, "5、业务流关系（操作流程）", bold=True)
    add_kv_table(doc, [("5、业务流关系（操作流程）", biz_flow)])
    if proto_link:
        add_para(doc, "6、原型参考链接", bold=True)
        add_kv_table(doc, [("原型参考链接", proto_link)])
    if field_tables:
        add_para(doc, "新增—字段信息表（按所在界面/区块列全）", bold=True)
        add_para(
            doc,
            "说明：下列字段表覆盖筛选区、列表、各弹窗与抽屉分区（含 Tab/步骤差异列）；"
            "每行「所在界面/区块」标明字段归属。",
        )
        if field_shared_note:
            add_para(doc, field_shared_note)
        for ft_title, ft_rows in field_tables:
            add_field_table(doc, ft_title, ft_rows)
        if field_hidden_note:
            add_para(doc, "本期隐藏/注释字段（不在主界面展示）", bold=True)
            add_para(doc, field_hidden_note)
    if functions:
        add_para(doc, "（3）菜单功能清单", bold=True)
        for item in functions:
            add_function_item(doc, *item)
    if non_goals:
        add_para(doc, "非目标", bold=True)
        add_para(doc, non_goals)
    if pending:
        add_para(doc, "待确认事项", bold=True)
        add_para(doc, pending)


# ─── 各菜单内容 ───────────────────────────────────────────────

DOCS = [
    {
        "folder": "开课时间设置",
        "filename": "开课时间设置20260803V1.docx",
        "meta": {
            "cn": "开课时间设置",
            "en": "Offering Time Setting",
            "parent_cn": "开课设置",
            "parent_en": "Offering Settings",
            "path": "开课管理 → 开课设置 → 开课时间设置",
            "version": "V1",
            "date": "2026 年 8 月 3 日",
            "purpose": "明确学年学期维度的开课操作时间窗口、默认展示学期、授课确认截止日期及单位级时间覆盖规则，作为专业开课等下游模块的前置配置。",
            "background": "开课业务依赖学期窗口与默认展示学期；本菜单维护全校及单位级开放时间，教学周按学期类型固定推导。",
        },
        "menu": dict(
            section_no="2.2.1.1 ",
            title="开课时间设置（英文名称：Offering Time Setting）",
            status="已确认",
            intro="按学年学期配置全校开课开放时间、默认展示学期、授课确认截止日期；支持按开课单位覆盖单位级开放时间（优先于全校）。"
            "支持新增学期配置（弹窗选择未配置学期并填写窗口与确认截止）。"
            "教学周数不在本页手工维护，按学期类型固定推导（原型当前：2 月短学期 5 周，4/9 月长学期 14 周）。",
            list_fields="勾选、展开（单位开课时间）、学年学期、教学周数、状态、开课时间范围、授课确认截止日期、默认展示学期（开关）、操作（修改 / 单位开课时间设置）。"
            "展开子表：序号、开课单位、状态、开放开始、开放截止、操作。",
            search_fields="学年学期、默认展示学期（是/否）；查询/重置。",
            flow_rel="配置「默认展示学期=是」后，专业开课「生成开课任务」方可进行；单位有独立时间则优先于全校；授课确认截止日期供教师端确认函 Notes。",
            flow_pre="已维护校历学年学期。",
            flow_out="开课操作窗口、默认展示学期、授课确认截止日期；供开课计划/安排/名单等学期筛选与下游确认信函。",
            biz_flow="进入【开课时间设置】→ 查询列表 → 修改全校窗口/确认截止 → 按需展开维护单位开课时间 → 切换默认展示学期（全局最多一个「是」）。",
            proto_link="prototype/index.html → page-course-time-setting",
            field_tables=[
                (
                    "开课时间设置——字段",
                    [
                        f("1", "学年学期", "Academic Term", "下拉框", "是", "—", "来自校历学期代码集", "是", "2026/04"),
                        f("2", "教学周数", "Teaching Weeks", "派生只读", "—", "按学期类型固定", "本页不可改；短学期5周/长学期14周（以校历决议为准）", "否", "14"),
                        f("3", "状态", "Window Status", "派生", "—", "—", "相对当前时刻：开放中/已截止", "是", "开放中"),
                        f("4", "开课开始", "School Open From", "日期时间", "是", "须早于截止", "dd/mm/yyyy HH:mm", "否", "02/03/2026 08:00"),
                        f("5", "开课截止", "School Open To", "日期时间", "是", "须晚于开始", "超出校历结束日可确认后保存", "否", "02/04/2026 18:00"),
                        f("6", "默认展示学期", "Default Display Term", "开关", "是", "全局最多一个「是」", "关掉当前项可为全部「否」", "是", "是"),
                        f("7", "授课确认截止日期", "Teacher Confirm Deadline", "日期", "否", "—", "仅日期；可过期；不随开课窗口锁定", "否", "15/01/2026"),
                        f("8", "备注", "Remark", "文本", "否", "≤200", "—", "否", ""),
                        f("9", "开课单位", "Offering Unit", "代码", "条件", "—", "单位覆盖行", "是", "AC"),
                        f("10", "单位开放开始/截止", "Unit Open From/To", "日期时间", "条件", "—", "有配置则优先于全校", "否", ""),
                    ],
                ),
            ],
            functions=[
                fn("1", "查询/重置", "Query/Reset", "按条件过滤列表。", "筛选区按钮。", "—", "无。"),
                fn("2", "新增", "Add", "新增学年学期开课时间配置。", "工具栏打开弹窗：选择未配置学期、填写开课起止与授课确认截止后保存。", "同学期不可重复；开始须早于截止。", "弹窗：新增开课时间。"),
                fn("3", "修改", "Edit", "修改全校开课窗口、授课确认截止、备注。", "行操作打开弹窗保存。", "开始须早于截止。", "弹窗：修改开课时间。"),
                fn("4", "删除", "Delete", "删除选中学期配置。", "多选确认。", "不可撤销。", "无。"),
                fn("5", "默认展示学期开关", "Toggle Default Term", "设为默认展示学期。", "列表开关。", "排他：开一个则关掉其他。", "无。"),
                fn("6", "单位开课时间设置", "Unit Offering Window", "维护单位级覆盖时间。", "展开区新增/行内编辑/删除。", "未配置继承全校。", "展开子表。"),
            ],
            non_goals="不维护校历本身与考试周；不替代执行计划锁定；不含操作流程图、校选课程管理。",
            pending="编辑弹窗文案若仍提示周数（6/17）与原型数据推导（5/14）不一致，正式以校历决议统一。",
        ),
    },
    {
        "folder": "特殊课程设置",
        "filename": "特殊课程设置20260803V1.docx",
        "meta": {
            "cn": "特殊课程设置",
            "en": "Special Course Settings",
            "parent_cn": "开课设置",
            "parent_en": "Offering Settings",
            "path": "开课管理 → 开课设置 → 特殊课程设置",
            "version": "V1",
            "date": "2026 年 8 月 3 日",
            "purpose": "提供课号级教室属性/偏好、默认 Support 学院与 Support 历史配置，供专业开课安排新建或仍为系统默认的教学班自动带出。",
            "background": "同一课号的教室与 Support 信息跨学期变化少，若每班重复维护成本高、易漏设；故在开课设置侧集中维护课号默认。",
        },
        "menu": dict(
            section_no="2.2.1.2 ",
            title="特殊课程设置（英文名称：Special Course Settings）",
            status="已确认",
            intro="从教务课程库手动纳入课号，维护四学时教室属性/偏好、默认 Support 学院、Support 历史（Y/N）。"
            "默认仅影响此后新建或仍为系统默认的教学班；已改过的班不覆盖；班级覆盖不回写课号配置。通识等可复用同一数据源。",
            list_fields="勾选、序号、课程代码、课程名称、开课单位、Support 学院、Support历史、操作（编辑）。教室仅在编辑抽屉维护，列表无教室摘要列。",
            search_fields="开课单位、课程号；查询/清除过滤。",
            flow_rel="清单内课号在专业开课安排新建/仅系统默认时带出教室与 Support；Support历史始终按课号同步（未配置为 N）。",
            flow_pre="教务课程库已存在对应课程。",
            flow_out="开课安排编辑抽屉教室区、Support、Support历史只读展示；通识等预留同一读取接口。",
            biz_flow="进入【特殊课程设置】→ 添加课程纳入清单 → 编辑教室/Support/历史 → 下游开课安排自动带出可覆盖项。",
            proto_link="prototype/index.html → page-special-course-settings",
            field_tables=[
                (
                    "特殊课程设置——字段",
                    [
                        f("1", "课程代码", "Course Code", "文本", "是", "课号唯一", "来自教务课程库", "否", "ACC301"),
                        f("2", "课程名称", "Course Name", "只读", "是", "—", "随课程库", "否", ""),
                        f("3", "开课单位", "Offering Unit", "只读", "是", "—", "—", "是", "AC"),
                        f("4", "学时类型", "Hour Type", "枚举", "是", "L/T/P/O", "理论/辅导/实践/其他", "是", "L"),
                        f("5", "教室属性", "Classroom Attribute", "下拉", "是", "—", "默认普通", "是", "普通"),
                        f("6", "参考教室偏好", "Reference Classrooms", "多选/文本", "否", "—", "偏好提示非强制占用", "否", ""),
                        f("7", "Support学院", "Support Units", "多选", "否", "不可含本开课单位", "—", "是", ""),
                        f("8", "Support历史", "Support History", "Y/N", "是", "—", "未配置课号下游按 N", "是", "N"),
                    ],
                ),
            ],
            functions=[
                fn("1", "查询/清除过滤", "Query/Clear", "按单位、课号过滤。", "筛选区。", "—", "无。"),
                fn("2", "添加课程", "Add Courses", "从课程库纳入清单。", "弹窗多选确定。", "课号去重；新纳入默认全普通无偏好、Support空、历史N。", "弹窗：添加特殊课程。"),
                fn("3", "编辑", "Edit", "维护教室、Support、历史。", "行操作打开抽屉。", "—", "抽屉：编辑特殊课程设置。"),
                fn("4", "删除", "Delete", "移出维护清单。", "多选确认。", "不回滚已落班数据。", "无。"),
            ],
            non_goals="不替换校选课程管理；不强制全库先配置才能开课；不改排课冲突算法；通识全链路 UI 默认同批落地可后续。",
            pending="列表是否增加教室摘要列；Support历史是否按学期细分。",
        ),
    },
    {
        "folder": "校选课程管理",
        "filename": "校选课程管理20260804V1.docx",
        "meta": {
            "cn": "校选课程管理",
            "en": "School Elective Course Management",
            "parent_cn": "开课设置",
            "parent_en": "Offering Settings",
            "path": "开课管理 → 开课设置 → 校选课程管理",
            "version": "V1",
            "date": "2026 年 8 月 4 日",
            "purpose": "维护校选课程主数据池（GE/ME）、修读范围与开课默认教学设置，供通识选修开课选用并带出默认值。",
            "background": "通识选修开课依赖校选课库；修读范围与教学开关需在开课前集中维护，计划侧可按学期覆盖修读范围。",
        },
        "menu": dict(
            section_no="2.2.1.3 ",
            title="校选课程管理（英文名称：School Elective Course Management）",
            status="已确认",
            intro="维护 GE/ME 校选课清单：新增 ME（课库 Major Elective，去重）、批量改状态、删除（仅移出校选池不删课库）；"
            "GE 由 Restrictions List 种子进入，本页不可新增。"
            "修读范围：可选/不可选专业互斥（皆空=全开放）；ME 初始化可选=开课单位下专业；GE arts/business/science 初始化不可选=同大类（science 剔除 MAT）；unrestricted 全开放。"
            "可选学生类型至少1种；不可选年级学期 YnSm（空=均可）；冲突取自校选库、先修取自课库+校选，均排除本课。"
            "开课默认：Support历史默认N；排课/排场地/考勤/录入成绩/排考无值默认是；选课类型固定开放选课；备注带入通识计划行。"
            "下游通识开课仅消费 GE+正常状态；同学期一门课仅可开一次；学期侧可覆盖修读范围（任务生效后锁定）。"
            "本页不维护计划人数/预留/起止周/分组/批次级范围；条件可选红星无UI。",
            list_fields="勾选、序号、课程代码、课程名称、类型、校选课类别、学分、总学时、开课单位、状态、备注、修读范围（查看/设置）、课程信息（修改）。",
            search_fields="课程管理单位、类型、校选课类别、状态、课程代码/名称；查询/清除过滤。",
            flow_rel="GE 种子/ME 课库纳入 → 维护范围与默认 → 通识选修开课添加候选并带出默认；学期侧可覆盖修读范围。",
            flow_pre="教务课程库存在；GE 有 Restrictions List（或等价同步）。",
            flow_out="通识选修计划/开课的候选课库、修读范围默认、教学开关默认、计划行备注。",
            biz_flow="进入【校选课程管理】→ 查询/新增ME → 设置修读范围 → 修改开课默认 → 下游通识开课选用。",
            proto_link="prototype/index.html → page-school-elective-courses",
            field_tables=[
                (
                    "校选课程管理——关键字段",
                    [
                        f("1", "类型", "Elective Type", "枚举", "是", "GE/ME", "—", "是", "GE"),
                        f("2", "校选课类别", "GE Category", "枚举", "条件", "GE必填", "ME为空", "是", "Arts"),
                        f("3", "状态", "Status", "枚举", "是", "active/suspended", "停课不可入通识计划", "是", "正常"),
                        f("4", "修读范围", "Programme Scope", "名单", "是", "可选/不可选互斥", "皆空=全开放", "否", "全开放"),
                        f("5", "可选学生类型", "Student Types", "多选", "是", "≥1", "Local/Chinese/International", "是", "全部"),
                        f("6", "不可选年级学期", "Excluded Year-Semester", "多选", "否", "YnSm", "空=均可选", "是", "Y2S1"),
                        f("7", "Support历史", "Support History", "Y/N", "是", "—", "默认N", "是", "N"),
                        f("8", "是否排课等开关", "Teaching Switches", "是/否", "是", "—", "无值默认是", "是", "是"),
                    ],
                ),
            ],
            functions=[
                fn("1", "查询", "Query", "按筛选刷列表。", "筛选区。", "含排序分页。", "无。"),
                fn("2", "清除过滤", "Reset Filters", "清空筛选并重置分页。", "清除过滤。", "—", "无。"),
                fn("3", "新增ME", "Add ME", "从课库纳入Major Elective。", "弹窗。", "不可新增GE；按catalogId/课号去重。", "弹窗：新增。"),
                fn("4", "批量改状态", "Batch Status", "正常/停课。", "需勾选。", "停课不可入通识计划。", "弹窗：批量状态。"),
                fn("5", "删除", "Delete", "移出校选池。", "二次确认。", "不删课库；原型未拦下游引用。", "无。"),
                fn("6", "查看修读范围", "View Scope", "只读展示限制。", "列表链接。", "—", "弹窗：查看修读范围。"),
                fn("7", "设置修读范围", "Edit Scope", "专业/学生类型/年级学期/冲突先修。", "抽屉。", "可选不可选互斥；保存学生类型≥1。", "抽屉：修读范围。"),
                fn("8", "修改开课默认", "Edit Defaults", "Support/状态/开关/备注。", "抽屉。", "下游GE默认值；排场地/考勤不按学时全0改否。", "抽屉：开课默认。"),
                fn("9", "专业选择", "Programme Picker", "可选/不可选Tab与初始化。", "嵌套弹窗。", "两边不可同时有值。", "弹窗：专业选择。"),
                fn("10", "冲突/先修选择", "Course Link Picker", "多选课号写回。", "嵌套弹窗。", "排除本课。", "弹窗：冲突/先修。"),
            ],
            non_goals="不含专业开课生成；不含计划人数/起止周/分组；不含条件可选红星UI；不含批次级修读范围；不替代课库。",
            pending="已引用课删除策略；conditionalProgrammeKeys；GE List同步方式；备注是否拆分。",
        ),
    },
    {
        "folder": "开课计划",
        "filename": "开课计划20260803V1.docx",
        "meta": {
            "cn": "开课计划",
            "en": "Major Offering Plan",
            "parent_cn": "专业开课",
            "parent_en": "Major Offering",
            "path": "开课管理 → 专业开课 → 开课计划",
            "version": "V1",
            "date": "2026 年 8 月 3 日",
            "purpose": "规范专业开课计划层：从已锁定执行计划生成学期开课任务、维护教学任务与合班（未生效时）、批量生效后进入开课安排。",
            "background": "主路径为「计划 → 安排 → 名单」。计划层负责生成与生效；合班与分组职责分离。"
            "生成默认不合班，可选用「生成并自动合班」或工具栏「一键合班」（仅未生效草稿互合）。",
        },
        "menu": dict(
            section_no="2.2.2.1 ",
            title="开课计划（英文名称：Major Offering Plan）",
            status="已确认",
            intro="从已锁定执行计划筛选本学期未开课专业课，按「1课×1专业×1批次」生成计划行与教学班；支持修改教学任务、合班（未生效）、批量生效进入开课安排。"
            "生成默认不合班；确认框可选「生成并自动合班」，工具栏提供「一键合班」（均仅未生效草稿互合，不与已生效班合并）。"
            "已生效后本页合班锁定，合班改在开课安排；编辑抽屉锁定规则见字段表。"
            "学时分类只读；排场地/考勤按学时默认（可改）；计划总人数三口径；退回计划保留合班；已生效计划不可删须先回退。",
            list_fields="勾选、开课学期、负责人、生效状态、课程代码、课程名称、课程类别、选课类型、上课专业、上课批次、是否新生、学分、总学时、起止周、开课单位、预置人数、操作记录、课程信息、操作（合班/回退说明）。",
            search_fields="学年学期、专业、开课单位、课程号、生效状态（草稿/已生效/回退）；查询/重置。",
            flow_rel="生成候选=已锁定执行计划×本学期未开课专业课；生效后 submitStatus=submitted 进入开课安排；回退来自开课安排「退回」。",
            flow_pre="开课时间设置存在默认展示学期=是；对应专业批次执行计划已锁定；课程未开课。",
            flow_out="已生效计划进入开课安排；schedulingVisible 供下游可见性。",
            biz_flow="进入【开课计划】→ 生成开课任务（可选并自动合班）→ 修改教学任务 → 必要时合班/一键合班 → 勾选生效 → 进入【开课安排】。",
            proto_link="prototype/index.html → page-course-offering-major",
            field_tables=[
                (
                    "开课计划列表——关键字段",
                    [
                        f("1", "生效状态", "Submit Status", "状态", "是", "draft/submitted/reverted", "UI：草稿/已生效/回退", "是", "草稿"),
                        f("2", "开课学期", "Offering Term", "代码", "是", "—", "—", "是", "2026/04"),
                        f("3", "负责人", "Owner", "文本", "否", "—", "生成时记入当前用户", "否", ""),
                        f("4", "课程代码", "Course Code", "文本", "是", "—", "—", "否", "ACC301"),
                        f("5", "课程名称", "Course Name", "文本", "是", "—", "—", "否", ""),
                        f("6", "选课类型", "Enrollment Type", "下拉", "是", "—", "默认不开放选课", "是", "不开放"),
                        f("7", "上课专业", "Programme", "只读", "是", "—", "合班可多专业", "是", "FIN"),
                        f("8", "上课批次", "Intake", "只读", "是", "—", "合班可多批次", "是", "2024/09"),
                        f("9", "学分", "Credits", "数字", "是", "—", "同课对齐", "否", "3"),
                        f("10", "总学时", "Total Hours", "数字", "是", "—", "—", "否", "42"),
                        f("11", "起止周", "Week Range", "文本", "是", "落在学期教学周", "—", "否", "1-14"),
                        f("12", "开课单位", "Offering Unit", "只读", "是", "—", "—", "是", ""),
                        f("13", "计划总人数", "Planned Headcount", "数字", "—", "不可直接改", "老生=预置可修人数；新生=招生计划", "否", "96"),
                        f("14", "预留名额", "Reserved Seats", "数字", "否", "≥0", "名额上限=计划总人数+预留", "否", "0"),
                        f("15", "Support历史", "Support History", "只读", "—", "—", "来自特殊课程设置；未配置为N", "是", "N"),
                    ],
                ),
                (
                    "合班——约束字段",
                    [
                        f("1", "课程号", "Course Code", "只读", "是", "双方须相同", "同课号方可合班", "否", "ACC301"),
                        f("2", "学分/总学时/起止周/教学周数", "—", "只读", "是", "双方须相同", "—", "否", ""),
                        f("3", "教学任务开关", "Teaching Switches", "下拉", "是", "双方须相同", "选课类型、是否排课/排场地/考勤/录入成绩/排考；排场地与考勤默认随学时分类（学时全0→否，否则→是）", "是", ""),
                        f("4", "专业批次", "Prog. Batch", "穿梭", "是", "至少保留1个", "合入后挂同一教学班；人数预留累加；清空分组教师", "否", ""),
                    ],
                ),
            ],
            functions=[
                fn("1", "生成开课任务", "Generate Offering Tasks", "从执行计划生成计划行与教学班。", "弹窗勾选或一键生成；默认不合班；可选「生成并自动合班」。", "已存在课+专业+批次跳过；自动合班仅未生效草稿互合。", "弹窗：生成开课任务。"),
                fn("2", "修改教学任务", "Edit Teaching Task", "维护选课类型、周次、开关、教室等。", "课程信息列打开抽屉。", "已生效后主要保留教室可改；安排生效后整抽屉只读。", "抽屉：修改教学任务。"),
                fn("3", "合班", "Merge Sections", "合并同课多专业批次。", "操作列；未生效可用。", "已生效改开课安排。", "合拆班抽屉。"),
                fn("4", "一键合班", "One-click Merge", "按合班规则批量合并未生效草稿。", "工具栏。", "不与已生效班合并；合班后清空分组与教师。", "确认框。"),
                fn("5", "生效", "Activate Plan", "批量将计划生效。", "工具栏勾选生效。", "进入开课安排可见。", "无。"),
                fn("6", "删除", "Delete", "删除计划行与关联数据。", "工具栏确认。", "可再生成。", "无。"),
                fn("7", "预置人数", "View Preset Count", "查看可修预置人数。", "列点击。", "—", "学生预置人数抽屉。"),
                fn("8", "操作记录", "Change Log", "查看字段变更。", "列→查看。", "—", "修改记录窗。"),
            ],
            non_goals="不含通识/特殊开课计划；不含任务安排三步本体；不含操作流程图、校选课程管理。",
            pending="合班后计划总人数汇总正式口径；学院与教务合班权限矩阵。",
        ),
    },
    {
        "folder": "开课安排",
        "filename": "开课安排20260803V1.docx",
        "meta": {
            "cn": "开课安排",
            "en": "Major Offering Task Arrangement",
            "parent_cn": "专业开课",
            "parent_en": "Major Offering",
            "path": "开课管理 → 专业开课 → 开课安排",
            "version": "V1",
            "date": "2026 年 8 月 3 日",
            "purpose": "规范专业开课任务安排层：对计划已生效教学班分三步完成前期设置、合班/Support、师资与授课确认，再批量生效进入开课名单。",
            "background": "导航以样式二（三步引导）为准；合班≠分组；任务生效与计划生效为双层状态。",
        },
        "menu": dict(
            section_no="2.2.2.2 ",
            title="开课安排（英文名称：Major Offering Task Arrangement）",
            status="已确认",
            intro="对计划已生效的专业教学班，按三步完成：①开课前期设置 ②课程班（合班·Support）③师资安排（安排教师·授课确认·共同授课）；"
            "各步可生效/撤回/退回。生效门槛含分组人数上限、学时软硬校验与全部教师已确认；零学时默认不豁免教师。"
            "共同授课须至少一名相同任课教师；授课确认侧学时「Combined with」仅展示给共同授课组内相同教师。"
            "修改教学任务与开课计划共用规则；退回计划保留合班（须手动拆）；撤回禁已有排课。",
            list_fields="勾选、开课详情、开课状态、生效状态、授课确认、开课学期、Course Code/Name、Classification、开课单位、选课类型、Credits、起止周、Prog. Batch、计划总人数、Lecturer、Coordinator、共同授课状态/课号、操作记录；"
            "步骤2含Support；步骤3含安排教师/共同授课/授课确认管理；展开子表为学时投递明细（教师名可点进课时详情）。",
            search_fields="学年学期、课程号/名称、生效状态、共同授课状态、授课确认；查询/重置。",
            flow_rel="输入为计划已生效班；任务生效后进入名单与开课清单；撤回保留结构但重置学生小组分配；退回使计划变回退并清空安排。",
            flow_pre="开课计划已生效；特殊课程设置可提供教室/Support默认。",
            flow_out="任务已生效数据进入开课名单、开课清单、排课/选课。",
            biz_flow="进入【开课安排】→ 步骤1设置 → 步骤2合班/Support → 步骤3安排教师与授课确认 → 勾选生效；需改计划则退回；需改安排则撤回。",
            proto_link="prototype/index.html → page-course-major-offering-task-style2；page-offering-grouping",
            field_tables=[
                (
                    "任务安排——关键字段",
                    [
                        f("1", "任务生效状态", "Task Arrangement Status", "状态", "是", "—", "草稿/已生效/回退", "是", "草稿"),
                        f("2", "开课状态", "Arrangement Completeness", "指示", "—", "—", "分组/学时完整度", "否", ""),
                        f("3", "授课确认", "Teacher Confirmation", "状态", "—", "—", "已确认/未确认/无教师；生效门槛", "是", "未确认"),
                        f("4", "Prog. Batch", "Programme Batch", "只读", "是", "—", "合班多行", "否", ""),
                        f("5", "Support历史", "Support History", "只读", "—", "—", "特殊课程设置同步", "是", "N"),
                        f("6", "Support学院", "Support Units", "多选", "否", "—", "可接课号默认", "是", ""),
                        f("7", "共同授课状态", "Shared Teaching", "Y/N", "—", "—", "须先有任课教师；同组须有相同教师", "是", "N"),
                        f("8", "退回说明", "Revert Remark", "文本", "条件", "退回计划必填", "—", "否", ""),
                    ],
                ),
                (
                    "分组与学时投递——字段",
                    [
                        f("1", "小组名称", "Group Name", "文本", "是", "同班不重复", "—", "否", "Group 1"),
                        f("2", "小组人数上限", "Group Capacity", "数字", "是", ">0", "—", "否", "40"),
                        f("3", "学时类型", "Hour Type", "下拉", "是", "L/T/P/O", "—", "是", "L"),
                        f("4", "投递范围", "Delivery Scope", "下拉", "是", "section/group", "—", "是", "section"),
                        f("5", "授课教师", "Lecturer", "选择器", "是", "至少1名", "多教师须Coordinator；列表/明细姓名可点", "否", ""),
                        f("6", "Course Coordinator", "Course Coordinator", "选择器", "条件", "多教师必填", "—", "否", ""),
                        f("7", "周次", "Week Range", "文本", "是", "落在班起止周", "—", "否", "1-14"),
                        f("8", "同时多组授课", "Simultaneous Groups", "开关", "否", "—", "开启后学时只计一次", "否", ""),
                    ],
                ),
            ],
            functions=[
                fn("1", "设置", "Pre-setup", "开课前期信息（教室等）。", "步骤1行操作。", "任务未生效可改范围见锁定规则。", "修改教学任务抽屉。"),
                fn("2", "合班", "Merge", "合并/拆分专业批次。", "步骤2。", "规则同开课计划合班。", "合拆班抽屉。"),
                fn("3", "Support管理", "Manage Support", "维护Support学院。", "步骤2。", "历史只读。", "Support管理。"),
                fn("4", "安排教师", "Arrange Teachers", "分组工作台。", "步骤3。", "—", "分组工作台。"),
                fn("5", "授课确认管理", "Teaching Confirmation", "查看/代确认；学时列含 Combined Groups / Combined with（同教师）。", "步骤3。", "生效门槛。", "授课确认管理弹窗。"),
                fn("6", "共同授课", "Shared Teaching", "配置共同授课课号。", "步骤3。", "须先有任课教师；须至少一名相同教师。", "共同授课设置。"),
                fn("7", "生效", "Activate Task", "批量生效任务安排。", "工具栏。", "须安排完整且教师已确认。", "无。"),
                fn("8", "撤回", "Withdraw", "任务回到未生效。", "工具栏。", "结构保留；名单分组重置；已排课禁撤。", "无。"),
                fn("9", "退回", "Revert to Plan", "退回开课计划。", "工具栏；填说明。", "已生效任务不可退回；清合分班分组教师。", "退回说明。"),
                fn("10", "操作记录", "Change Log", "变更历史。", "列操作。", "—", "修改记录窗。"),
                fn("11", "导入", "Import", "批量导入。", "工具栏。", "原型占位。", "无。"),
            ],
            non_goals="不含计划生成；不含通识安排合页细节；样式一仅对照，主路径为样式二。",
            pending="学时软/硬校验边界；导入规则；排课后撤回退回拦截矩阵。",
        ),
    },
    {
        "folder": "开课名单",
        "filename": "开课名单20260803V1.docx",
        "meta": {
            "cn": "开课名单",
            "en": "Major Offering Roster",
            "parent_cn": "专业开课",
            "parent_en": "Major Offering",
            "path": "开课管理 → 专业开课 → 开课名单",
            "version": "V1",
            "date": "2026 年 8 月 3 日",
            "purpose": "规范任务安排已生效后按开课任务维护各小组学生名单，以及退回开课安排、与开放选课应用的边界。",
            "background": "名单维护在分组工作台学生名单 Tab 完成；支持预置分配、手工维护与一键分配。",
        },
        "menu": dict(
            section_no="2.2.2.3 ",
            title="开课名单（英文名称：Major Offering Roster）",
            status="已确认",
            intro="任务安排生效后，按开课任务展示列表；「管理名单」进入分组工作台学生名单 Tab。"
            "不开放选课：预置+手工+一键分配；开放选课：默认由选课应用生成。支持批量退回至开课安排（名单可保留）。",
            list_fields="勾选、开课学期、课程号、课程名称、选课类型、教师、教师数、小组数、Prog. Batch、学分、总学时、开课单位、起止周、人数、操作记录、操作（管理名单）。",
            search_fields="学年学期、上课学院、上课专业、开课单位、课程号；查询/重置。",
            flow_rel="读取任务已生效教学班；名单变更可审计；退回后任务回草稿且名单保留（不同于安排页撤回重置分组）。",
            flow_pre="计划已生效且任务安排已生效；允许空态进入页面。",
            flow_out="稳定名单供排课/考勤等；退回说明供安排侧查看。",
            biz_flow="进入【开课名单】→ 管理名单 → 预置/一键分配/添加/调组/移除 → 保存；必要时勾选退回开课安排。",
            proto_link="prototype/index.html → page-course-major-offering-roster；page-offering-grouping（学生名单 Tab）",
            field_tables=[
                (
                    "开课任务列表——字段",
                    [
                        f("1", "开课学期", "Offering Term", "代码", "是", "—", "—", "是", "2026/04"),
                        f("2", "课程号", "Course Code", "文本", "是", "—", "—", "否", "ACC301"),
                        f("3", "选课类型", "Enrollment Type", "枚举", "是", "—", "开放/不开放", "是", "不开放"),
                        f("4", "教师数/小组数", "—", "派生", "—", "—", "—", "否", ""),
                        f("5", "Prog. Batch", "Programme Batch", "只读", "是", "—", "合班多批次", "否", ""),
                        f("6", "人数", "Roster Headcount", "数字", "—", "—", "当前名单人数", "否", "40"),
                    ],
                ),
                (
                    "学生名单——字段",
                    [
                        f("1", "学号", "Student No", "文本", "是", "—", "—", "否", ""),
                        f("2", "姓名", "Name", "文本", "是", "—", "—", "否", ""),
                        f("3", "性别", "Gender", "代码", "否", "—", "—", "是", ""),
                        f("4", "国籍类型", "Nationality Type", "代码", "否", "—", "Local/Chinese/International等", "是", ""),
                        f("5", "专业", "Programme", "只读", "是", "—", "—", "是", ""),
                        f("6", "入学批次", "Intake", "只读", "是", "—", "—", "是", ""),
                        f("7", "课程班/小组", "Group", "选择器", "是", "须属本班", "可调组", "否", ""),
                        f("8", "来源", "Source", "只读", "—", "—", "预置/手工/选课应用/一键分配等", "是", "预置"),
                    ],
                ),
            ],
            functions=[
                fn("1", "管理名单", "Manage Roster", "进入工作台维护名单。", "行操作。", "主入口。", "分组工作台-学生名单。"),
                fn("2", "一键分配", "One-click Assign", "按规则批量入组。", "工作台工具栏。", "无平行组按专业+intake导入；有平行组按规则拆组。", "规则配置弹窗。"),
                fn("3", "预置名单", "Assign Preset", "从预置可修名单分配。", "工作台。", "与一键分配并存。", "无。"),
                fn("4", "添加学生", "Add Students", "手工加入。", "工作台。", "受组容量约束。", "添加学生选择器。"),
                fn("5", "调整分组", "Reassign Group", "学生调组。", "工作台。", "—", "无。"),
                fn("6", "移除", "Remove", "移出名单。", "工作台。", "保留移除记录。", "名单移除记录。"),
                fn("7", "保存名单", "Save Roster", "保存维护结果。", "工作台。", "可自动保存提示。", "无。"),
                fn("8", "退回", "Revert to Task Arrangement", "批量退回开课安排。", "列表工具栏；填说明。", "名单可保留。", "退回说明。"),
                fn("9", "操作记录", "Change Log", "查看变更。", "列操作。", "—", "修改记录窗。"),
            ],
            non_goals="不含计划生成与任务三步本体；不含通识/特殊开课名单；不含学籍异动主数据维护。",
            pending="开放选课是否允许开课侧手工改名单的最终边界；学籍异动生效学期；一键分配与预置冲突提示文案。",
        ),
    },
]


def build_one(spec: dict, *, version: str, date_ymd: str, allow_overwrite: bool = False) -> Path:
    meta = dict(spec["meta"])
    menu = dict(spec["menu"])
    pack = MENU_FIELD_PACKS[spec["folder"]]
    # 全量字段表覆盖 DOCS 内旧的「关键字段」摘要表
    menu["field_tables"] = pack["tables"]
    menu["field_shared_note"] = pack.get("shared_note", "")
    menu["field_hidden_note"] = pack.get("hidden", "")

    version = version if version.startswith("V") else f"V{version}"
    meta["version"] = version
    # 展示用中文日期
    y, m, d = date_ymd[:4], str(int(date_ymd[4:6])), str(int(date_ymd[6:8]))
    meta["date"] = f"{y} 年 {m} 月 {d} 日"

    filename = f"{spec['folder']}{date_ymd}{version}.docx"
    parent = MENU_PARENT[spec["folder"]]
    ver_dir = version_dir(spec["folder"], date_ymd, version, parent)
    out = ver_dir / filename
    ver_dir.mkdir(parents=True, exist_ok=True)
    if out.exists() and not allow_overwrite:
        raise SystemExit(
            f"拒绝覆盖已有版本：{out}\n"
            f"请使用 --version 升版（如 V2）并换用新 --date，或参阅「需求文档版本与变更说明规范」。"
        )

    doc = Document()
    set_document_landscape(doc)

    add_heading(doc, "厦大马来分校本科教务系统产品需求文档", 0)
    add_para(doc, f"模块：开课管理 · {meta['cn']}", bold=True)
    add_para(doc, f"文档版本：{meta['version']}　　创建日期：{meta['date']}")
    add_para(doc, f"菜单路径：{meta['path']}")
    add_para(
        doc,
        "依据：可交互原型（prototype/index.html、app.js）；"
        "模板结构对齐《厦大马来分校本科教务系统产品需求文档模板（空白模板）0610》。"
        "升版时不得覆盖旧版，并须产出变更说明："
        "正文若按 PRD 模板 V3，用变更说明简版 V3；仍为旧 V2 结构可用简版 V2。",
    )
    doc.add_paragraph()

    add_heading(doc, "文档概述", 1)
    add_heading(doc, "1.1 文档目的", 2)
    add_para(doc, meta["purpose"])
    add_heading(doc, "1.2 开发背景", 2)
    add_para(doc, meta["background"])
    add_para(doc, "开发模式：边分析边迭代，分模块生成可交互原型；覆盖范围：本科生。")
    add_heading(doc, "1.3 文档说明", 2)
    add_para(doc, "本文档为单个二级菜单需求说明；已确认需求标记为「已确认」。不含操作流程图。")
    add_para(doc, "字段信息表按「所在界面/区块」列全（筛选、列表、弹窗、抽屉分区、步骤/Tab 差异列）；隐藏字段单独说明。")
    add_para(doc, "术语：计划/任务「生效」对应原型 submitStatus / 任务安排提交态；UI 文案为草稿 / 已生效 / 回退。")

    add_heading(doc, "系统分析", 1)
    add_heading(doc, "2.1 应用目录——开课管理（节选）", 2)
    add_grid_table(
        doc,
        ["一级目录", "一级目录英文名称", "二级目录", "二级目录英文名称", "三级目录", "备注说明"],
        [
            [
                meta["parent_cn"],
                meta["parent_en"],
                meta["cn"],
                meta["en"],
                "—",
                meta["path"],
            ]
        ],
        col_widths=[3.2, 3.8, 3.2, 4.2, 2.0, 8.0],
    )

    add_heading(doc, f"2.2 开课管理——系统需求（{meta['cn']}）", 2)
    add_heading(doc, f"2.2.1 {meta['parent_cn']}（英文名称：{meta['parent_en']}）", 3)
    add_menu_block(doc, **menu)

    doc.save(out)
    return out


def _latest_version_num(menu_folder: Path, menu: str) -> int:
    import re

    best = 0
    # 版本文件夹：<菜单><日期>V<n>；兼容旧平铺 docx
    dir_pat = re.compile(rf"^{re.escape(menu)}\d{{8}}V(\d+)$")
    file_pat = re.compile(rf"^{re.escape(menu)}\d{{8}}V(\d+)\.docx$")
    if not menu_folder.is_dir():
        return 0
    for p in menu_folder.iterdir():
        if p.is_dir():
            m = dir_pat.match(p.name)
            if m:
                best = max(best, int(m.group(1)))
            continue
        m = file_pat.match(p.name)
        if m:
            best = max(best, int(m.group(1)))
    return best


def main():
    import argparse
    from datetime import date

    if str(SCRIPT_DIR) not in sys.path:
        sys.path.insert(0, str(SCRIPT_DIR))

    ap = argparse.ArgumentParser(
        description="生成二级菜单 PRD docx。默认不覆盖已有版本；升版请用 --bump 或 --version。"
    )
    ap.add_argument(
        "--version",
        help="显式版本号，如 V2。与 --bump 二选一；省略且未 --bump 时仅允许首次写入不存在的目标。",
    )
    ap.add_argument(
        "--bump",
        action="store_true",
        help="在各菜单已有最高版本基础上 +1（推荐升版方式）。",
    )
    ap.add_argument(
        "--date",
        default=date.today().strftime("%Y%m%d"),
        help="版本日期 YYYYMMDD（默认今天）",
    )
    ap.add_argument(
        "--menu",
        action="append",
        dest="menus",
        help="只生成指定菜单（可多次）；默认全部",
    )
    ap.add_argument(
        "--allow-overwrite",
        action="store_true",
        help="危险：允许覆盖同名文件（规范禁止；仅应急）",
    )
    args = ap.parse_args()

    if not re_match_ymd(args.date):
        raise SystemExit("--date 须为 YYYYMMDD")

    written = []
    specs = DOCS
    if args.menus:
        want = set(args.menus)
        specs = [s for s in DOCS if s["folder"] in want]
        missing = want - {s["folder"] for s in specs}
        if missing:
            raise SystemExit(f"未知菜单：{', '.join(sorted(missing))}")

    for spec in specs:
        folder = menu_dir(spec["folder"])
        if args.bump:
            n = _latest_version_num(folder, spec["folder"]) + 1
            if n < 1:
                n = 1
            ver = f"V{n}"
        elif args.version:
            ver = args.version if args.version.startswith("V") else f"V{args.version}"
        else:
            # 无升版参数：沿用 DOCS 默认文件名中的版本，但仍禁止覆盖
            ver = spec["meta"].get("version", "V1")
            if not ver.startswith("V"):
                ver = f"V{ver}"

        path = build_one(
            spec,
            version=ver,
            date_ymd=args.date,
            allow_overwrite=args.allow_overwrite,
        )
        written.append(path)
        print(f"Wrote: {path}")

    print(f"Done: {len(written)} files")
    if args.bump or (args.version and args.version not in ("V1", "1")):
        print(
            "提醒：升版后请运行 generate-submenu-prd-change-note.py 产出变更说明"
            "（默认 V3；旧 PRD 结构加 --template-version v2），并补全「本次改了什么」。"
        )


def re_match_ymd(s: str) -> bool:
    import re

    return bool(re.fullmatch(r"\d{8}", s))


if __name__ == "__main__":
    main()
