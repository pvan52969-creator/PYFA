#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「需求调整变更说明模板（简版）V3」docx + md 样例。

对齐 PRD 模板 V3 的界面/字段/功能/规则 ID；不覆盖 V2 模板文件。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "参考文档" / "0、模板"
STEM = "需求调整变更说明模板（简版）20260901V3"

CN_FONT = "宋体"
CN_HEADING = "黑体"
EN_FONT = "Times New Roman"
HEADER_FILL = "D9E2F3"
HINT_COLOR = RGBColor(0x80, 0x80, 0x80)
HEADING_COLOR = RGBColor(0x1F, 0x29, 0x37)
BORDER_COLOR = "000000"


def set_run_font(run, *, cn=CN_FONT, en=EN_FONT, size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if color is not None:
        run.font.color.rgb = color


def apply_line_spacing(p):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.line_spacing = 1.5


def add_para(doc, text="", *, bold=False, size=11, color=None, align=None, space_after=8):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    apply_line_spacing(p)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_hint(doc, text):
    return add_para(doc, text, size=10, color=HINT_COLOR, space_after=8)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 14, 3: 12}
    for run in h.runs:
        set_run_font(run, cn=CN_HEADING, size=sizes.get(level, 12), bold=True, color=HEADING_COLOR)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(8)
    apply_line_spacing(h)
    return h


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def set_table_cell_margins(tbl_pr, top=80, bottom=80, left=100, right=100):
    for old in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER_COLOR)
        borders.append(el)
    tbl_pr.append(borders)
    set_table_cell_margins(tbl_pr)


def set_col_widths(table, widths_cm):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def fill_cell(cell, text, *, bold=False, size=9, align=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    apply_line_spacing(p)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)


def add_kv_table(doc, rows, col_widths=(4.5, 21.1)):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_col_widths(table, col_widths)
    for i, (k, v) in enumerate(rows):
        fill_cell(table.rows[i].cells[0], k, bold=True, size=10.5)
        set_cell_shading(table.rows[i].cells[0], "F2F2F2")
        fill_cell(table.rows[i].cells[1], v, size=10.5)
    doc.add_paragraph()
    return table


def add_grid_table(doc, headers, rows, col_widths, *, header_fill=HEADER_FILL):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        fill_cell(cell, h, bold=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, header_fill)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            fill_cell(
                table.rows[ri + 1].cells[ci],
                val,
                size=8,
                align=WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 2, 3, 5, 7) else None,
            )
    set_col_widths(table, col_widths)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    return add_para(doc, text, size=10, bold=True, space_after=4)


def add_page_field(paragraph, kind="PAGE"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {kind} "
    run2._r.append(instr)
    run3 = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run3._r.append(end)
    for r in (run, run2, run3):
        set_run_font(r, size=9, color=HINT_COLOR)


def setup_section(doc, header_text):
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21.0)
        sec.left_margin = Cm(1.6)
        sec.right_margin = Cm(1.6)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(1.8)
        sec.header_distance = Cm(0.8)
        sec.footer_distance = Cm(0.6)

        header = sec.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(header_text)
        set_run_font(run, size=9, color=HINT_COLOR)

        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run("第 ")
        set_run_font(r1, size=9, color=HINT_COLOR)
        add_page_field(fp, "PAGE")
        r2 = fp.add_run(" 页 / 共 ")
        set_run_font(r2, size=9, color=HINT_COLOR)
        add_page_field(fp, "NUMPAGES")
        r3 = fp.add_run(" 页")
        set_run_font(r3, size=9, color=HINT_COLOR)


def add_title_block(doc, title, subtitle):
    p = add_para(doc, title, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    for run in p.runs:
        set_run_font(run, cn=CN_HEADING, size=18, bold=True, color=HEADING_COLOR)
    add_para(doc, subtitle, size=10, color=HINT_COLOR, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def build_docx() -> Path:
    doc = Document()
    setup_section(doc, "厦大马来分校本科教务系统——需求调整变更说明（简版）· V3")
    add_title_block(
        doc,
        "厦大马来分校本科教务系统——需求调整变更说明（简版模板）",
        "模板版本：V3　　创建日期：2026 年 9 月 1 日　　"
        "对照：需求调整变更说明模板（简版）20260821V2 + 厦大马来分校本科教务系统PRD模板_V3",
    )
    add_para(
        doc,
        "相对 V2：版本信息增加本版 PRD 模板 / 菜单路径 / 原型地址；"
        "总览表增加对象类型与对象ID，便于对齐 PRD V3 的 P/FD/F/BR/UC 编号。"
        "不含填写人/确认人/确认状态签核栏（整份确认记在 PRD 文档信息）。"
        "打开后请更新页码域：全选 → F9。",
        size=10,
        color=HINT_COLOR,
    )
    add_para(doc, "一眼对照：哪个模块 / 哪个对象ID / 改了什么", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)

    add_heading(doc, "1 版本信息", 1)
    add_hint(
        doc,
        "对照上一有效版本（中间版本已废弃则直接写 V2→V4）。"
        "版本文件夹命名：＜二级菜单名称＞＜YYYYMMDD＞＜Vn＞；变更说明放在新版本文件夹内。"
        "无实质需求变更不要升版。"
        "本版 PRD 模板填 V3（按 PRD 模板 V3 写正文）或 V2（正文仍为旧结构，过渡期可用变更说明 V2）。",
    )
    add_kv_table(
        doc,
        [
            ("需求文档名称", "厦大马来分校本科教务系统产品需求文档 — ＜二级菜单名称＞"),
            ("菜单路径", "＜如：开课管理 → 专业开课 → 开课安排＞"),
            ("上一有效版 → 本版", "V__ → V__"),
            ("上一版文件", "＜菜单＞＜旧日期＞＜旧Vn＞.docx"),
            ("本版文件", "＜菜单＞＜新日期＞＜新Vn＞.docx"),
            ("本版 PRD 模板", "V3（或过渡期写 V2）"),
            ("原型地址", "prototype/index.html → page-xxx（可选）"),
            ("变更日期", "YYYY-MM-DD"),
        ],
    )

    add_heading(doc, "2 本次改了什么（总览）", 1)
    add_para(
        doc,
        "按「模块 → 对象 → 变动」填写；一行一事。变更类型只填：新增 / 修改 / 删除 / 澄清。"
        "无变动的模块不必出现。有 PRD ID 的对象必须填对象ID；纯文案澄清可填 —。",
    )
    add_caption(doc, "表 2-1 变更总览")
    overview_headers = [
        "序号",
        "模块（菜单路径）",
        "对象类型",
        "对象ID",
        "功能 / 位置（可读名）",
        "变更类型",
        "调整内容（从什么变成什么）",
        "状态",
    ]
    # 横向 A4 宽约 26.5cm 可用
    overview_widths = [1.0, 4.2, 1.6, 2.0, 3.2, 1.6, 9.0, 2.0]
    overview_rows = [
        [
            "1",
            "例：开课管理 → 专业开课 → 开课安排",
            "字段",
            "FD012",
            "例：选课方式",
            "修改",
            "例：Edit：标签「选修/必修」→「选课方式」",
            "已明确",
        ],
        [
            "2",
            "例：开课管理 → 专业开课 → 开课安排",
            "功能",
            "F007",
            "例：停课",
            "删除",
            "例：工具栏去掉停课；改由课程班统一入口",
            "已明确",
        ],
    ] + [[str(i), "", "", "", "", "", "", "未确认"] for i in range(3, 9)]
    add_grid_table(doc, overview_headers, overview_rows, overview_widths)

    add_para(doc, "填写提示：", bold=True)
    add_para(doc, "· 模块：一级/二级/三级菜单路径，如「开课管理 → 专业开课 → 开课安排」。")
    add_para(
        doc,
        "· 对象类型：界面 / 字段 / 功能 / 规则 / 数据流 / 其他"
        "（对应 PRD V3：P01… / FD… / F… / BR… / 第 8 章 / 文档结构等）。",
    )
    add_para(
        doc,
        "· 对象ID：如 P01-M01、FD003、F012、BR001；未编号对象填 —。"
        "未确认且 PRD 已建 UC 时，在调整内容末尾标注 UC00x。",
    )
    add_para(doc, "· 功能/位置：可读名（按钮、字段、列表列、筛选项、步骤、弹窗/抽屉等）。")
    add_para(
        doc,
        "· 调整内容：一句话说清「从什么变成什么」；字段状态变更须写明模式，"
        "如「Edit：可编辑 → 只读」。新增写「新增xxx」，删除写「删除xxx」。",
    )
    add_para(
        doc,
        "· 状态：已明确 / 未确认。"
        "已明确=用户指定且已讲清；未确认=未聊过、未拍板、或由 AI 补写/可能遗漏。",
    )

    add_heading(doc, "3 前后对照（可选，复杂变更再填）", 1)
    add_para(doc, "总览表说不清时展开对照；简单修改可跳过。序号对应上表；须能落到对象ID。")
    add_caption(doc, "表 3-1 前后对照")
    add_grid_table(
        doc,
        ["序号（对应上表）", "对象ID", "变更前（上一有效版）", "变更后（本版）"],
        [
            ["1", "FD012", "例：Edit 可编辑；标签「选修/必修」", "例：Edit 可编辑；标签「选课方式」"],
            ["", "", "", ""],
            ["", "", "", ""],
            ["", "", "", ""],
        ],
        [2.4, 2.4, 10.0, 10.0],
    )

    add_heading(doc, "4 连带影响（可选）", 1)
    add_para(
        doc,
        "仅当本次调整会影响其他模块、原型或测试时填写。"
        "「影响到哪里」鼓励写下游菜单路径或界面 ID（对齐 PRD V3 第 8 章下游影响）。",
    )
    add_caption(doc, "表 4-1 连带影响")
    add_grid_table(
        doc,
        ["序号", "影响到哪里", "影响说明", "需同步"],
        [
            ["1", "例：课程班 / P01", "例：停课入口迁至课程班", "□ 原型　□ 开发　□ 测试"],
            ["", "", "", "□ 原型　□ 开发　□ 测试"],
            ["", "", "", "□ 原型　□ 开发　□ 测试"],
        ],
        [1.4, 6.0, 12.6, 5.2],
    )

    add_heading(doc, "5 过渡策略（模板说明，成文时可删）", 1)
    add_para(doc, "· 不删除、不覆盖「需求调整变更说明模板（简版）20260821V2」。")
    add_para(doc, "· 历史已产出的 V2 结构变更说明不必批量回改。")
    add_para(doc, "· 正文仍按旧 PRD V2 结构升版时，变更说明可继续用 V2。")
    add_para(doc, "· 自首份按《厦大马来分校本科教务系统PRD模板_V3》写正文的升版起，强制使用本 V3 变更说明。")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{STEM}.docx"
    if path.exists():
        raise SystemExit(f"拒绝覆盖已有模板：{path}")
    doc.save(path)
    return path


def build_md_sample() -> Path:
    path = OUT_DIR / f"{STEM}.md"
    if path.exists():
        raise SystemExit(f"拒绝覆盖已有样例：{path}")
    text = f"""# 需求调整变更说明模板（简版）· V3 样例

> 模板文件：`参考文档/0、模板/{STEM}.docx`  
> 对照：变更说明简版 20260821V2 + 厦大马来分校本科教务系统PRD模板_V3  
> 创建日期：2026-09-01

一眼对照：哪个模块 / 哪个对象ID / 改了什么。

## 1 版本信息

| 项目 | 内容 |
|------|------|
| 需求文档名称 | 厦大马来分校本科教务系统产品需求文档 — ＜二级菜单名称＞ |
| 菜单路径 | ＜如：开课管理 → 专业开课 → 开课安排＞ |
| 上一有效版 → 本版 | V__ → V__ |
| 上一版文件 | ＜菜单＞＜旧日期＞＜旧Vn＞.docx |
| 本版文件 | ＜菜单＞＜新日期＞＜新Vn＞.docx |
| 本版 PRD 模板 | V3（或过渡期写 V2） |
| 原型地址 | prototype/index.html → page-xxx（可选） |
| 变更日期 | YYYY-MM-DD |

不含填写人 / 确认人 / 确认状态签核栏；整份确认记在 PRD 文档信息。

## 2 本次改了什么（总览）

按「模块 → 对象 → 变动」填写；一行一事。变更类型只填：新增 / 修改 / 删除 / 澄清。  
有 PRD ID 的对象必须填对象ID；纯文案澄清可填 —。

| 序号 | 模块（菜单路径） | 对象类型 | 对象ID | 功能 / 位置（可读名） | 变更类型 | 调整内容（从什么变成什么） | 状态 |
|------|------------------|----------|--------|---------------------|----------|------------------------------|------|
| 1 | 例：开课管理 → 专业开课 → 开课安排 | 字段 | FD012 | 例：选课方式 | 修改 | 例：Edit：标签「选修/必修」→「选课方式」 | 已明确 |
| 2 | 例：开课管理 → 专业开课 → 开课安排 | 功能 | F007 | 例：停课 | 删除 | 例：工具栏去掉停课；改由课程班统一入口 | 已明确 |

填写提示：

- 对象类型：界面 / 字段 / 功能 / 规则 / 数据流 / 其他
- 对象ID：如 P01-M01、FD003、F012、BR001；未编号填 —；未确认可在调整内容末标注 UC00x
- 字段状态变更须写明模式，如「Edit：可编辑 → 只读」
- 状态：已明确 / 未确认

## 3 前后对照（可选，复杂变更再填）

| 序号（对应上表） | 对象ID | 变更前（上一有效版） | 变更后（本版） |
|------------------|--------|----------------------|----------------|
| 1 | FD012 | 例：Edit 可编辑；标签「选修/必修」 | 例：Edit 可编辑；标签「选课方式」 |

## 4 连带影响（可选）

| 序号 | 影响到哪里 | 影响说明 | 需同步 |
|------|------------|----------|--------|
| 1 | 例：课程班 / P01 | 例：停课入口迁至课程班 | □原型 □开发 □测试 |

## 5 过渡策略

- 不删除、不覆盖 V2 变更说明模板。
- 历史 V2 变更说明不必批量回改。
- 正文仍按旧 PRD V2 结构升版 → 可用变更说明 V2。
- 自首份按 PRD 模板 V3 写正文的升版起 → 强制使用本 V3 变更说明。
"""
    path.write_text(text, encoding="utf-8")
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    docx = build_docx()
    md = build_md_sample()
    print(docx.relative_to(ROOT), docx.stat().st_size)
    print(md.relative_to(ROOT), md.stat().st_size)


if __name__ == "__main__":
    main()
