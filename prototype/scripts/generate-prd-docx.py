#!/usr/bin/env python3
"""Generate 培养方案管理 PRD docx with bordered tables matching XMUM template."""

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path

from prd_build import build_document

OUT = Path(__file__).resolve().parent.parent / "参考文档" / "培养方案管理系统产品需求文档.docx"
SCRIPT_DIR = Path(__file__).resolve().parent

# A4 横向：可排版宽度约 25.7 cm（页宽 29.7 − 左右边距）
LANDSCAPE_SCALE = 1.55
FIELD_TABLE_COL_WIDTHS = [1.2, 3.1, 3.4, 1.9, 1.1, 3.9, 4.4, 1.9, 2.5]


def scale_widths(widths):
    if isinstance(widths, tuple):
        return tuple(round(w * LANDSCAPE_SCALE, 1) for w in widths)
    return [round(w * LANDSCAPE_SCALE, 1) for w in widths]


def set_document_landscape(doc):
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

FIELD_HEADERS = [
    "序号", "字段中文名称", "英文名称", "字段类型", "必填",
    "校验规则\n（没有明确规则时先限制字符即可，放宽）",
    "备注（如下拉框取值来源、说明等）", "是否代码集取值", "数据格式样例\n（无特定格式留空）",
]


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return h


def add_kv_table(doc, rows, col_widths=(5, 18.5)):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_col_widths(table, col_widths)
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = k
        c1.text = v
        for cell in (c0, c1):
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        c0.paragraphs[0].runs[0].bold = True
    doc.add_paragraph()
    return table


def add_grid_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if col_widths:
        set_col_widths(table, col_widths)
    doc.add_paragraph()
    return table


def add_field_table(doc, title, rows):
    add_para(doc, title, bold=True)
    add_grid_table(doc, FIELD_HEADERS, rows, col_widths=FIELD_TABLE_COL_WIDTHS)


def add_function_item(doc, num, cn_name, en_name, desc, interaction, notes, drill_down, source=""):
    """参考模板格式：a 节纯中文业务说明；b 节技术标识后置为「中文术语（编码）」。"""
    title = f"{num}、功能按钮——{cn_name}（英文名称：{en_name}）"
    if source:
        title += f"【来源：{source}】"
    add_para(doc, title, bold=True)
    add_para(doc, "a. 功能说明（描述、业务的事件交互、备注信息等）：")
    add_para(doc, f"· 描述：{desc}")
    add_para(doc, f"· 业务的事件交互：{interaction}")
    add_para(doc, f"· 备注信息（校验规则补充、其他说明等）：{notes}")
    add_para(doc, f"b. 下钻页面说明（备注说明）：{drill_down}")
    doc.add_paragraph()


def add_menu_block(doc, section_no, title, status, intro, list_fields, search_fields,
                   flow_rel, flow_pre, flow_out, biz_flow, proto_link="",
                   field_tables=None, functions=None):
    add_heading(doc, f"{section_no}{title}（需求确认状态：{status}）", level=4)
    add_para(doc, "菜单介绍", bold=True)
    add_kv_table(doc, [("菜单介绍", ""), ("菜单内容简介", intro)])
    add_para(doc, "页面展示字段信息", bold=True)
    add_kv_table(doc, [("页面展示字段信息", list_fields)])
    add_para(doc, "支持查询检索的字段信息", bold=True)
    add_kv_table(doc, [("支持查询检索的字段信息", search_fields)])
    add_para(doc, "数据前后流转关系（说明、前置条件、下游输出等）", bold=True)
    add_kv_table(doc, [
        ("1）关系说明", flow_rel),
        ("2）前置条件", flow_pre),
        ("3）下游输出", flow_out),
    ])
    add_para(doc, "5、业务流关系（操作流程）", bold=True)
    add_kv_table(doc, [("5、业务流关系（操作流程）", biz_flow)])
    if proto_link:
        add_para(doc, "6、原型参考链接", bold=True)
        add_kv_table(doc, [("原型参考链接", proto_link)])
    if field_tables:
        add_para(doc, "新增—字段信息表", bold=True)
        for ft_title, ft_rows in field_tables:
            add_field_table(doc, ft_title, ft_rows)
    if functions:
        add_para(doc, "（3）菜单功能清单", bold=True)
        for fn in functions:
            add_function_item(doc, *fn)


# 校验提示附录（业务文案纯中文）
VALIDATION_PROMPTS = [
    ["标签页一·分类", "未完成分类即进入标签页二", "轻提示", "请先在标签页一完成课程分类建设（各级分类及学分要求），再进入课程设置。"],
    ["标签页一·分类", "未选择一级或二级分类", "轻提示", "请选择一级分类 / 请选择二级分类。"],
    ["标签页一·分类", "同级分类重名", "轻提示", "该分类下已存在同名分类。"],
    ["标签页一·分类", "学分未填或非法", "轻提示", "请填写最低/最高学分；学分须为非负数字；最低不能大于最高。"],
    ["标签页一·分类", "三级最低学分突破二级上限", "轻提示", "三级分类最低学分合计将超过所属二级分类学分上限。"],
    ["标签页一·分类", "必修二三级学分结构不一致", "轻提示", "必修：三级最低学分合计须等于二级最低学分且等于二级最高学分。"],
    ["标签页一·分类", "选修三级合计超过二级最低", "轻提示", "选修：三级最低学分合计不能超过二级最低学分。"],
    ["标签页一·分类", "选修矩阵为空", "轻提示", "修读类型为选修时，请至少添加一条学期修读要求。"],
    ["标签页一·分类", "矩阵学期未选择", "轻提示", "请为每条学期要求选择对应学期。"],
    ["标签页一·分类", "矩阵单学期最高低于最低", "轻提示", "最高学分限制须大于等于最低学分要求。"],
    ["标签页一·分类", "矩阵各学期学分之和超限", "轻提示", "各学期最低/最高学分要求之和不能超过二级分类最高学分限制。"],
    ["标签页二·课程", "无分类时新增课程", "轻提示", "请先在标签页一建设课程分类，否则无法新增课程。"],
    ["标签页二·课程", "未选课程库课程", "轻提示", "请先从课程库选择一门课程。"],
    ["标签页二·课程", "分类未选全", "轻提示", "请选择一级/二级分类；有三级子分类时须选择三级分类。"],
    ["标签页二·课程", "必修未指定开课学期", "轻提示", "必修课程须指定开课学期。"],
    ["标签页二·课程", "学分无效", "轻提示", "请填写有效的课程学分。"],
    ["标签页二·课程", "必修添加后超二级最高学分", "轻提示", "课程学分之和超过该二级分类的最高学分限制。"],
    ["标签页二·课程", "批量未选课程", "轻提示", "请至少选择一门课程。"],
    ["标签页二·课程", "批量课程均已存在", "轻提示", "所选课程均已存在于本方案中。"],
    ["标签页二·课程", "进入第三步无学习成果", "轻提示", "请先在第二步创建至少一条课程学习成果。"],
    ["标签页二·课程", "移除课程", "确认弹窗", "确定从本方案中移除该课程吗？此操作不可撤销。"],
    ["提交审批", "已配学分低于最低要求", "轻提示", "学分校验未通过：某分类课程学分之和小于最低学分要求。"],
    ["提交审批", "已配学分超过最高限制", "轻提示", "学分校验未通过：某必修分类已配置学分超过最高限制。"],
    ["提交审批", "选修矩阵或结构校验失败", "轻提示", "学分校验未通过：选修矩阵或二三级学分结构不符合规则。"],
    ["列表·版本", "新增版本开始批次过早", "轻提示", "新版本开始入学批次须晚于上一版本截止批次。"],
    ["列表·版本", "批量提交含不可提交项", "轻提示", "仅草稿或已驳回状态的版本可提交审批。"],
    ["执行计划", "已开课不可提交或撤回", "轻提示/确认弹窗", "已开课的执行计划无法提交或撤回。"],
    ["变更申请", "目标版本未通过审批", "轻提示", "只能选择已审批通过的培养方案版本。"],
    ["变更申请", "已有进行中变更", "轻提示", "该版本已有进行中的变更申请，请完成或撤回后再新建。"],
]


def build():
    import sys
    if str(SCRIPT_DIR) not in sys.path:
        sys.path.insert(0, str(SCRIPT_DIR))
    build_document(
        add_heading, add_para, add_kv_table, add_grid_table, add_field_table,
        add_menu_block, scale_widths, VALIDATION_PROMPTS, OUT, Document,
        set_document_landscape,
    )


if __name__ == "__main__":
    build()

