#!/usr/bin/env python3
"""按 V2 模板为用户点名的 9 个侧栏菜单生成新版 PRD + 变更说明（不覆盖旧版）。"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(SCRIPT_DIR))

from docx import Document


def load_mod(name: str, filename: str):
    path = SCRIPT_DIR / filename
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


tpl = load_mod("generate_prd_templates_v2", "generate-prd-templates-v2.py")
FIELD_HEADERS = tpl.FIELD_HEADERS
FIELD_WIDTHS = tpl.FIELD_WIDTHS
add_caption = tpl.add_caption
add_function_item = tpl.add_function_item
add_grid_table = tpl.add_grid_table
add_heading = tpl.add_heading
add_hint = tpl.add_hint
add_kv_table = tpl.add_kv_table
add_para = tpl.add_para
add_title_block = tpl.add_title_block
configure_styles = tpl.configure_styles
setup_section = tpl.setup_section

old_gen = load_mod("submenu_prd_old", "generate-submenu-prd-docx.py")
from prd_submenu_fields import MENU_FIELD_PACKS  # noqa: E402

DATE = "20260821"
DATE_DISP = "2026 年 8 月 21 日"
CHANGE_DATE = "2026-08-21"
BASE = ROOT / "参考文档" / "2、开课管理"

DOCS_BY_CN = {d["meta"]["cn"]: d for d in old_gen.DOCS}

MENUS = [
    {
        "cn": "开课时间设置",
        "en": "Offering Time Setting",
        "parent": "开课设置",
        "parent_en": "Offering Settings",
        "path": "开课管理 → 开课设置 → 开课时间设置",
        "page": "page-course-time-setting",
        "prev": ("20260804", "V2"),
        "ver": "V3",
        "kind": "bump",
    },
    {
        "cn": "校选课程管理",
        "en": "School Elective Course Management",
        "parent": "开课设置",
        "parent_en": "Offering Settings",
        "path": "开课管理 → 开课设置 → 校选课程管理",
        "page": "page-school-elective-courses",
        "prev": ("20260804", "V1"),
        "ver": "V2",
        "kind": "bump",
    },
    {
        "cn": "特殊课程设置",
        "en": "Special Course Settings",
        "parent": "开课设置",
        "parent_en": "Offering Settings",
        "path": "开课管理 → 开课设置 → 特殊课程设置",
        "page": "page-special-course-settings",
        "prev": ("20260804", "V2"),
        "ver": "V3",
        "kind": "bump",
    },
    {
        "cn": "开课计划",
        "en": "Major Offering Plan",
        "parent": "专业开课",
        "parent_en": "Major Offering",
        "path": "开课管理 → 专业开课 → 开课计划",
        "page": "page-course-offering-major",
        "prev": ("20260804", "V4"),
        "ver": "V5",
        "kind": "bump",
    },
    {
        "cn": "开课安排",
        "en": "Major Offering Task Arrangement",
        "parent": "专业开课",
        "parent_en": "Major Offering",
        "path": "开课管理 → 专业开课 → 开课安排",
        "page": "page-course-major-offering-task-style2",
        "prev": ("20260804", "V4"),
        "ver": "V5",
        "kind": "bump",
    },
    {
        "cn": "开课名单",
        "en": "Major Offering Roster",
        "parent": "专业开课",
        "parent_en": "Major Offering",
        "path": "开课管理 → 专业开课 → 开课名单",
        "page": "page-course-major-offering-roster",
        "prev": ("20260804", "V2"),
        "ver": "V3",
        "kind": "bump",
    },
    {
        "cn": "Teaching Load",
        "en": "Teaching Load (hour)",
        "parent": "开课清单",
        "parent_en": "Offering Manifest",
        "path": "开课管理 → 开课清单 → Teaching Load",
        "page": "page-course-teacher-teaching-load",
        "prev": None,
        "ver": "V1",
        "kind": "new",
    },
    {
        "cn": "授课确认管理",
        "en": "Teacher Course Confirmation (Admin)",
        "parent": "开课清单",
        "parent_en": "Offering Manifest",
        "path": "开课管理 → 开课清单 → 授课确认管理",
        "page": "page-course-teacher-confirmation-admin",
        "prev": None,
        "ver": "V1",
        "kind": "new",
    },
    {
        "cn": "授课确认（教师端）",
        "en": "Teacher Course Confirmation (Portal)",
        "parent": "开课清单",
        "parent_en": "Offering Manifest",
        "path": "开课管理 → 开课清单 → 授课确认（教师端）",
        "page": "page-teacher-course-confirmation",
        "prev": None,
        "ver": "V1",
        "kind": "new",
    },
]


def folder_name(m):
    return f"{m['cn']}{DATE}{m['ver']}"


def out_dir(m):
    return BASE / m["parent"] / m["cn"] / folder_name(m)


def patch_fields(cn: str, tables):
    out = []
    for title, rows in tables:
        new_rows = []
        for row in rows:
            cells = [str(c) for c in row]
            if cn == "开课安排":
                cells = [c.replace("Support 学院", "Support 专业").replace("Support学院", "Support 专业") for c in cells]
            if cn == "校选课程管理":
                cells = [c.replace("通识选修开课", "选修开课") for c in cells]
            new_rows.append(cells)
        out.append((title, new_rows))
    return out


NEW_SPECS = {
    "Teaching Load": {
        "purpose": "按教师汇总已安排授课学时（hour），支持按周次 / 学期 / 学年切换查看。",
        "background": "侧栏「开课清单」分组下的学时统计页；页面标题为「教师开课学时统计」。",
        "intro": "页面标题：教师开课学时统计。副标题：Teaching Load (hour) · 按教师汇总各学期已安排授课学时。"
        "统计维度三个 Tab：按周次、按学期（默认）、按学年。",
        "list_fields": "表头/表体由脚本按当前维度动态生成；底部为学期合计与人均均值。按学年时可出现「本学期计入」开关。",
        "search_fields": "开课学期（多选，提示最多 6 个学期）、所属部门、教师类型、课程类型（多选）、教师工号、教师姓名；查询 / 重置。",
        "flow_rel": "统计对象为已安排授课学时。工具栏可打开「超限阈值」、导出。可从本页跳转到授课确认管理（原型函数存在）。",
        "flow_pre": "依赖开课学期与教师安排数据（未确认与开课清单「已提交」是否强制联动）。",
        "flow_out": "学时统计表、导出文件（格式未确认）。",
        "biz_flow": "进入 Teaching Load → 选维度 → 设筛选 → 查询 → 按需设超限阈值 / 导出。",
        "roles": [("未指定角色（未确认）", "页面可见查询、阈值、导出；权限边界未确认", "未确认")],
        "non_goals": "不在本页维护开课安排或发送授课确认。",
        "pending": "· 与开课清单已提交安排是否强制联动（未确认）。\n· 导出文件格式（未确认）。\n· 使用角色与权限（未确认）。",
        "functions": [
            ("1", "查询 / 重置", "Query / Reset", "按筛选刷新统计表。", "筛选区按钮。", "学期多选最多 6 个（界面提示）。", "无下钻页面。"),
            ("2", "统计维度 Tab", "Dimension", "在按周次 / 按学期 / 按学年之间切换。", "顶部 Tab。", "按学年显示「本学期计入」开关。", "无下钻页面。"),
            ("3", "超限阈值", "Threshold", "设置各维度超限标红阈值。", "打开弹窗。", "阈值规则细节未确认。", "弹窗：超限阈值。"),
            ("4", "导出", "Export", "导出当前统计。", "工具栏按钮。", "导出格式未确认。", "无下钻页面。"),
        ],
        "fields": [
            (
                "字段信息表——筛选与工具栏",
                [
                    ["1", "页面·维度 Tab", "按周次", "Week", "Tab", "否", "—", "data-dimension=week", "否", ""],
                    ["2", "页面·维度 Tab", "按学期", "Term", "Tab", "否", "—", "默认选中", "否", ""],
                    ["3", "页面·维度 Tab", "按学年", "Academic Year", "Tab", "否", "—", "—", "否", ""],
                    ["4", "页面·筛选区", "开课学期", "Term", "多选", "否", "最多 6 个学期（提示）", "默认当前开课学期", "是", ""],
                    ["5", "页面·筛选区", "所属部门", "Department", "下拉", "否", "—", "全部所属部门", "是", ""],
                    ["6", "页面·筛选区", "教师类型", "Teacher Type", "下拉", "否", "—", "全部教师类型", "是", ""],
                    ["7", "页面·筛选区", "课程类型", "Course Type", "多选", "否", "—", "默认全部", "是", ""],
                    ["8", "页面·筛选区", "教师工号", "Staff ID", "输入框", "否", "—", "—", "否", ""],
                    ["9", "页面·筛选区", "教师姓名", "Name", "输入框", "否", "—", "—", "否", ""],
                    ["10", "页面·工具栏", "本学期计入", "Include current term", "开关", "条件", "按学年时出现", "是否计入本学期", "是", ""],
                    ["11", "页面·工具栏", "超限阈值", "Threshold", "按钮+展示", "否", "—", "标红阈值", "否", ""],
                    ["12", "页面·工具栏", "导出", "Export", "按钮", "否", "—", "—", "否", ""],
                    ["13", "页面·表", "合计 / 人均均值", "Totals / Average", "表尾", "—", "—", "底部行", "否", ""],
                ],
            )
        ],
    },
    "授课确认管理": {
        "purpose": "按教师/课程行管理授课确认：发送、代确认、查看历史。",
        "background": "从专业开课分组迁至侧栏「开课清单」。按教师维度下发系统内授课确认，支持管理端代确认。",
        "intro": "列表一行对应教师+开课任务。状态含草稿/待确认/已确认/不同意等（以原型状态胶囊为准）。TBC 教师无需确认。",
        "list_fields": "勾选｜Status｜Lecturer｜Course Code｜Course Name｜Classification｜Credits｜Teaching Weeks｜Teaching Method｜No. of Group｜No. of Combined Groups｜Total Student No.｜Weekly Teaching Hours｜Total Teaching Hours｜Course Coordinator｜Email｜Co-teaching Staff｜Co-teaching Email｜所属部门｜History｜Actions（发送 · 代确认）",
        "search_fields": "开课学期、所属部门、教师姓名、教师号、课程号；搜索 / 重置。",
        "flow_rel": "发送后教师端可见待确认。代确认仅待确认状态。历史复制未点确认的课程不可发送。",
        "flow_pre": "开课安排已有真实教师（非 TBC）。",
        "flow_out": "教师端待确认任务；授课确认进度可供开课安排列表展示。",
        "biz_flow": "选学期与部门 → 搜索 → 勾选 → 发送授课确认或批量代确认；行内可发送/代确认/查看 History。",
        "roles": [("教务/管理员（未确认正式角色名）", "发送、代确认、查看历史（对照原型按钮）", "未确认")],
        "non_goals": "教师本人确认在「授课确认（教师端）」；不在本页改课表。",
        "pending": "· 正式角色与学院管理员范围（未确认）。\n· 状态全集与流转图（未确认，仅按原型按钮可用性）。",
        "functions": [
            ("1", "搜索 / 重置", "Search / Reset", "按条件过滤列表。", "筛选区。", "—", "无下钻页面。"),
            ("2", "发送 / 发送授课确认", "Issue", "草稿或不同意可发送；批量按钮「发送授课确认」。", "行内或批量。", "历史复制未确认不可发。TBC 不发。", "二次确认弹窗。"),
            ("3", "代确认 / 批量代确认", "Proxy confirm", "仅待确认可代确认。", "行内或批量。", "记录操作人（原型提示）。", "二次确认弹窗。"),
            ("4", "History 查看", "History", "查看该行确认历史。", "History 列。", "—", "弹窗：确认历史。"),
        ],
        "fields": [
            (
                "字段信息表——筛选",
                [
                    ["1", "页面·筛选区", "开课学期", "Term", "下拉", "否", "—", "—", "是", ""],
                    ["2", "页面·筛选区", "所属部门", "Department", "下拉", "否", "含全部", "—", "是", ""],
                    ["3", "页面·筛选区", "教师", "Name", "搜索框", "否", "—", "姓名", "否", ""],
                    ["4", "页面·筛选区", "教师号", "Staff ID", "搜索框", "否", "—", "工号", "否", ""],
                    ["5", "页面·筛选区", "课程号", "Course Code", "搜索框", "否", "—", "—", "否", ""],
                ],
            ),
            (
                "字段信息表——列表（表头以原型英文为准）",
                [
                    ["1", "页面·列表", "勾选", "—", "复选框", "否", "—", "批量发送/代确认", "否", ""],
                    ["2", "页面·列表", "Status", "Status", "状态", "—", "—", "决定发送/代确认是否可点", "是", ""],
                    ["3", "页面·列表", "Lecturer", "Lecturer", "文本", "—", "—", "教师姓名", "否", ""],
                    ["4", "页面·列表", "Course Code", "Course Code", "文本", "—", "—", "可排序", "否", ""],
                    ["5", "页面·列表", "Course Name", "Course Name", "文本", "—", "—", "—", "否", ""],
                    ["6", "页面·列表", "Classification", "Classification", "文本", "—", "—", "—", "是", ""],
                    ["7", "页面·列表", "Credits", "Credits", "数字", "—", "—", "—", "否", ""],
                    ["8", "页面·列表", "Teaching Weeks", "Teaching Weeks", "文本", "—", "—", "—", "否", ""],
                    ["9", "页面·列表", "Teaching Method", "Teaching Method", "文本", "—", "—", "—", "否", ""],
                    ["10", "页面·列表", "No. of Group", "Group Count", "数字", "—", "—", "表头有规则提示", "否", ""],
                    ["11", "页面·列表", "No. of Combined Groups", "Combined Groups", "数字", "—", "—", "表头有规则提示", "否", ""],
                    ["12", "页面·列表", "Total Student No.", "Total Student No.", "数字", "—", "—", "—", "否", ""],
                    ["13", "页面·列表", "Weekly Teaching Hours", "Weekly Hours", "数字", "—", "—", "—", "否", ""],
                    ["14", "页面·列表", "Total Teaching Hours", "Total Hours", "数字", "—", "—", "—", "否", ""],
                    ["15", "页面·列表", "Course Coordinator", "Coordinator", "文本", "—", "—", "—", "否", ""],
                    ["16", "页面·列表", "Email Address", "Coordinator Email", "文本", "—", "—", "—", "否", ""],
                    ["17", "页面·列表", "Co-teaching Staff", "Co-teaching Staff", "文本", "—", "—", "可叠行", "否", ""],
                    ["18", "页面·列表", "Co-teaching Staff Email", "Co-teaching Email", "文本", "—", "—", "可叠行", "否", ""],
                    ["19", "页面·列表", "所属部门", "Department", "代码", "—", "—", "—", "是", ""],
                    ["20", "页面·列表", "History", "History", "链接", "—", "—", "查看", "否", ""],
                    ["21", "页面·列表", "Actions", "Actions", "链接", "—", "—", "发送 · 代确认", "否", ""],
                ],
            ),
        ],
    },
    "授课确认（教师端）": {
        "purpose": "教师在系统内确认本学期授课安排，或提交不同意及诉求。",
        "background": "教师端入口；页头有「演示身份」下拉（仅原型）。",
        "intro": "筛选开课学期。Tab：Pending / History。Pending 可确认；安排已变更须等教务重新下发。表格形态为确认信样式（含问候与 Notes）。",
        "list_fields": "确认信表格列与管理端信函表同源（课程、学时、小组等，以原型渲染为准）；Pending 显示确认操作。",
        "search_fields": "开课学期；演示身份（原型）。",
        "flow_rel": "任务来自管理端发送。不同意后管理端可再发送。授课确认截止日期来自开课时间设置（Notes）。",
        "flow_pre": "管理端已发送；教师为真实教师非 TBC。",
        "flow_out": "确认/不同意结果回写管理端状态。",
        "biz_flow": "选择学期 → Pending 查看安排 → 确认或不同意；History 查看已处理记录。",
        "roles": [("教师", "确认或提交不同意（对照原型文案）", "未确认正式权限模型")],
        "non_goals": "不在本页改开课安排。演示身份仅原型。",
        "pending": "· 不同意后的审批流（未确认）。\n· 与授课确认截止日期的拦截关系（未确认，Notes 有引用）。",
        "functions": [
            ("1", "开课学期", "Term", "切换本学期确认任务。", "下拉。", "—", "无下钻页面。"),
            ("2", "Pending / History", "Tabs", "待确认与历史。", "Tab。", "Pending 有数量角标。", "无下钻页面。"),
            ("3", "确认 / 不同意", "Confirm / Disagree", "仅 Pending 且快照仍有效时可操作。", "确认信表格操作。", "安排已变更则提示等待重新下发。", "未确认是否另有弹窗。"),
        ],
        "fields": [
            (
                "字段信息表",
                [
                    ["1", "页面·顶栏", "演示身份", "Demo teacher", "下拉", "否", "仅原型", "切换教师", "否", ""],
                    ["2", "页面·筛选", "开课学期", "Term", "下拉", "否", "—", "—", "是", ""],
                    ["3", "页面·Tab", "Pending", "Pending", "Tab", "否", "—", "待确认", "否", ""],
                    ["4", "页面·Tab", "History", "History", "Tab", "否", "—", "历史", "否", ""],
                ],
            )
        ],
    },
}


def extra_pending(cn: str) -> str:
    extra = {
        "开课安排": "· Support 管理弹窗已按用户指定改为「Support 专业」（已明确）。安排教师筛选仍有「Support 学院」文案（未确认是否同步改）。",
        "特殊课程设置": "· 本页列表/抽屉仍显示「默认 Support 学院」；安排侧 Support 管理已改为专业。两边是否统一（未确认）。",
        "校选课程管理": "· 下游侧栏现为「选修开课」；旧版写「通识选修开课」。本版正文改为选修开课（已明确对照侧栏）。",
        "开课时间设置": "· 教学周 5/14 仍为原型推导，正式以校历决议为准（未确认）。",
        "开课计划": "· 相对 V4 业务规则未在本轮对话中重开，正文沿用上一有效版并改套 V2 模板（规则口径：上一版已写明且仍对照原型的视为已明确；未再讨论的权限边界仍未确认）。",
        "开课名单": "· 开放选课名单与选课应用边界仍以产品最终口径为准（上一版已标注，未确认）。",
    }
    return extra.get(cn, "")


def bump_changes(m):
    rows = [
        {
            "where": "文档模板",
            "type": "修改",
            "content": "改为 V2 模板：文档信息表、条款（已明确）/（未确认）、1.5 倍行距、三级标题、字段表框线与跨页表头",
            "status": "已明确",
        },
        {
            "where": "文档信息·确认状态",
            "type": "澄清",
            "content": "整份文档签核仍用待确认/已确认；条款不用「已确认/待确认」",
            "status": "已明确",
        },
    ]
    if m["cn"] == "开课安排":
        rows.append({
            "where": "弹窗·Support 管理",
            "type": "修改",
            "content": "「当前/选择 Support 学院」改为「Support 专业」（用户指定，对照现行原型）",
            "status": "已明确",
        })
    if m["cn"] == "校选课程管理":
        rows.append({
            "where": "下游表述",
            "type": "修改",
            "content": "下游由「通识选修开课」改为侧栏现行「选修开课」",
            "status": "已明确",
        })
    if m["cn"] == "特殊课程设置":
        rows.append({
            "where": "Support 用词",
            "type": "澄清",
            "content": "本页仍为 Support 学院；安排侧已改 Support 专业。是否统一未确认",
            "status": "未确认",
        })
    return rows


def new_changes(m):
    return [
        {
            "where": "文档",
            "type": "新增",
            "content": f"首版需求文档，依据现行原型 `{m['page']}`；无上一版可对照",
            "status": "未确认",
        },
        {
            "where": "角色与权限",
            "type": "澄清",
            "content": "未在对话中指定正式角色，标（未确认）",
            "status": "未确认",
        },
        {
            "where": "字段/规则",
            "type": "澄清",
            "content": "仅收录原型可见控件与文案；未聊过的流转细节不编造",
            "status": "已明确",
        },
    ]


def write_md(path: Path, m: dict, spec: dict, field_tables, functions, hidden=""):
    prev = m["prev"]
    lines = [
        f"# 厦大马来分校本科教务系统产品需求文档 — {m['cn']}",
        "",
        f"> **文档版本：** {m['ver']} · **创建日期：** {DATE_DISP}  ",
        f"> **模板：** 产品需求文档模板（空白/带参考数据版）20260821V2  ",
        f"> **原型页面：** `{m['page']}`  ",
        "> **条款口径：** 用户指定且已讲清 →（已明确）；未聊过、AI 补写或可能遗漏 →（未确认）",
        "",
        "| 项 | 内容 |",
        "|---|---|",
        f"| 菜单路径 | {m['path']} |",
        f"| 英文名称 | {m['en']} |",
        f"| 文档版本 | {m['ver']} |",
        f"| 创建日期 | {CHANGE_DATE} |",
        f"| 文档签核 | □ 待确认　□ 已确认 |",
        f"| 原型页面 | `{m['page']}` |",
        "",
        "## 1. 文档概述",
        "",
        "### 1.1 文档目的",
        "",
        spec["purpose"],
        "",
        "### 1.2 开发背景",
        "",
        spec["background"],
        "",
        "开发模式：边分析边迭代，分模块生成可交互原型。覆盖范围：本科生（若本菜单另有范围未指定，标未确认）。",
        "",
        "### 1.3 填写说明",
        "",
        "- 条款：（已明确）/（未确认）。不得把 AI 发挥标成已明确。",
        "- 文档信息「确认状态」是整份签核，与条款口径不是同一件事。",
        "",
        "## 2. 模块需求",
        "",
        "### 2.1 菜单路径与范围",
        "",
        f"- 一级菜单：{m['parent']}（{m['parent_en']}）",
        f"- 二级菜单：{m['cn']}（{m['en']}）",
        f"- 范围内：{spec['intro']}",
        "",
        "### 2.2 建设目标与非目标",
        "",
        f"目标：{spec['purpose']}",
        "",
        f"非目标：{spec.get('non_goals') or '见正文；未列出的不做。'}",
        "",
        "### 2.3 角色与权限",
        "",
        "| 角色 | 场景与能力 | 口径 |",
        "|------|------------|------|",
    ]
    for role, cap, mark in spec.get("roles") or [("见上一版", "以原型按钮为准", "未确认")]:
        lines.append(f"| {role} | {cap} | {mark} |")
    lines += [
        "",
        "### 2.4 菜单介绍",
        "",
        spec["intro"],
        "",
        "### 2.5 界面说明",
        "",
        "**筛选区**",
        "",
        spec["search_fields"],
        "",
        "**列表列**",
        "",
        spec["list_fields"],
        "",
        "### 2.6 数据前后流转",
        "",
        f"1）关系说明：{spec['flow_rel']}",
        "",
        f"2）前置条件：{spec['flow_pre']}",
        "",
        f"3）下游输出：{spec['flow_out']}",
        "",
        "### 2.7 业务流",
        "",
        spec["biz_flow"],
        "",
        "### 2.8 原型参考",
        "",
        f"`prototype/index.html` → `{m['page']}`",
        "",
        "### 2.9 字段信息表",
        "",
        "字段以当前原型可见控件为准（已明确对照原型）。",
        "",
    ]
    for title, rows in field_tables:
        lines.append(f"#### {title}")
        lines.append("")
        lines.append("| " + " | ".join(h.replace("\n", " ") for h in FIELD_HEADERS) + " |")
        lines.append("|" + "|".join(["---"] * len(FIELD_HEADERS)) + "|")
        for row in rows:
            cells = list(row) + [""] * (len(FIELD_HEADERS) - len(row))
            lines.append("| " + " | ".join(str(c).replace("\n", " ") for c in cells[: len(FIELD_HEADERS)]) + " |")
        lines.append("")
    if hidden:
        lines += ["**本期隐藏/注释字段**", "", hidden, ""]
    lines += ["### 2.10 功能清单", ""]
    for item in functions:
        num, cn_name, en_name, desc, interaction, notes, drill = item[:7]
        lines += [
            f"**{num}、功能按钮——{cn_name}（英文名称：{en_name}）**",
            "",
            f"- 描述：{desc}",
            f"- 业务的事件交互：{interaction}",
            f"- 备注：{notes}",
            f"- 下钻页面说明：{drill}",
            "",
        ]
    pending = spec.get("pending") or ""
    extra = extra_pending(m["cn"])
    lines += ["### 2.11 未确认事项", "", pending or "· 无则填无。", ""]
    if extra:
        lines += [extra, ""]
    if prev:
        lines += [
            "### 附录 上一有效版",
            "",
            f"{m['cn']}{prev[0]}{prev[1]}",
            "",
        ]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_prd_docx(path: Path, m: dict, spec: dict, field_tables, functions, hidden=""):
    doc = Document()
    configure_styles(doc)
    setup_section(doc, f"产品需求文档 — {m['cn']} · {m['ver']}")
    add_title_block(
        doc,
        f"厦大马来分校本科教务系统产品需求文档 — {m['cn']}",
        f"模板版本：V2　　文档版本：{m['ver']}　　创建日期：{DATE_DISP}",
    )
    add_heading(doc, "0. 文档信息", 1)
    prev_txt = f"{m['cn']}{m['prev'][0]}{m['prev'][1]}" if m["prev"] else "无（本菜单首版）"
    add_kv_table(
        doc,
        [
            ("需求文档名称", f"厦大马来分校本科教务系统产品需求文档 — {m['cn']}"),
            ("菜单路径", m["path"]),
            ("英文名称", m["en"]),
            ("文档版本", m["ver"]),
            ("创建/发布日期", CHANGE_DATE),
            ("上一有效版", prev_txt),
            ("填写人", ""),
            ("确认人", ""),
            ("确认状态（整份签核）", "□ 待确认　　□ 已确认"),
            ("原型页面", f"prototype/index.html → {m['page']}"),
        ],
    )
    add_heading(doc, "1. 文档概述", 1)
    add_heading(doc, "1.1 文档目的", 2)
    add_para(doc, spec["purpose"])
    add_heading(doc, "1.2 开发背景", 2)
    add_para(doc, spec["background"])
    add_para(doc, "开发模式：边分析边迭代，分模块生成可交互原型。覆盖范围：本科生。")
    add_heading(doc, "1.3 填写说明", 2)
    add_para(
        doc,
        "条款口径：用户指定且已讲清的标（已明确）；没有明确、没有聊过、由 AI 补写或可能遗漏的标（未确认）。不得把 AI 发挥标成已明确。",
    )
    add_para(doc, "文档信息中的确认状态表示整份是否经确认人核过，与条款口径不是同一件事。")

    add_heading(doc, "2. 模块需求", 1)
    add_heading(doc, "2.1 菜单路径与范围", 2)
    add_kv_table(
        doc,
        [
            ("一级菜单", f"{m['parent']}（{m['parent_en']}）"),
            ("二级菜单", f"{m['cn']}（{m['en']}）"),
            ("三级菜单", "无"),
            ("范围内", spec["intro"]),
        ],
    )
    add_heading(doc, "2.2 建设目标与非目标", 2)
    add_para(doc, "目标：" + spec["purpose"])
    add_para(doc, "非目标：" + (spec.get("non_goals") or "未列出的能力不在本菜单。"))
    add_heading(doc, "2.3 角色与权限", 2)
    roles = spec.get("roles") or [("见上一版/原型按钮", "未在本轮指定正式角色名", "未确认")]
    add_grid_table(doc, ["角色", "场景与能力", "口径"], roles, [4.0, 17.0, 4.6])
    add_heading(doc, "2.4 菜单介绍", 2)
    add_para(doc, spec["intro"])
    add_heading(doc, "2.5 界面说明", 2)
    add_para(doc, "筛选区", bold=True)
    add_para(doc, spec["search_fields"])
    add_para(doc, "列表列", bold=True)
    add_para(doc, spec["list_fields"])
    add_heading(doc, "2.6 数据前后流转", 2)
    add_kv_table(
        doc,
        [
            ("1）关系说明", spec["flow_rel"]),
            ("2）前置条件", spec["flow_pre"]),
            ("3）下游输出", spec["flow_out"]),
        ],
    )
    add_heading(doc, "2.7 业务流", 2)
    add_para(doc, spec["biz_flow"])
    add_heading(doc, "2.8 原型参考", 2)
    add_kv_table(doc, [("原型页面", f"prototype/index.html → {m['page']}")])
    add_heading(doc, "2.9 字段信息表", 2)
    add_hint(doc, "以当前原型可见控件为准（已明确对照原型）。")
    for i, (title, rows) in enumerate(field_tables, 1):
        add_caption(doc, f"表 2-{i} {title}")
        add_grid_table(doc, FIELD_HEADERS, rows, FIELD_WIDTHS)
    if hidden:
        add_para(doc, "本期隐藏/注释字段", bold=True)
        add_para(doc, hidden)
    add_heading(doc, "2.10 功能清单", 2)
    for item in functions:
        add_function_item(doc, *item[:7], filled=True)
    add_heading(doc, "2.11 未确认事项", 2)
    add_para(doc, spec.get("pending") or "无。")
    extra = extra_pending(m["cn"])
    if extra:
        add_para(doc, extra)
    add_heading(doc, "附录 A 本模块菜单路径", 1)
    add_grid_table(
        doc,
        ["一级目录", "一级目录英文名称", "二级目录", "二级目录英文名称", "三级目录", "备注说明"],
        [[m["parent"], m["parent_en"], m["cn"], m["en"], "—", m["path"]]],
        [3.2, 3.8, 3.6, 4.4, 2.0, 8.6],
    )
    doc.save(path)


def write_change_md(path: Path, m: dict, rows: list[dict]):
    prev = m["prev"]
    from_ver = prev[1] if prev else "无"
    lines = [
        f"# {m['cn']}——需求调整变更说明",
        "",
        f"> 模板：需求调整变更说明模板（简版）20260821V2  ",
        f"> 对照：{from_ver} → {m['ver']}",
        "",
        "## 1 版本信息",
        "",
        "| 项目 | 内容 |",
        "|------|------|",
        f"| 需求文档名称 | 厦大马来分校本科教务系统产品需求文档 — {m['cn']} |",
        f"| 上一有效版 → 本版 | {from_ver} → {m['ver']} |",
        f"| 变更日期 | {CHANGE_DATE} |",
        "| 填写人 | |",
        "| 确认人 | |",
        "| 确认状态（整份签核） | □ 待确认　□ 已确认 |",
        "",
        "## 2 本次改了什么（总览）",
        "",
        "| 序号 | 模块（菜单路径） | 功能 / 位置 | 变更类型 | 调整内容 | 状态 |",
        "|------|------------------|-------------|---------|----------|------|",
    ]
    for i, r in enumerate(rows, 1):
        lines.append(
            f"| {i} | {m['path']} | {r['where']} | {r['type']} | {r['content']} | {r['status']} |"
        )
    lines += [
        "",
        "## 3 前后对照",
        "",
        "总览已能说明的条目不展开。条款标记由「已确认/待确认」改为「已明确/未确认」。",
        "",
        "## 4 连带影响",
        "",
        "| 序号 | 影响到哪里 | 影响说明 | 需同步 |",
        "|------|------------|----------|--------|",
        "| 1 | 原型 | 以现行原型为准整理文档 | ☑原型 □开发 □测试 |",
        "| 2 | 开发/测试 | 本批以文档套版与已核对文案为主，不代替开发任务 | □原型 ☑开发 ☑测试 |",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_change_docx(path: Path, m: dict, rows: list[dict]):
    doc = Document()
    configure_styles(doc)
    setup_section(doc, f"需求调整变更说明 — {m['cn']} · {m['ver']}")
    prev = m["prev"]
    from_ver = prev[1] if prev else "无"
    add_title_block(
        doc,
        f"厦大马来分校本科教务系统——需求调整变更说明（{m['cn']}）",
        f"模板版本：V2　　{from_ver} → {m['ver']}　　{CHANGE_DATE}",
    )
    add_heading(doc, "1 版本信息", 1)
    add_kv_table(
        doc,
        [
            ("需求文档名称", f"厦大马来分校本科教务系统产品需求文档 — {m['cn']}"),
            ("上一有效版 → 本版", f"{from_ver} → {m['ver']}"),
            ("上一版文件", f"{m['cn']}{prev[0]}{prev[1]}" if prev else "无"),
            ("本版文件", folder_name(m)),
            ("变更日期", CHANGE_DATE),
            ("填写人", ""),
            ("确认人", ""),
            ("确认状态（整份签核）", "□ 待确认　　□ 已确认"),
        ],
        col_widths=(4.5, 21.1),
    )
    add_heading(doc, "2 本次改了什么（总览）", 1)
    add_caption(doc, "表 2-1 变更总览")
    table_rows = [
        [str(i), m["path"], r["where"], r["type"], r["content"], r["status"]]
        for i, r in enumerate(rows, 1)
    ]
    add_grid_table(
        doc,
        ["序号", "模块（菜单路径）", "功能 / 位置", "变更类型", "调整内容（从什么变成什么）", "状态"],
        table_rows,
        [1.4, 5.0, 4.0, 2.2, 9.6, 2.4],
    )
    add_heading(doc, "3 前后对照（可选）", 1)
    add_para(doc, "条款标记由「已确认/待确认」改为「已明确/未确认」。其余见总览。")
    add_heading(doc, "4 连带影响（可选）", 1)
    add_grid_table(
        doc,
        ["序号", "影响到哪里", "影响说明", "需同步"],
        [
            ["1", "原型", "以现行原型为准整理文档", "☑原型 □开发 □测试"],
            ["2", "开发/测试", "本批以文档套版与已核对文案为主", "□原型 ☑开发 ☑测试"],
        ],
        [1.6, 6.0, 12.4, 5.6],
    )
    doc.save(path)


def _rewrite_copy(cn: str, text: str) -> str:
    s = text or ""
    if cn == "校选课程管理":
        s = s.replace("通识选修开课", "选修开课").replace("通识选修计划", "选修开课计划")
    if cn == "开课安排":
        s = s.replace("Support 学院", "Support 专业")
    return s


def spec_from_old(m, old):
    menu = old["menu"]
    meta = old["meta"]
    cn = m["cn"]
    roles = [
        ("教务/学院/教师（上一版写「原型约定」）", "能力见上一版；本轮未重开权限评审", "未确认"),
    ]
    functions = menu.get("functions") or []
    return {
        "purpose": _rewrite_copy(cn, meta["purpose"]),
        "background": _rewrite_copy(cn, meta["background"]),
        "intro": _rewrite_copy(cn, menu["intro"]),
        "list_fields": _rewrite_copy(cn, menu["list_fields"]),
        "search_fields": _rewrite_copy(cn, menu["search_fields"]),
        "flow_rel": _rewrite_copy(cn, menu["flow_rel"]),
        "flow_pre": _rewrite_copy(cn, menu["flow_pre"]),
        "flow_out": _rewrite_copy(cn, menu["flow_out"]),
        "biz_flow": _rewrite_copy(cn, menu["biz_flow"]),
        "roles": roles,
        "non_goals": _rewrite_copy(cn, menu.get("non_goals") or ""),
        "pending": menu.get("pending") or "",
        "functions": functions,
    }


def main():
    written = []
    overview_rows = []
    for m in MENUS:
        d = out_dir(m)
        d.mkdir(parents=True, exist_ok=True)
        if m["kind"] == "bump":
            old = DOCS_BY_CN[m["cn"]]
            spec = spec_from_old(m, old)
            pack = MENU_FIELD_PACKS[m["cn"]]
            fields = patch_fields(m["cn"], pack["tables"])
            hidden = pack.get("hidden") or ""
            functions = spec["functions"]
            changes = bump_changes(m)
        else:
            spec = NEW_SPECS[m["cn"]]
            fields = spec["fields"]
            hidden = ""
            functions = spec["functions"]
            changes = new_changes(m)

        md_prd = d / f"{folder_name(m)}.md"
        docx_prd = d / f"{folder_name(m)}.docx"
        write_md(md_prd, m, spec, fields, functions, hidden)
        write_prd_docx(docx_prd, m, spec, fields, functions, hidden)
        written.append(docx_prd)

        if m["prev"]:
            ch_stem = f"{m['cn']}{m['prev'][0]}{m['prev'][1]}→{DATE}{m['ver']}变更说明"
        else:
            ch_stem = f"{m['cn']}无→{DATE}{m['ver']}变更说明"
        write_change_md(d / f"{ch_stem}.md", m, changes)
        write_change_docx(d / f"{ch_stem}.docx", m, changes)
        written.append(d / f"{ch_stem}.docx")
        for r in changes:
            overview_rows.append((m["path"], r["where"], r["type"], r["content"], r["status"]))

    batch = BASE / f"指定侧栏菜单需求{DATE}调整说明"
    batch.mkdir(parents=True, exist_ok=True)
    batch_md = batch / f"指定侧栏菜单需求{DATE}调整说明.md"
    lines = [
        "# 开课管理指定侧栏菜单——需求调整说明（汇总）",
        "",
        f"> 日期：{CHANGE_DATE}　模板：需求调整变更说明模板（简版）20260821V2",
        "",
        "本批覆盖用户点名的 9 个侧栏菜单。各菜单完整 PRD 与分册变更说明在各自版本文件夹内，本文件只做总览。",
        "",
        "## 1 版本信息",
        "",
        "| 项目 | 内容 |",
        "|------|------|",
        "| 范围 | 开课时间设置、校选课程管理、特殊课程设置、开课计划、开课安排、开课名单、Teaching Load、授课确认管理、授课确认（教师端） |",
        f"| 日期 | {CHANGE_DATE} |",
        "| 确认状态（整份签核） | □ 待确认　□ 已确认 |",
        "",
        "## 2 本次改了什么（总览）",
        "",
        "| 序号 | 模块 | 功能/位置 | 变更类型 | 调整内容 | 状态 |",
        "|------|------|-----------|---------|----------|------|",
    ]
    for i, (mod, where, typ, content, st) in enumerate(overview_rows, 1):
        lines.append(f"| {i} | {mod} | {where} | {typ} | {content} | {st} |")
    lines += [
        "",
        "## 3 产出路径",
        "",
        "开课设置 / 专业开课 / 开课清单 下各二级菜单的 `…20260821Vn/` 文件夹。",
        "",
        "## 4 连带影响",
        "",
        "文档套 V2 模板与条款口径；除已写明的 Support 专业、选修开课用词外，不把未讨论规则写成已明确。",
        "",
    ]
    batch_md.write_text("\n".join(lines), encoding="utf-8")

    doc = Document()
    configure_styles(doc)
    setup_section(doc, "指定侧栏菜单需求调整说明 · 20260821")
    add_title_block(
        doc,
        "厦大马来分校本科教务系统——需求调整变更说明（指定侧栏菜单汇总）",
        f"模板版本：V2　　日期：{CHANGE_DATE}",
    )
    add_heading(doc, "1 版本信息", 1)
    add_kv_table(
        doc,
        [
            ("范围", "用户点名的 9 个侧栏二级菜单"),
            ("变更日期", CHANGE_DATE),
            ("填写人", ""),
            ("确认人", ""),
            ("确认状态（整份签核）", "□ 待确认　　□ 已确认"),
        ],
    )
    add_heading(doc, "2 本次改了什么（总览）", 1)
    add_grid_table(
        doc,
        ["序号", "模块（菜单路径）", "功能 / 位置", "变更类型", "调整内容", "状态"],
        [
            [str(i), mod, where, typ, content, st]
            for i, (mod, where, typ, content, st) in enumerate(overview_rows, 1)
        ],
        [1.2, 5.2, 3.8, 2.0, 10.0, 2.4],
    )
    add_heading(doc, "3 分册位置", 1)
    add_para(doc, "各菜单 PRD 与变更说明见 `参考文档/2、开课管理/<一级>/<二级>/<菜单20260821Vn>/`。")
    add_heading(doc, "4 连带影响", 1)
    add_para(doc, "以现行原型为准落档；未讨论规则标（未确认）。开发任务不因本文自动成立。")
    batch_docx = batch / f"指定侧栏菜单需求{DATE}调整说明.docx"
    doc.save(batch_docx)
    written.append(batch_docx)
    for p in written:
        print(p.relative_to(ROOT), p.stat().st_size)


if __name__ == "__main__":
    main()
