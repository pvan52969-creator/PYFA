#!/usr/bin/env python3
"""2026-08-25 批次：V2 模板 PRD + 变更说明（开课设置/专业开课升版 + 课程班首版）。"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(SCRIPT_DIR))

from docx import Document

from prd_folder_paths import BASE, menu_dir

DATE = "20260825"
DATE_DISP = "2026 年 8 月 25 日"
CHANGE_DATE = "2026-08-25"


def load_mod(name: str, filename: str):
    path = SCRIPT_DIR / filename
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


tpl = load_mod("generate_prd_templates_v2", "generate-prd-templates-v2.py")
old_gen = load_mod("submenu_prd_old", "generate-submenu-prd-docx.py")
from prd_submenu_fields import MENU_FIELD_PACKS  # noqa: E402

FIELD_HEADERS = tpl.FIELD_HEADERS
FIELD_WIDTHS = tpl.FIELD_WIDTHS
add_caption = tpl.add_caption
add_function_item = tpl.add_function_item
add_grid_table = tpl.add_grid_table
add_heading = tpl.add_heading
add_hint = tpl.add_hint
add_kv_table = tpl.add_kv_table
add_para = tpl.add_para
add_title_block = tpl.add_title_block
configure_styles = tpl.configure_styles
setup_section = tpl.setup_section
fn = old_gen.fn

DOCS_BY_CN = {d["meta"]["cn"]: d for d in old_gen.DOCS}

MENUS = [
    {
        "cn": "特殊课程设置",
        "en": "Special Course Settings",
        "parent": "开课设置",
        "parent_en": "Offering Settings",
        "path": "开课管理 → 开课设置 → 特殊课程设置",
        "page": "page-special-course-settings",
        "prev": ("20260821", "V3"),
        "ver": "V4",
        "kind": "bump",
    },
    {
        "cn": "开课计划",
        "en": "Major Offering Plan",
        "parent": "专业开课",
        "parent_en": "Major Offering",
        "path": "开课管理 → 专业开课 → 开课计划",
        "page": "page-course-offering-major",
        "prev": ("20260821", "V5"),
        "ver": "V6",
        "kind": "bump",
    },
    {
        "cn": "开课安排",
        "en": "Major Offering Task Arrangement",
        "parent": "专业开课",
        "parent_en": "Major Offering",
        "path": "开课管理 → 专业开课 → 开课安排",
        "page": "page-course-major-offering-task-style2",
        "prev": ("20260821", "V5"),
        "ver": "V6",
        "kind": "bump",
    },
    {
        "cn": "课程班",
        "en": "Course Section Manifest",
        "parent": "课程班管理",
        "parent_en": "Course Section Management",
        "path": "开课管理 → 课程班管理 → 课程班",
        "page": "page-course-offering-manifest",
        "prev": None,
        "ver": "V1",
        "kind": "new",
    },
]

EXTRA_TASK_FUNCTIONS = [
    fn(
        "12",
        "一键复制",
        "One-click Copy",
        "按所选学期同课号复制分组与学时/教师/教室配置。",
        "工具栏打开弹窗。",
        "已生效或授课确认中任务跳过。",
        "弹窗：一键复制以往学期安排。",
    ),
    fn(
        "13",
        "一键同步共同授课",
        "Sync Shared Teaching",
        "按当前列表可见任务，依据特殊课程关联课号与共同教师同步共同授课绑定。",
        "工具栏（步骤3）。",
        "仅无共同教师时解绑；保护手工非预置绑定。",
        "确认框。",
    ),
]

NEW_SPECS = {
    "课程班": {
        "purpose": "汇总专业开课、选修开课、特殊开课中已开出的课程班，支持统一查阅、修读范围查看、名单下钻与批量停课。",
        "background": "侧栏「课程班管理」分组下的汇总页；原「开课清单」能力迁移并扩展。"
        "各开课安排页不再提供停课，停课统一在本页操作。",
        "intro": "自动汇总三类开课流程中已开出的课程班（含草稿、已生效与停课），一门课一行。"
        "支持展开查看分组/学时安排子表；修读范围：选修(GE)走完整 GE 弹窗，专业走简化弹窗（含开课模块、选课方式、专业批次），特殊开课不适用。"
        "当前学生名单人数可点击打开名单抽屉；课程信息列查看各模块课程设置；操作列进入分组详情。",
        "list_fields": "勾选、开课详情（展开）、开课状态、授课确认、生效状态、开课学期、开课模块、课程号、课程名称、所属专业、选课类型、学分、小组数、总学时、课程人数上限、当前学生名单人数、修读范围、课程信息、操作（分组详情）。"
        "左侧冻结至课程号；右侧冻结：修读范围、课程信息、操作。",
        "search_fields": "开课学期、开课模块（专业/选修GE/选修ME/特殊）、所属专业、课号、课名；查询/重置。",
        "flow_rel": "读取各学期 COURSE_OFFERING_PLAN 中 major/ge/other 已开出 section；停课回写各开课模块状态；名单抽屉复用专业开课名单组件。",
        "flow_pre": "至少一条已开出课程班（计划/任务生成后即有草稿行）。",
        "flow_out": "停课结果同步至对应开课安排页；名单查看供教务核对；不修改变更计划/安排本体（除停课）。",
        "biz_flow": "进入【课程班】→ 筛选 → 展开查看安排 / 查看修读范围 / 点开名单人数 → 勾选已安排教师的班 → 停课。",
        "roles": [("教务/学院（未确认正式角色名）", "查询、查看、停课（对照原型按钮）", "未确认")],
        "non_goals": "不在本页维护计划生成、任务三步安排、共同授课码批量设置（工具栏已移除）。",
        "pending": "· 停课与排课/选课应用的拦截矩阵（未确认）。\n· ME/GE 筛选文案是否需产品统一（未确认）。",
        "functions": [
            fn("1", "查询/重置", "Query/Reset", "按筛选刷新列表。", "筛选区按钮。", "—", "无。"),
            fn(
                "2",
                "停课",
                "Cancel Offering",
                "对已安排教师（含未生效）的课程班批量停课；停课后课时不再计入教师开课学时。",
                "工具栏左侧红色按钮；需勾选。",
                "未安排教师不可停；停课后各开课模块列表同步。",
                "确认框。",
            ),
            fn(
                "3",
                "展开开课详情",
                "Expand Arrangement",
                "展开分组/学时投递子表。",
                "开课详情列 +/- 按钮。",
                "—",
                "内嵌安排子表。",
            ),
            fn(
                "4",
                "查看修读范围",
                "View Scope",
                "GE 完整弹窗；专业简化弹窗；特殊开课禁用。",
                "修读范围列链接。",
                "—",
                "弹窗：修读范围查看。",
            ),
            fn(
                "5",
                "当前学生名单人数",
                "Roster Count",
                "打开学生名单抽屉。",
                "名单人数列链接。",
                "—",
                "抽屉：学生名单。",
            ),
            fn(
                "6",
                "课程信息",
                "Course Settings",
                "按开课模块打开对应只读/编辑课程设置视图。",
                "课程信息列「查看」。",
                "—",
                "各模块任务查看抽屉。",
            ),
            fn(
                "7",
                "分组详情",
                "Grouping Detail",
                "进入分组工作台（课程班/小组学时详情 Tab）。",
                "操作列链接。",
                "—",
                "分组工作台。",
            ),
        ],
    },
}

CHANGE_ROWS = {
    "特殊课程设置": [
        {
            "where": "列表上方说明条",
            "type": "删除",
            "content": "去掉 legend 说明条；默认生效范围保留在页头 subtitle",
            "status": "已明确",
        },
    ],
    "开课计划": [
        {
            "where": "修改教学任务抽屉",
            "type": "修改",
            "content": "字段标签「选修/必修」→「选课方式」",
            "status": "已明确",
        },
    ],
    "开课安排": [
        {
            "where": "工具栏·停课",
            "type": "删除",
            "content": "去掉「停课」按钮；停课改由「课程班」页统一操作",
            "status": "已明确",
        },
        {
            "where": "工具栏",
            "type": "澄清",
            "content": "保留：生效、撤回、退回、导入、一键复制、一键同步共同授课",
            "status": "已明确",
        },
        {
            "where": "数据流转",
            "type": "修改",
            "content": "下游「开课清单」表述改为「课程班」汇总；停课入口迁至课程班",
            "status": "已明确",
        },
        {
            "where": "功能清单",
            "type": "新增",
            "content": "补充「一键复制」「一键同步共同授课」功能说明",
            "status": "已明确",
        },
    ],
    "课程班": [
        {
            "where": "文档",
            "type": "新增",
            "content": "首版 PRD；汇总已开出课程班，承担统一停课入口（对照 page-course-offering-manifest）",
            "status": "已明确",
        },
        {
            "where": "侧栏结构",
            "type": "澄清",
            "content": "一级菜单「课程班管理」下二级菜单「课程班」；原开课清单侧栏分组更名",
            "status": "已明确",
        },
        {
            "where": "列表",
            "type": "新增",
            "content": "修读范围查看、名单人数抽屉、开课详情展开、左右列冻结等（对照现行原型）",
            "status": "已明确",
        },
    ],
}


def folder_name(m):
    return f"{m['cn']}{DATE}{m['ver']}"


def out_dir(m):
    return menu_dir(m["cn"])


def patch_fields(cn: str, tables):
    out = []
    for title, rows in tables:
        new_rows = []
        for row in rows:
            cells = [str(c) for c in row]
            if cn == "开课安排":
                cells = [
                    c.replace("Support 学院", "Support 专业").replace("Support学院", "Support 专业")
                    for c in cells
                ]
            new_rows.append(cells)
        out.append((title, new_rows))
    return out


def spec_from_old(m, old):
    menu = old["menu"]
    meta = old["meta"]
    cn = m["cn"]
    functions = list(menu.get("functions") or [])
    if cn == "开课安排":
        existing = {f[1] for f in functions}
        for extra in EXTRA_TASK_FUNCTIONS:
            if extra[1] not in existing:
                functions.append(extra)
    return {
        "purpose": meta["purpose"],
        "background": meta["background"],
        "intro": menu["intro"],
        "list_fields": menu["list_fields"],
        "search_fields": menu["search_fields"],
        "flow_rel": menu["flow_rel"],
        "flow_pre": menu["flow_pre"],
        "flow_out": menu["flow_out"],
        "biz_flow": menu["biz_flow"],
        "roles": [
            ("教务/学院/教师（上一版写「原型约定」）", "能力见上一版；本轮未重开权限评审", "未确认"),
        ],
        "non_goals": menu.get("non_goals") or "",
        "pending": menu.get("pending") or "",
        "functions": functions,
    }


def apply_bump_overrides(cn: str, spec: dict) -> dict:
    if cn == "特殊课程设置":
        spec["intro"] = (
            "从教务课程库手动纳入课号，维护四学时教室属性/偏好、默认 Support 学院、Support 历史（Y/N）。"
            "默认仅影响此后新建或仍为系统默认的教学班；已改过的班不覆盖；班级覆盖不回写课号配置。"
            "列表区上方不再单独展示 legend 说明条，默认生效范围说明保留在页头 subtitle。"
        )
        spec["list_fields"] = (
            "勾选、序号、课程代码、课程名称、开课单位、Support 学院、Support历史、操作（编辑）。"
            "教室仅在编辑抽屉维护，列表无教室摘要列；列表上方无 legend 说明条。"
        )
    if cn == "开课计划":
        spec["pending"] = (
            (spec.get("pending") or "")
            + "\n· 修改教学任务抽屉字段标签已改为「选课方式」（已明确）。"
        ).strip()
        new_functions = []
        for item in spec["functions"]:
            num, cn_name, en_name, desc, interaction, notes, drill = item[:7]
            if cn_name == "修改教学任务":
                notes = (notes or "") + " 抽屉内原「选修/必修」标签改为「选课方式」。"
            new_functions.append((num, cn_name, en_name, desc, interaction, notes, drill))
        spec["functions"] = new_functions
    if cn == "开课安排":
        spec["intro"] = (
            spec["intro"]
            + " 工具栏含生效、撤回、退回、导入、一键复制、一键同步共同授课；不含停课（停课改在「课程班」页）。"
        )
        spec["flow_out"] = spec["flow_out"].replace("开课清单", "课程班")
        spec["pending"] = (
            (spec.get("pending") or "")
            + "\n· 停课不在本页维护；统一改由「课程班」页操作（已明确）。"
        ).strip()
    return spec


def write_md(path: Path, m: dict, spec: dict, field_tables, functions, hidden=""):
    prev = m["prev"]
    lines = [
        f"# 厦大马来分校本科教务系统产品需求文档 — {m['cn']}",
        "",
        f"> **文档版本：** {m['ver']} · **创建日期：** {DATE_DISP}  ",
        f"> **模板：** 产品需求文档模板（空白/带参考数据版）20260821V2  ",
        f"> **原型页面：** `{m['page']}`  ",
        "> **条款口径：** 用户指定且已讲清 →（已明确）；未聊过、AI 补写或可能遗漏 →（未确认）",
        "",
        "| 项 | 内容 |",
        "|---|---|",
        f"| 菜单路径 | {m['path']} |",
        f"| 英文名称 | {m['en']} |",
        f"| 文档版本 | {m['ver']} |",
        f"| 创建日期 | {CHANGE_DATE} |",
        f"| 原型页面 | `{m['page']}` |",
        "",
        "## 1. 文档概述",
        "",
        "### 1.1 文档目的",
        "",
        spec["purpose"],
        "",
        "### 1.2 开发背景",
        "",
        spec["background"],
        "",
        "开发模式：边分析边迭代，分模块生成可交互原型。覆盖范围：本科生（若本菜单另有范围未指定，标未确认）。",
        "",
        "### 1.3 填写说明",
        "",
        "- 条款：（已明确）/（未确认）。不得把 AI 发挥标成已明确。",
        "",
        "## 2. 模块需求",
        "",
        "### 2.1 菜单路径与范围",
        "",
        f"- 一级菜单：{m['parent']}（{m['parent_en']}）",
        f"- 二级菜单：{m['cn']}（{m['en']}）",
        f"- 范围内：{spec['intro']}",
        "",
        "### 2.2 建设目标与非目标",
        "",
        f"目标：{spec['purpose']}",
        "",
        f"非目标：{spec.get('non_goals') or '见正文；未列出的不做。'}",
        "",
        "### 2.3 角色与权限",
        "",
        "| 角色 | 场景与能力 | 口径 |",
        "|------|------------|------|",
    ]
    for role, cap, mark in spec.get("roles") or [("见上一版", "以原型按钮为准", "未确认")]:
        lines.append(f"| {role} | {cap} | {mark} |")
    lines += [
        "",
        "### 2.4 菜单介绍",
        "",
        spec["intro"],
        "",
        "### 2.5 界面说明",
        "",
        "**筛选区**",
        "",
        spec["search_fields"],
        "",
        "**列表列**",
        "",
        spec["list_fields"],
        "",
        "### 2.6 数据前后流转",
        "",
        f"1）关系说明：{spec['flow_rel']}",
        "",
        f"2）前置条件：{spec['flow_pre']}",
        "",
        f"3）下游输出：{spec['flow_out']}",
        "",
        "### 2.7 业务流",
        "",
        spec["biz_flow"],
        "",
        "### 2.8 原型参考",
        "",
        f"`prototype/index.html` → `{m['page']}`",
        "",
        "### 2.9 字段信息表",
        "",
        "字段以当前原型可见控件为准（已明确对照原型）。",
        "",
    ]
    for title, rows in field_tables:
        lines.append(f"#### {title}")
        lines.append("")
        lines.append("| " + " | ".join(h.replace("\n", " ") for h in FIELD_HEADERS) + " |")
        lines.append("|" + "|".join(["---"] * len(FIELD_HEADERS)) + "|")
        for row in rows:
            cells = list(row) + [""] * (len(FIELD_HEADERS) - len(row))
            lines.append("| " + " | ".join(str(c).replace("\n", " ") for c in cells[: len(FIELD_HEADERS)]) + " |")
        lines.append("")
    if hidden:
        lines += ["**本期隐藏/注释字段**", "", hidden, ""]
    lines += ["### 2.10 功能清单", ""]
    for item in functions:
        num, cn_name, en_name, desc, interaction, notes, drill = item[:7]
        lines += [
            f"**{num}、功能按钮——{cn_name}（英文名称：{en_name}）**",
            "",
            f"- 描述：{desc}",
            f"- 业务的事件交互：{interaction}",
            f"- 备注：{notes}",
            f"- 下钻页面说明：{drill}",
            "",
        ]
    lines += ["### 2.11 未确认事项", "", spec.get("pending") or "· 无则填无。", ""]
    if prev:
        lines += ["### 附录 上一有效版", "", f"{m['cn']}{prev[0]}{prev[1]}", ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_prd_docx(path: Path, m: dict, spec: dict, field_tables, functions, hidden=""):
    doc = Document()
    configure_styles(doc)
    setup_section(doc, f"产品需求文档 — {m['cn']} · {m['ver']}")
    add_title_block(
        doc,
        f"厦大马来分校本科教务系统产品需求文档 — {m['cn']}",
        f"模板版本：V2　　文档版本：{m['ver']}　　创建日期：{DATE_DISP}",
    )
    add_heading(doc, "0. 文档信息", 1)
    prev_txt = f"{m['cn']}{m['prev'][0]}{m['prev'][1]}" if m["prev"] else "无（本菜单首版）"
    add_kv_table(
        doc,
        [
            ("需求文档名称", f"厦大马来分校本科教务系统产品需求文档 — {m['cn']}"),
            ("菜单路径", m["path"]),
            ("英文名称", m["en"]),
            ("文档版本", m["ver"]),
            ("创建/发布日期", CHANGE_DATE),
            ("上一有效版", prev_txt),
            ("原型页面", f"prototype/index.html → {m['page']}"),
        ],
    )
    add_heading(doc, "1. 文档概述", 1)
    add_heading(doc, "1.1 文档目的", 2)
    add_para(doc, spec["purpose"])
    add_heading(doc, "1.2 开发背景", 2)
    add_para(doc, spec["background"])
    add_para(doc, "开发模式：边分析边迭代，分模块生成可交互原型。覆盖范围：本科生。")
    add_heading(doc, "1.3 填写说明", 2)
    add_para(
        doc,
        "条款口径：用户指定且已讲清的标（已明确）；没有明确、没有聊过、由 AI 补写或可能遗漏的标（未确认）。不得把 AI 发挥标成已明确。",
    )

    add_heading(doc, "2. 模块需求", 1)
    add_heading(doc, "2.1 菜单路径与范围", 2)
    add_kv_table(
        doc,
        [
            ("一级菜单", f"{m['parent']}（{m['parent_en']}）"),
            ("二级菜单", f"{m['cn']}（{m['en']}）"),
            ("三级菜单", "无"),
            ("范围内", spec["intro"]),
        ],
    )
    add_heading(doc, "2.2 建设目标与非目标", 2)
    add_para(doc, "目标：" + spec["purpose"])
    add_para(doc, "非目标：" + (spec.get("non_goals") or "未列出的能力不在本菜单。"))
    add_heading(doc, "2.3 角色与权限", 2)
    roles = spec.get("roles") or [("见上一版/原型按钮", "未在本轮指定正式角色名", "未确认")]
    add_grid_table(doc, ["角色", "场景与能力", "口径"], roles, [4.0, 17.0, 4.6])
    add_heading(doc, "2.4 菜单介绍", 2)
    add_para(doc, spec["intro"])
    add_heading(doc, "2.5 界面说明", 2)
    add_para(doc, "筛选区", bold=True)
    add_para(doc, spec["search_fields"])
    add_para(doc, "列表列", bold=True)
    add_para(doc, spec["list_fields"])
    add_heading(doc, "2.6 数据前后流转", 2)
    add_kv_table(
        doc,
        [
            ("1）关系说明", spec["flow_rel"]),
            ("2）前置条件", spec["flow_pre"]),
            ("3）下游输出", spec["flow_out"]),
        ],
    )
    add_heading(doc, "2.7 业务流", 2)
    add_para(doc, spec["biz_flow"])
    add_heading(doc, "2.8 原型参考", 2)
    add_kv_table(doc, [("原型页面", f"prototype/index.html → {m['page']}")])
    add_heading(doc, "2.9 字段信息表", 2)
    add_hint(doc, "以当前原型可见控件为准（已明确对照原型）。")
    for i, (title, rows) in enumerate(field_tables, 1):
        add_caption(doc, f"表 2-{i} {title}")
        add_grid_table(doc, FIELD_HEADERS, rows, FIELD_WIDTHS)
    if hidden:
        add_para(doc, "本期隐藏/注释字段", bold=True)
        add_para(doc, hidden)
    add_heading(doc, "2.10 功能清单", 2)
    for item in functions:
        add_function_item(doc, *item[:7], filled=True)
    add_heading(doc, "2.11 未确认事项", 2)
    add_para(doc, spec.get("pending") or "无。")
    add_heading(doc, "附录 A 本模块菜单路径", 1)
    add_grid_table(
        doc,
        ["一级目录", "一级目录英文名称", "二级目录", "二级目录英文名称", "三级目录", "备注说明"],
        [[m["parent"], m["parent_en"], m["cn"], m["en"], "—", m["path"]]],
        [3.2, 3.8, 3.6, 4.4, 2.0, 8.6],
    )
    doc.save(path)


def write_change_md(path: Path, m: dict, rows: list[dict]):
    prev = m["prev"]
    from_ver = prev[1] if prev else "无"
    lines = [
        f"# {m['cn']}——需求调整变更说明",
        "",
        f"> 模板：需求调整变更说明模板（简版）20260821V2  ",
        f"> 对照：{from_ver} → {m['ver']}",
        "",
        "## 1 版本信息",
        "",
        "| 项目 | 内容 |",
        "|------|------|",
        f"| 需求文档名称 | 厦大马来分校本科教务系统产品需求文档 — {m['cn']} |",
        f"| 上一有效版 → 本版 | {from_ver} → {m['ver']} |",
        f"| 变更日期 | {CHANGE_DATE} |",
        "",
        "## 2 本次改了什么（总览）",
        "",
        "| 序号 | 模块（菜单路径） | 功能 / 位置 | 变更类型 | 调整内容 | 状态 |",
        "|------|------------------|-------------|---------|----------|------|",
    ]
    for i, r in enumerate(rows, 1):
        lines.append(
            f"| {i} | {m['path']} | {r['where']} | {r['type']} | {r['content']} | {r['status']} |"
        )
    lines += [
        "",
        "## 3 前后对照（可选）",
        "",
        "总览已能说明的条目不展开。",
        "",
        "## 4 连带影响（可选）",
        "",
        "| 序号 | 影响到哪里 | 影响说明 | 需同步 |",
        "|------|------------|----------|--------|",
        "| 1 | 原型 | 以现行原型为准整理文档 | ☑原型 □开发 □测试 |",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_change_docx(path: Path, m: dict, rows: list[dict]):
    doc = Document()
    configure_styles(doc)
    setup_section(doc, f"需求调整变更说明 — {m['cn']} · {m['ver']}")
    prev = m["prev"]
    from_ver = prev[1] if prev else "无"
    add_title_block(
        doc,
        f"厦大马来分校本科教务系统——需求调整变更说明（{m['cn']}）",
        f"模板版本：V2　　{from_ver} → {m['ver']}　　{CHANGE_DATE}",
    )
    add_heading(doc, "1 版本信息", 1)
    add_kv_table(
        doc,
        [
            ("需求文档名称", f"厦大马来分校本科教务系统产品需求文档 — {m['cn']}"),
            ("上一有效版 → 本版", f"{from_ver} → {m['ver']}"),
            ("上一版文件", f"{m['cn']}{prev[0]}{prev[1]}" if prev else "无"),
            ("本版文件", folder_name(m)),
            ("变更日期", CHANGE_DATE),
        ],
        col_widths=(4.5, 21.1),
    )
    add_heading(doc, "2 本次改了什么（总览）", 1)
    add_caption(doc, "表 2-1 变更总览")
    table_rows = [
        [str(i), m["path"], r["where"], r["type"], r["content"], r["status"]]
        for i, r in enumerate(rows, 1)
    ]
    add_grid_table(
        doc,
        ["序号", "模块（菜单路径）", "功能 / 位置", "变更类型", "调整内容（从什么变成什么）", "状态"],
        table_rows,
        [1.4, 5.0, 4.0, 2.2, 9.6, 2.4],
    )
    add_heading(doc, "3 前后对照（可选）", 1)
    add_para(doc, "见总览；复杂条目可在评审时补充。")
    add_heading(doc, "4 连带影响（可选）", 1)
    add_grid_table(
        doc,
        ["序号", "影响到哪里", "影响说明", "需同步"],
        [["1", "原型", "以现行原型为准整理文档", "☑原型 □开发 □测试"]],
        [1.6, 6.0, 12.4, 5.6],
    )
    doc.save(path)


def write_summary():
    batch = BASE / "00_规范与变更说明" / "20260825_开课设置与专业开课需求调整说明"
    batch.mkdir(parents=True, exist_ok=True)
    md = batch / "开课设置与专业开课需求20260825调整说明.md"
    lines = [
        "# 开课设置 · 专业开课 · 课程班管理——需求调整说明（汇总）",
        "",
        f"> 日期：{CHANGE_DATE}　模板：需求调整变更说明模板（简版）20260821V2",
        "",
        "本批对照近期原型调整，更新有改动的二级菜单 PRD 与变更说明。目录结构：一级菜单一个文件夹，二级菜单其下再建文件夹。",
        "",
        "## 1 版本信息",
        "",
        "| 项目 | 内容 |",
        "|------|------|",
        f"| 范围 | **开课设置**：特殊课程设置（V3→V4）<br>**专业开课**：开课计划（V5→V6）、开课安排（V5→V6）<br>**课程班管理**：课程班（首版 V1） |",
        "| 未升版 | 开课时间设置（V3）、校选课程管理（V2）、开课名单（V3）——本批原型无变更 |",
        f"| 日期 | {CHANGE_DATE} |",
        "",
        "## 2 本次改了什么（总览）",
        "",
        "| 序号 | 模块 | 功能/位置 | 变更类型 | 调整内容 | 状态 |",
        "|------|------|-----------|---------|----------|------|",
        "| 1 | 开课管理 → 开课设置 → 特殊课程设置 | 列表上方说明条 | 删除 | 去掉 legend 说明条 | 已明确 |",
        "| 2 | 开课管理 → 专业开课 → 开课计划 | 修改教学任务抽屉 | 修改 | 「选修/必修」→「选课方式」 | 已明确 |",
        "| 3 | 开课管理 → 专业开课 → 开课安排 | 工具栏·停课 | 删除 | 停课改在「课程班」页 | 已明确 |",
        "| 4 | 开课管理 → 课程班管理 → 课程班 | 文档 | 新增 | 首版 PRD：汇总已开出课程班、统一停课 | 已明确 |",
        "",
        "## 3 产出路径",
        "",
        "| 一级菜单 | 二级菜单 | 新版本文件夹 |",
        "|----------|----------|-------------|",
        "| 01_开课设置 | 03_特殊课程设置 | `01_开课设置/03_特殊课程设置/特殊课程设置20260825V4/` |",
        "| 02_专业开课 | 01_开课计划 | `02_专业开课/01_开课计划/开课计划20260825V6/` |",
        "| 02_专业开课 | 02_开课安排 | `02_专业开课/02_开课安排/开课安排20260825V6/` |",
        "| 05_课程班管理 | 01_课程班 | `05_课程班管理/01_课程班/课程班20260825V1/` |",
        "",
        "各文件夹内含 PRD（md + docx，V2 模板）与变更说明（md + docx）。",
        "",
        "## 4 连带影响",
        "",
        "- 选修/特殊开课安排工具栏亦去掉停课，逻辑与专业开课安排一致。",
        "- Teaching Load、授课确认等已迁至 `05_课程班管理/` 下对应序号文件夹；历史 PRD 菜单路径文案待升版时更正。",
    ]
    md.write_text("\n".join(lines), encoding="utf-8")

    doc = Document()
    configure_styles(doc)
    setup_section(doc, "开课设置与专业开课需求20260825调整说明")
    add_title_block(
        doc,
        "开课设置 · 专业开课 · 课程班管理——需求调整说明（汇总）",
        f"模板版本：V2　　日期：{CHANGE_DATE}",
    )
    add_heading(doc, "1 版本信息", 1)
    add_para(doc, "范围：特殊课程设置 V4；开课计划/开课安排 V6；课程班 V1（首版）。")
    add_heading(doc, "2 本次改了什么", 1)
    add_para(doc, "详见各二级菜单版本文件夹内变更说明。")
    add_heading(doc, "3 产出路径", 1)
    add_para(doc, "一级菜单一个文件夹；二级菜单其下再建版本文件夹。")
    doc.save(batch / "开课设置与专业开课需求20260825调整说明.docx")
    return md


def main():
    written = []
    for m in MENUS:
        d = out_dir(m)
        d.mkdir(parents=True, exist_ok=True)
        if m["kind"] == "bump":
            old = DOCS_BY_CN[m["cn"]]
            spec = apply_bump_overrides(m["cn"], spec_from_old(m, old))
            pack = MENU_FIELD_PACKS[m["cn"]]
            fields = patch_fields(m["cn"], pack["tables"])
            hidden = pack.get("hidden") or ""
            functions = spec["functions"]
            changes = CHANGE_ROWS[m["cn"]]
        else:
            spec = NEW_SPECS[m["cn"]]
            pack = MENU_FIELD_PACKS[m["cn"]]
            fields = pack["tables"]
            hidden = pack.get("hidden") or ""
            functions = spec["functions"]
            changes = CHANGE_ROWS[m["cn"]]

        md_prd = d / f"{folder_name(m)}.md"
        docx_prd = d / f"{folder_name(m)}.docx"
        write_md(md_prd, m, spec, fields, functions, hidden)
        write_prd_docx(docx_prd, m, spec, fields, functions, hidden)
        written.extend([md_prd, docx_prd])

        if m["prev"]:
            ch_stem = f"{m['cn']}{m['prev'][0]}{m['prev'][1]}→{DATE}{m['ver']}变更说明"
        else:
            ch_stem = f"{m['cn']}无→{DATE}{m['ver']}变更说明"
        ch_md = d / f"{ch_stem}.md"
        ch_docx = d / f"{ch_stem}.docx"
        write_change_md(ch_md, m, changes)
        write_change_docx(ch_docx, m, changes)
        written.extend([ch_md, ch_docx])
        print(f"OK {m['cn']} {m['ver']}")

    summary = write_summary()
    written.append(summary)
    print(f"Done: {len(written)} files")


if __name__ == "__main__":
    main()
