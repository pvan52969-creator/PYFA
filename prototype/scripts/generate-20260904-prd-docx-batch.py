#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成本批 PRD / 变更说明 docx：md→docx + 变更说明填 V3 模板。"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent

TEMPLATE = ROOT / "参考文档" / "0、模板" / "需求调整变更说明模板（简版）20260901V3.docx"


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text or ""))
    run.bold = bold
    run.font.size = Pt(10.5)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def fill_kv(table, mapping: dict[str, str]):
    for row in table.rows:
        label = (row.cells[0].text or "").strip().replace("\n", "")
        for key, val in mapping.items():
            if key in label:
                set_cell_text(row.cells[1], val)
                break


def parse_overview_rows(md_text: str) -> list[dict]:
    """从变更说明 md「## 2 本次改了什么」表解析行。"""
    m = re.search(r"## 2 .*?\n\n(\|.+\n\|[-| :]+\n(?:\|.+\n)+)", md_text)
    if not m:
        return []
    lines = [ln for ln in m.group(1).splitlines() if ln.startswith("|")]
    if len(lines) < 3:
        return []
    rows = []
    for ln in lines[2:]:
        cells = [c.strip() for c in ln.strip("|").split("|")]
        if len(cells) < 8:
            continue
        rows.append(
            {
                "module": cells[1],
                "obj_type": cells[2],
                "obj_id": cells[3],
                "where": cells[4],
                "type": cells[5],
                "content": cells[6],
                "status": cells[7],
            }
        )
    return rows


def write_change_note_docx(md_path: Path, out_path: Path, meta: dict, rows: list[dict]):
    if out_path.exists():
        raise SystemExit(f"拒绝覆盖：{out_path}")
    doc = Document(str(TEMPLATE))
    fill_kv(
        doc.tables[0],
        {
            "需求文档名称": meta["name"],
            "菜单路径": meta["path"],
            "上一有效版": meta["versions"],
            "上一版 → 本版": meta["versions"],
            "上一版文件": meta["from_file"],
            "本版文件": meta["to_file"],
            "本版 PRD 模板": meta.get("prd_template", "V3.1"),
            "原型地址": meta.get("prototype", "—"),
            "变更日期": meta.get("date", "2026-09-04"),
        },
    )
    t1 = doc.tables[1]
    data_rows = t1.rows[1:]
    fill = rows or [
        {
            "module": meta["path"],
            "obj_type": "",
            "obj_id": "—",
            "where": "",
            "type": "",
            "content": "（空）",
            "status": "未确认",
        }
    ]
    # 模板行不够时追加
    while len(data_rows) < len(fill):
        t1.add_row()
        data_rows = t1.rows[1:]
    for i, row in enumerate(data_rows):
        if i >= len(fill):
            for c in row.cells:
                set_cell_text(c, "")
            continue
        r = fill[i]
        vals = [
            str(i + 1),
            r.get("module", ""),
            r.get("obj_type", ""),
            r.get("obj_id", "—"),
            r.get("where", ""),
            r.get("type", ""),
            r.get("content", ""),
            r.get("status", "未确认"),
        ]
        for cell, val in zip(row.cells, vals):
            set_cell_text(cell, val)
    doc.save(out_path)


def main():
    # import md_to_docx from generate-design-docx (hyphenated module name)
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "generate_design_docx", SCRIPT_DIR / "generate-design-docx.py"
    )
    mod = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(mod)
    md_to_docx = mod.md_to_docx

    jobs = [
        ROOT
        / "参考文档/2、开课管理/03_选修开课/02_选修开课安排/选修开课安排20260904V1/选修开课安排20260904V1.md",
        ROOT
        / "参考文档/2、开课管理/01_开课设置/02_校选课程管理/校选课程管理20260904V5/校选课程管理20260904V5.md",
    ]
    for md in jobs:
        out = md.with_suffix(".docx")
        if out.exists():
            print(f"SKIP exists {out.relative_to(ROOT)}")
        else:
            md_to_docx(md, out)
            print(f"OK PRD {out.relative_to(ROOT)}")

    notes = [
        (
            ROOT
            / "参考文档/2、开课管理/03_选修开课/02_选修开课安排/选修开课安排20260904V1/选修开课安排无→20260904V1变更说明.md",
            {
                "name": "厦大马来分校本科教务系统产品需求文档 — 选修开课安排",
                "path": "开课管理 → 选修开课 → 选修开课安排",
                "versions": "无（开课0807 无本菜单正式稿）→ 20260904V1",
                "from_file": "开课0807/…/校选课程管理20260804V1.md（对比基线）",
                "to_file": "选修开课安排20260904V1.docx",
                "prototype": "page-course-offering-ge / drawer-ge-offering-scope",
                "date": "2026-09-04",
            },
        ),
        (
            ROOT
            / "参考文档/2、开课管理/01_开课设置/02_校选课程管理/校选课程管理20260904V5/校选课程管理20260903V4→20260904V5变更说明.md",
            {
                "name": "厦大马来分校本科教务系统产品需求文档 — 校选课程管理",
                "path": "开课管理 → 开课设置 → 校选课程管理",
                "versions": "开课0807·20260804V1 / 链上V4 → 20260904V5",
                "from_file": "开课0807/…/校选课程管理20260804V1.md",
                "to_file": "校选课程管理20260904V5.docx",
                "prototype": "page-school-elective-courses",
                "date": "2026-09-04",
            },
        ),
    ]
    for md_path, meta in notes:
        out = md_path.with_suffix(".docx")
        if out.exists():
            print(f"SKIP exists {out.relative_to(ROOT)}")
            continue
        rows = parse_overview_rows(md_path.read_text(encoding="utf-8"))
        write_change_note_docx(md_path, out, meta, rows)
        print(f"OK note {out.relative_to(ROOT)} ({len(rows)} rows)")


if __name__ == "__main__":
    main()
