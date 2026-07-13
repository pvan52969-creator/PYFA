#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将「开课管理与排课管理原型评审0710汇总.md」转为排版 Word 文档。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "参考文档" / "开课管理与排课管理原型评审0710汇总.md"
OUT_PATH = ROOT / "参考文档" / "开课管理与排课管理原型评审0710汇总.docx"

FONT_BODY = "宋体"
FONT_HEADING = "黑体"
FONT_CODE = "Courier New"
COLOR_PRIMARY = "1E4E8C"
COLOR_HEADER_BG = "E8F0FA"
COLOR_QUOTE_BG = "F5F7FA"
COLOR_QUOTE_BORDER = "1E4E8C"


def set_run_font(run, name=FONT_BODY, size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "B8C4D0")
        borders.append(el)
    tbl_pr.append(borders)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = FONT_BODY
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(6)


def add_title_block(doc: Document, title: str, subtitle_lines: list[str]):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(title)
    set_run_font(run, FONT_HEADING, 22, bold=True, color=COLOR_PRIMARY)

    for line in subtitle_lines:
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sp.paragraph_format.space_after = Pt(4)
        sr = sp.add_run(line)
        set_run_font(sr, FONT_BODY, 10.5, color="666666")

    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int):
    h = doc.add_heading(text, level=min(level, 3))
    sizes = {1: 16, 2: 14, 3: 12}
    for run in h.runs:
        set_run_font(run, FONT_HEADING, sizes.get(level, 12), bold=True, color=COLOR_PRIMARY)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(8)
    return h


def add_body_paragraph(doc: Document, text: str, bold_parts: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0)
    if bold_parts:
        parts = re.split(r"(\*\*[^*]+\*\*)", text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                set_run_font(run, FONT_BODY, 11, bold=True)
            elif part:
                run = p.add_run(part)
                set_run_font(run, FONT_BODY, 11)
    else:
        run = p.add_run(text)
        set_run_font(run, FONT_BODY, 11, bold=bold_parts)
    return p


def add_bullet(doc: Document, text: str, level: int = 0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.8 + level * 0.6)
    p.paragraph_format.space_after = Pt(3)
    parts = re.split(r"(\*\*[^*]+\*\*)", text.lstrip("- ").strip())
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, FONT_BODY, 11, bold=True)
        else:
            run = p.add_run(part)
            set_run_font(run, FONT_BODY, 11)
    return p


def add_quote_box(doc: Document, lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, COLOR_QUOTE_BG)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8" if edge == "left" else "4")
        el.set(qn("w:color"), COLOR_QUOTE_BORDER if edge == "left" else "D0D7DE")
        borders.append(el)
    tc_pr.append(borders)
    cell.text = ""
    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(line)
        set_run_font(run, FONT_BODY, 10.5, color="444444")
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F6F8FA")
    set_table_borders(table)
    cell.text = ""
    for i, line in enumerate(lines):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.15
        run = para.add_run(line.rstrip())
        set_run_font(run, FONT_CODE, 9.5, color="333333")
    doc.add_paragraph()


def strip_md_inline(text: str) -> str:
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def parse_table_rows(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    headers = [c.strip() for c in lines[0].strip("|").split("|")]
    rows = []
    for line in lines[2:]:
        if not line.strip():
            continue
        rows.append([strip_md_inline(c.strip()) for c in line.strip("|").split("|")])
    return headers, rows


def guess_col_widths(headers: list[str], rows: list[list[str]], total_cm: float = 16.0):
    n = len(headers)
    if n == 2:
        return [4.5, total_cm - 4.5]
    if n == 3:
        return [5.5, 5.0, total_cm - 10.5]
    weights = []
    for i in range(n):
        max_len = len(headers[i])
        for row in rows:
            if i < len(row):
                max_len = max(max_len, len(row[i]))
        weights.append(max(max_len, 4))
    s = sum(weights) or 1
    return [round(total_cm * w / s, 2) for w in weights]


def add_data_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_borders(table)
    widths = guess_col_widths(headers, rows)
    set_col_widths(table, widths)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = strip_md_inline(h)
        set_cell_shading(hdr_cells[i], COLOR_HEADER_BG)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                set_run_font(run, FONT_BODY, 10, bold=True, color=COLOR_PRIMARY)

    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci in range(len(headers)):
            val = row[ci] if ci < len(row) else ""
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                if ci == 0 and len(headers) <= 3:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, FONT_BODY, 10)
                    if val in ("✅", "⚠️", "❌", "🔄"):
                        pass
                    elif val.startswith("P0"):
                        set_run_font(run, FONT_BODY, 10, bold=True, color="C0392B")

    doc.add_paragraph()


def add_numbered_list(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        parts = re.split(r"(\*\*[^*]+\*\*)", item)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                set_run_font(run, FONT_BODY, 11, bold=True)
            else:
                run = p.add_run(part)
                set_run_font(run, FONT_BODY, 11)


def build_from_markdown(md_text: str) -> Document:
    doc = Document()
    configure_document(doc)

    lines = md_text.splitlines()
    i = 0
    title = ""
    quote_buffer: list[str] = []
    code_buffer: list[str] = []
    in_code = False

    # 首行标题
    if lines and lines[0].startswith("# "):
        title = lines[0][2:].strip()
        i = 1

    meta_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("> "):
            meta_lines.append(line[2:].strip())
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            break
        i += 1

    add_title_block(doc, title, meta_lines)

    while i < len(lines):
        line = lines[i]

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_buffer)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            if len(table_lines) >= 2:
                headers, rows = parse_table_rows(table_lines)
                add_data_table(doc, headers, rows)
            continue

        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
            i += 1
            continue

        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1 if re.match(r"^\d+\.", line[3:]) else 2)
            i += 1
            continue

        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
            i += 1
            continue

        if line.startswith("- "):
            add_bullet(doc, line)
            i += 1
            continue

        if re.match(r"^\d+\.\s", line):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i]):
                items.append(re.sub(r"^\d+\.\s", "", lines[i]).strip())
                i += 1
            add_numbered_list(doc, items)
            continue

        if line.startswith("> "):
            quote_buffer.append(line[2:])
            i += 1
            if i >= len(lines) or not lines[i].startswith(">"):
                add_quote_box(doc, quote_buffer)
                quote_buffer = []
            continue

        if line.strip().startswith("*文档生成时间"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(line.strip("*"))
            set_run_font(run, FONT_BODY, 9, color="888888")
            i += 1
            continue

        if line.strip():
            add_body_paragraph(doc, line.strip(), bold_parts=True)
        i += 1

    return doc


def main():
    if not MD_PATH.exists():
        raise SystemExit(f"找不到源文件：{MD_PATH}")
    md_text = MD_PATH.read_text(encoding="utf-8")
    doc = build_from_markdown(md_text)
    doc.save(OUT_PATH)
    print(f"已生成：{OUT_PATH}")


if __name__ == "__main__":
    main()
