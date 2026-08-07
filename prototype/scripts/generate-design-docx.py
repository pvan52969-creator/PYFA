#!/usr/bin/env python3
"""将概要/详细设计 Markdown 生成为版式整洁的 Word（真表格、边框、表头底纹）。

用法：
  python3 scripts/generate-design-docx.py              # 模板 + 全部 *20260807V1.md
  python3 scripts/generate-design-docx.py path/to.md   # 指定文件
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent

FONT_BODY = "宋体"
FONT_HEAD = "黑体"
SIZE_BODY = Pt(10.5)
SIZE_TABLE = Pt(9)
SIZE_H1 = Pt(18)
SIZE_H2 = Pt(14)
SIZE_H3 = Pt(12)


def set_run_font(run, *, name=FONT_BODY, size=SIZE_BODY, bold=False, color=None):
    run.bold = bold
    run.font.size = size
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill_hex: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    # 去掉旧 borders
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "94A3B8")
        borders.append(el)
    tblPr.append(borders)


def set_table_full_width(table):
    """让表格尽量占满版心。"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn("w:tblW"):
            tblPr.remove(child)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)


def clear_cell(cell):
    cell.text = ""


def write_cell_text(cell, text: str, *, bold=False, center=False, header=False):
    clear_cell(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 简单处理行内 **bold** 与 `code`
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text or "")
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=SIZE_TABLE, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=Pt(8.5))
        else:
            run = p.add_run(part)
            set_run_font(run, size=SIZE_TABLE, bold=bold or header)
    if header:
        set_cell_shading(cell, "E2E8F0")


def add_rich_paragraph(doc, text: str, *, style="body"):
    p = doc.add_paragraph()
    if style == "body":
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        size = SIZE_BODY
    elif style == "quote":
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(8)
        size = Pt(9.5)
    elif style == "code":
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        size = Pt(9)
    else:
        size = SIZE_BODY

    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text or "")
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=Pt(max(8, size.pt - 1)))
        else:
            run = p.add_run(part)
            if style == "quote":
                set_run_font(run, size=size, color=RGBColor(0x47, 0x55, 0x69))
            elif style == "code":
                set_run_font(run, name="Consolas", size=size, color=RGBColor(0x1E, 0x29, 0x3B))
            else:
                set_run_font(run, size=size)
    return p


def add_heading_cn(doc, text: str, level: int):
    # 去掉标题里的 markdown 强调
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text).strip()
    h = doc.add_heading(text, level=min(level, 3))
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6)
    size = {1: SIZE_H1, 2: SIZE_H2, 3: SIZE_H3}.get(level, SIZE_H3)
    for run in h.runs:
        set_run_font(run, name=FONT_HEAD, size=size, bold=True, color=RGBColor(0x0F, 0x17, 0x2A))
    return h


def add_grid_table(doc, rows: list[list[str]]):
    if not rows:
        return None
    cols = max(len(r) for r in rows)
    # 补齐列数
    norm = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(norm), cols=cols)
    set_table_borders(table)
    set_table_full_width(table)
    for ri, row in enumerate(norm):
        is_header = ri == 0
        for ci, val in enumerate(row):
            write_cell_text(
                table.rows[ri].cells[ci],
                val.strip(),
                header=is_header,
                center=is_header,
                bold=is_header,
            )
    # 表后空一行
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def parse_md_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith("|"):
            continue
        # 分隔行 |---|---|
        if re.match(r"^\|[\s:\-|]+\|$", line):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
    return rows


def setup_doc() -> Document:
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
    # 默认正文字体
    style = doc.styles["Normal"]
    style.font.name = FONT_BODY
    style.font.size = SIZE_BODY
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    return doc


def add_picture(doc, img_path: Path, *, max_width_cm: float = 17.0):
    if not img_path.exists():
        add_rich_paragraph(doc, f"［缺图：{img_path.name}］", style="quote")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    # 按页宽缩放
    from PIL import Image as PILImage

    with PILImage.open(img_path) as im:
        w_px, h_px = im.size
    # 目标宽度
    target_w = Cm(max_width_cm)
    # python-docx 用英寸；保持比例
    aspect = h_px / max(w_px, 1)
    target_h = Cm(max_width_cm * aspect)
    # 高度过大则再压
    max_h = Cm(20)
    if target_h > max_h:
        scale = max_h / target_h
        target_w = Cm(max_width_cm * scale)
        target_h = max_h
    run.add_picture(str(img_path), width=target_w, height=target_h)


def md_to_docx(md_path: Path, out_path: Path | None = None) -> Path:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    doc = setup_doc()
    out_path = out_path or md_path.with_suffix(".docx")
    base_dir = md_path.parent

    i = 0
    n = len(lines)
    title_done = False
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            i += 1
            continue

        # 图片 ![alt](path)
        m_img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if m_img:
            rel = m_img.group(2).strip().strip('"').strip("'")
            img_path = (base_dir / rel).resolve()
            add_picture(doc, img_path)
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            if i < n:
                i += 1
            for code_line in buf:
                add_rich_paragraph(doc, code_line if code_line else " ", style="code")
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if stripped.startswith("|"):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append(lines[i])
                i += 1
            rows = parse_md_table(buf)
            if rows:
                add_grid_table(doc, rows)
            continue

        m = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            if level == 1 and not title_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                run = p.add_run(re.sub(r"\*\*([^*]+)\*\*", r"\1", heading_text))
                set_run_font(run, name=FONT_HEAD, size=Pt(20), bold=True)
                title_done = True
            else:
                add_heading_cn(doc, heading_text, level)
            i += 1
            continue

        if stripped.startswith(">"):
            quote = re.sub(r"^>\s?", "", stripped)
            add_rich_paragraph(doc, quote, style="quote")
            i += 1
            continue

        m_ul = re.match(r"^[-*]\s+(.+)$", stripped)
        m_ol = re.match(r"^(\d+)[.)、]\s+(.+)$", stripped)
        if m_ul or m_ol:
            content = m_ul.group(1) if m_ul else m_ol.group(2)
            prefix = "· " if m_ul else f"{m_ol.group(1)}. "
            add_rich_paragraph(doc, prefix + content, style="body")
            i += 1
            continue

        add_rich_paragraph(doc, stripped, style="body")
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def collect_targets(extra: list[str]) -> list[Path]:
    if extra:
        return [Path(p).resolve() for p in extra]
    targets: list[Path] = []
    tpl = ROOT / "参考文档" / "0、模板"
    for name in ("概要设计文档模板.md", "详细设计文档模板.md"):
        p = tpl / name
        if p.exists():
            targets.append(p)
    base = ROOT / "参考文档" / "2、开课管理"
    targets.extend(sorted(base.rglob("*20260807V1.md")))
    return targets


def main():
    targets = collect_targets(sys.argv[1:])
    if not targets:
        raise SystemExit("未找到待转换的 Markdown")
    for md in targets:
        if not md.exists():
            print(f"SKIP missing {md}")
            continue
        out = md_to_docx(md)
        print(f"OK {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
