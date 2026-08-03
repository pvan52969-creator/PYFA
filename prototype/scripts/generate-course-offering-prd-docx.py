#!/usr/bin/env python3
"""Generate 开课管理 PRD docx matching XMUM product requirement template."""

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path

from prd_course_offering_build import build_document

OUT = (
    Path(__file__).resolve().parent.parent
    / "参考文档"
    / "2、开课管理"
    / "开课管理系统产品需求文档.docx"
)
SCRIPT_DIR = Path(__file__).resolve().parent

LANDSCAPE_SCALE = 1.55
FIELD_TABLE_COL_WIDTHS = [1.2, 3.1, 3.4, 1.9, 1.1, 3.9, 4.4, 1.9, 2.5]

FIELD_HEADERS = [
    "序号",
    "字段中文名称",
    "英文名称",
    "字段类型",
    "必填",
    "校验规则\n（没有明确规则时先限制字符即可，放宽）",
    "备注（如下拉框取值来源、说明等）",
    "是否代码集取值",
    "数据格式样例\n（无特定格式留空）",
]


def scale_widths(widths):
    if isinstance(widths, tuple):
        return tuple(round(w * LANDSCAPE_SCALE, 1) for w in widths)
    return [round(w * LANDSCAPE_SCALE, 1) for w in widths]


def set_document_landscape(doc):
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def add_para(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "黑体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    return h


def add_kv_table(doc, rows, col_widths=(5, 18.5)):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_col_widths(table, col_widths)
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = k
        c1.text = v
        for cell in (c0, c1):
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        c0.paragraphs[0].runs[0].bold = True
    doc.add_paragraph()
    return table


def add_grid_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_borders(table)
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if col_widths:
        set_col_widths(table, col_widths)
    doc.add_paragraph()
    return table


def add_field_table(doc, title, rows):
    add_para(doc, title, bold=True)
    add_grid_table(doc, FIELD_HEADERS, rows, col_widths=FIELD_TABLE_COL_WIDTHS)


def add_function_item(doc, num, cn_name, en_name, desc, interaction, notes, drill_down, source=""):
    title = f"{num}、功能按钮——{cn_name}（英文名称：{en_name}）"
    if source:
        title += f"【来源：{source}】"
    add_para(doc, title, bold=True)
    add_para(doc, "a. 功能说明（描述、业务的事件交互、备注信息等）：")
    add_para(doc, f"· 描述：{desc}")
    add_para(doc, f"· 业务的事件交互：{interaction}")
    add_para(doc, f"· 备注信息（校验规则补充、其他说明等）：{notes}")
    add_para(doc, f"b. 下钻页面说明（备注说明）：{drill_down}")
    doc.add_paragraph()


def add_menu_block(
    doc,
    section_no,
    title,
    status,
    intro,
    list_fields,
    search_fields,
    flow_rel,
    flow_pre,
    flow_out,
    biz_flow,
    proto_link="",
    field_tables=None,
    functions=None,
):
    add_heading(doc, f"{section_no}{title}（需求确认状态：{status}）", level=4)
    add_para(doc, "菜单介绍", bold=True)
    add_kv_table(doc, [("菜单介绍", ""), ("菜单内容简介", intro)])
    add_para(doc, "页面展示字段信息", bold=True)
    add_kv_table(doc, [("页面展示字段信息", list_fields)])
    add_para(doc, "支持查询检索的字段信息", bold=True)
    add_kv_table(doc, [("支持查询检索的字段信息", search_fields)])
    add_para(doc, "数据前后流转关系（说明、前置条件、下游输出等）", bold=True)
    add_kv_table(
        doc,
        [
            ("1）关系说明", flow_rel),
            ("2）前置条件", flow_pre),
            ("3）下游输出", flow_out),
        ],
    )
    add_para(doc, "5、业务流关系（操作流程）", bold=True)
    add_kv_table(doc, [("5、业务流关系（操作流程）", biz_flow)])
    if proto_link:
        add_para(doc, "6、原型参考链接", bold=True)
        add_kv_table(doc, [("原型参考链接", proto_link)])
    if field_tables:
        add_para(doc, "新增—字段信息表", bold=True)
        for ft_title, ft_rows in field_tables:
            add_field_table(doc, ft_title, ft_rows)
    if functions:
        add_para(doc, "（3）菜单功能清单", bold=True)
        for fn in functions:
            add_function_item(doc, *fn)


def build():
    import sys

    if str(SCRIPT_DIR) not in sys.path:
        sys.path.insert(0, str(SCRIPT_DIR))
    build_document(
        add_heading,
        add_para,
        add_kv_table,
        add_grid_table,
        add_field_table,
        add_menu_block,
        scale_widths,
        OUT,
        Document,
        set_document_landscape,
    )


if __name__ == "__main__":
    build()
    print(f"Wrote: {OUT}")
