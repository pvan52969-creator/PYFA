#!/usr/bin/env python3
"""从 参考文档/参考文件/计算机教室软件.xlsx 生成 classroom-software-data.js。"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "参考文档/参考文件/计算机教室软件.xlsx"
OUT = ROOT / "classroom-software-data.js"


def norm_venue(raw: str) -> str:
    parts = str(raw).strip().split("-", 1)
    return f"{parts[0]}#{parts[1]}" if len(parts) == 2 else str(raw).strip()


def is_mark(value) -> bool:
    if value is None:
        return False
    return str(value).strip().upper() in ("X", "✓", "√", "Y", "YES", "1")


def extract() -> dict[str, list[str]]:
    wb = load_workbook(SRC, data_only=True)
    ws = wb.active
    venues: list[tuple[int, str]] = []
    for col in range(1, ws.max_column + 1):
        header = ws.cell(4, col).value
        if header and str(header).strip() not in ("Venue", ""):
            venues.append((col, norm_venue(header)))

    start = None
    for row in range(1, ws.max_row + 1):
        label = ws.cell(row, 2).value
        if label and str(label).strip().lower() == "softwares":
            start = row + 1
            break
    if start is None:
        raise RuntimeError("未找到 Softwares 起始行")

    result: dict[str, list[str]] = {}
    for col, code in venues:
        names: list[str] = []
        seen: set[str] = set()
        for row in range(start, ws.max_row + 1):
            name = ws.cell(row, 2).value
            if not name:
                continue
            name = str(name).strip()
            if not name or not is_mark(ws.cell(row, col).value):
                continue
            key = name.lower()
            if key in seen:
                continue
            seen.add(key)
            names.append(name)
        result[code] = names

    # xlsx 列 A1-G11 对应教室名单 A1#111
    if "A1#G11" in result and "A1#111" not in result:
        result["A1#111"] = result["A1#G11"]
    return result


def main() -> None:
    data = extract()
    body = json.dumps(data, ensure_ascii=False, indent=2)
    OUT.write_text(
        "/** 计算机教室软件清单（来源：参考文档/参考文件/计算机教室软件.xlsx；绿色 X = 已安装） */\n"
        f"const CLASSROOM_SOFTWARE_BY_ROOM = {body};\n",
        encoding="utf-8",
    )
    print(f"Wrote {OUT} ({len(data)} rooms)")


if __name__ == "__main__":
    main()
