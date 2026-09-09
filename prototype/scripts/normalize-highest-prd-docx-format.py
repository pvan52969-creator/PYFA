#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""将开课管理现行最高版 PRD docx 统一为：
- A4 横向打开
- 表格内字体一律小五（9pt）宋体
- 表格表头与内容一律水平+垂直居中
- 相对开课0807有差异的内容用黄色底色（单元格底纹）标注

差异判定（按优先级）：
1. 同文件夹变更说明总览中的对象ID（FD/F/BR/P…）命中表格行 → 标黄
2. 变更说明含「完整重写」或无0807基线（首版/课程班系列等）→ 字段/功能/规则/界面类表数据行整行标黄
3. 正文/代码块含上述对象ID或变更说明可读名关键词 → 段落字符黄底（highlight）

用法：
  python3 scripts/normalize-highest-prd-docx-format.py
  python3 scripts/normalize-highest-prd-docx-format.py --dry-run
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
OPEN_ROOT = ROOT / "参考文档" / "2、开课管理"

ASCII = "Times New Roman"
EAST = "宋体"
SIZE_TABLE = Pt(9)  # 小五
SIZE_BODY = Pt(10.5)  # 五号
SIZE_H1 = Pt(16)
SIZE_H2 = Pt(14)
SIZE_H3 = Pt(12)
YELLOW = "FFFF00"
HEADER_FILL = "D9E2F3"

VER_RE = re.compile(r"^(.+?)(\d{8})(V\d+)$")
OID_RE = re.compile(r"\b((?:FD|F|BR|UC|P)\d+(?:-[A-Z]\d+)*)\b")

# 菜单名 → 开课0807 基线 md（无则视为相对0807整份新增）
BASELINE_0807 = {
    "开课时间设置": OPEN_ROOT
    / "开课0807/开课设置/开课时间设置/开课时间设置20260804V2/开课时间设置20260804V2.md",
    "校选课程管理": OPEN_ROOT
    / "开课0807/开课设置/校选课程管理/校选课程管理20260804V1/校选课程管理20260804V1.md",
    "特殊课程设置": OPEN_ROOT
    / "开课0807/开课设置/特殊课程设置/特殊课程设置20260804V2/特殊课程设置20260804V2.md",
    "开课计划": OPEN_ROOT
    / "开课0807/专业开课/开课计划/开课计划20260804V4/开课计划20260804V4.md",
    "开课安排": OPEN_ROOT
    / "开课0807/专业开课/开课安排/开课安排20260804V4/开课安排20260804V4.md",
    "开课名单": OPEN_ROOT
    / "开课0807/专业开课/开课名单/开课名单20260804V2/开课名单20260804V2.md",
}


def latest_prd_folders() -> list[tuple[str, Path, Path]]:
    """返回 [(菜单名, 版本文件夹, prd.md), ...]"""
    latest: dict[str, tuple[int, str, Path, str]] = {}
    for p in OPEN_ROOT.rglob("*"):
        if not p.is_dir():
            continue
        m = VER_RE.match(p.name)
        if not m:
            continue
        rel = str(p.relative_to(OPEN_ROOT))
        if rel.startswith("开课0807") or rel.startswith("00_"):
            continue
        name, date, vn = m.group(1), m.group(2), int(m.group(3)[1:])
        key = str(p.parent)
        cur = latest.get(key)
        if not cur or (vn, date) > (cur[0], cur[1]):
            latest[key] = (vn, date, p, name)
    out = []
    for vn, date, folder, name in sorted(latest.values(), key=lambda t: str(t[2])):
        md = folder / f"{folder.name}.md"
        if md.exists():
            out.append((name, folder, md))
    return out


def set_run_font(run, size=SIZE_BODY, bold=False, highlight=False):
    run.font.name = ASCII
    run._element.rPr.rFonts.set(qn("w:ascii"), ASCII)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ASCII)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EAST)
    run.font.size = size
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    for old in rPr.findall(qn("w:highlight")):
        rPr.remove(old)
    if highlight:
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        rPr.append(hl)


def set_cell_shading(cell, fill: str | None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    if not fill:
        return
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    tcPr.append(shd)


def set_document_landscape(doc: Document):
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21)
        sec.left_margin = Cm(1.8)
        sec.right_margin = Cm(1.8)
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def parse_blocks(md: str):
    lines = md.splitlines()
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("h", level, line.lstrip("#").strip()))
            i += 1
            continue
        if line.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", "\n".join(buf)))
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[\s\-:|]+\|$", lines[i + 1].strip()
        ):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if re.match(r"^\|[\s\-:|]+\|$", lines[i].strip()):
                    i += 1
                    continue
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            blocks.append(("table", rows))
            continue
        if line.strip().startswith(">"):
            blocks.append(("quote", line.strip().lstrip("> ").strip()))
            i += 1
            continue
        if line.strip():
            blocks.append(("p", line.strip()))
        i += 1
    return blocks


def find_change_note(folder: Path) -> Path | None:
    cands = sorted(folder.glob("*变更说明.md"))
    return cands[0] if cands else None


def load_diff_meta(folder: Path, menu_name: str) -> dict:
    """从变更说明解析差异对象；判断是否整份相对0807重写/新增。"""
    note = find_change_note(folder)
    ids: set[str] = set()
    keywords: set[str] = set()
    full_rewrite = False
    has_baseline = menu_name in BASELINE_0807 and BASELINE_0807[menu_name].exists()

    if note and note.exists():
        text = note.read_text(encoding="utf-8")
        if "完整重写" in text or "无→" in note.name:
            full_rewrite = True
        # 总览表
        m = re.search(
            r"## 2 .*?\n\n(\|.+\n\|[-| :]+\n(?:\|.+\n)+)", text, re.S
        )
        if m:
            lines = [ln for ln in m.group(1).splitlines() if ln.startswith("|")]
            for ln in lines[2:]:
                cells = [c.strip() for c in ln.strip("|").split("|")]
                if len(cells) < 7:
                    continue
                oid = cells[3]
                where = cells[4]
                content = cells[6]
                if oid and oid not in ("—", "-"):
                    ids.add(oid)
                    ids.update(OID_RE.findall(oid))
                for piece in (where, content):
                    ids.update(OID_RE.findall(piece))
                    # 短可读名
                    if where and where not in ("—", "文档结构", "文档结构 / PRD 模板"):
                        keywords.add(where.strip())
                if "完整重写" in content or "首版" in content:
                    full_rewrite = True

    # 无0807基线：整份视为相对0807新增
    if not has_baseline:
        full_rewrite = True

    return {
        "ids": ids,
        "keywords": keywords,
        "full_rewrite": full_rewrite,
        "has_baseline": has_baseline,
        "note": note,
    }


def table_kind(header_row: list[str]) -> str:
    h = "|".join(header_row)
    if "字段ID" in h:
        return "field"
    if "功能ID" in h:
        return "func"
    if "规则ID" in h:
        return "rule"
    if "界面ID" in h or "页签/区块ID" in h:
        return "ui"
    if "编号" in h and ("事项" in h or "未确认" in h):
        return "uc"
    return "other"


def row_should_yellow(row: list[str], kind: str, meta: dict) -> bool:
    joined = "|".join(row)
    # 对象ID命中
    for oid in meta["ids"]:
        if oid and oid in joined:
            return True
    # 关键词（可读名）命中整格相等或包含
    for kw in meta["keywords"]:
        if len(kw) >= 2 and kw in joined:
            return True
    # 整份重写/无基线：业务表数据行全标
    if meta["full_rewrite"] and kind in ("field", "func", "rule", "ui"):
        return True
    return False


def text_should_yellow(text: str, meta: dict) -> bool:
    for oid in meta["ids"]:
        if oid and oid in text:
            return True
    for kw in meta["keywords"]:
        if len(kw) >= 4 and kw in text:
            return True
    # 界面树在完整重写时标黄
    if meta["full_rewrite"] and (
        text.startswith("P01") or "├──" in text or "└──" in text
    ):
        return True
    return False


def add_para(doc, text, *, size=SIZE_BODY, bold=False, highlight=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, highlight=highlight)
    return p


def add_table(doc, rows, meta: dict):
    if not rows:
        return
    cols = max(len(r) for r in rows)

    def trim(r):
        if cols <= 12:
            return r + [""] * (cols - len(r))
        # 过宽表：保留前若干列 + 末列备注
        if len(r) <= 12:
            return r + [""] * (12 - len(r))
        return r[:11] + [r[-1]]

    use_rows = rows
    use_cols = cols
    if cols > 12:
        use_rows = [trim(r) for r in rows]
        use_cols = 12

    kind = table_kind(use_rows[0]) if use_rows else "other"
    table = doc.add_table(rows=len(use_rows), cols=use_cols)
    table.style = "Table Grid"
    set_table_borders(table)

    for ri, row in enumerate(use_rows):
        is_header = ri == 0
        yellow = (not is_header) and row_should_yellow(row, kind, meta)
        for ci in range(use_cols):
            cell = table.rows[ri].cells[ci]
            text = row[ci] if ci < len(row) else ""
            # 清空单元格后只保留一个带字号的 run，避免残留空 run → 字号 None
            # 表头与内容：水平+垂直居中（后续默认）
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for pi, p in enumerate(cell.paragraphs):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not p.runs:
                    run = p.add_run("" if pi > 0 else text)
                    set_run_font(run, size=SIZE_TABLE, bold=is_header, highlight=False)
                else:
                    for ri_run, run in enumerate(p.runs):
                        if pi == 0 and ri_run == 0:
                            run.text = text
                        else:
                            run.text = ""
                        set_run_font(run, size=SIZE_TABLE, bold=is_header, highlight=False)
            if is_header:
                set_cell_shading(cell, HEADER_FILL)
            elif yellow:
                set_cell_shading(cell, YELLOW)
            else:
                set_cell_shading(cell, None)
    doc.add_paragraph()


def build_docx(md_text: str, out: Path, title: str, meta: dict, ver_label: str):
    doc = Document()
    set_document_landscape(doc)

    add_para(doc, title, size=SIZE_H1, bold=True)
    legend = (
        f"文档版本：{ver_label}　　模板：PRD V3.1　　纸张：A4 横向　　"
        f"表格：小五·水平垂直居中　　黄色底色=相对开课0807有差异"
    )
    add_para(doc, legend, size=SIZE_TABLE, highlight=False)

    for b in parse_blocks(md_text):
        if b[0] == "h":
            _, level, title_h = b
            if title_h.startswith("厦大马来"):
                continue
            size = SIZE_H1 if level <= 1 else (SIZE_H2 if level == 2 else SIZE_H3)
            add_para(
                doc,
                title_h,
                size=size,
                bold=True,
                highlight=text_should_yellow(title_h, meta),
            )
        elif b[0] == "quote":
            add_para(
                doc,
                b[1],
                size=SIZE_TABLE,
                highlight=text_should_yellow(b[1], meta),
            )
        elif b[0] == "code":
            # 界面树：完整重写或含差异ID时标黄
            hl = meta["full_rewrite"] or text_should_yellow(b[1], meta)
            add_para(doc, b[1], size=SIZE_TABLE, highlight=hl)
        elif b[0] == "p":
            add_para(
                doc,
                b[1],
                size=SIZE_BODY,
                highlight=text_should_yellow(b[1], meta),
            )
        elif b[0] == "table":
            add_table(doc, b[1], meta)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    items = latest_prd_folders()
    print(f"最高版菜单数：{len(items)}")
    for menu_name, folder, md in items:
        meta = load_diff_meta(folder, menu_name)
        docx = folder / f"{folder.name}.docx"
        ver_label = folder.name.replace(menu_name, "", 1) if folder.name.startswith(menu_name) else folder.name
        # folder name is like 校选课程管理20260904V5
        ver_label = folder.name[len(menu_name) :] if folder.name.startswith(menu_name) else folder.name
        mode = (
            "整份标黄(无0807/完整重写)"
            if meta["full_rewrite"]
            else f"按对象ID标黄×{len(meta['ids'])}"
        )
        print(
            f"- {menu_name} {ver_label}: baseline={'Y' if meta['has_baseline'] else 'N'} "
            f"{mode} note={meta['note'].name if meta['note'] else '—'}"
        )
        if args.dry_run:
            continue
        md_text = md.read_text(encoding="utf-8")
        title = f"厦大马来分校本科教务系统产品需求文档 — {menu_name}"
        build_docx(md_text, docx, title, meta, ver_label)
        print(f"  → wrote {docx.relative_to(ROOT)}")

    if args.dry_run:
        print("(dry-run，未写文件)")


if __name__ == "__main__":
    main()
