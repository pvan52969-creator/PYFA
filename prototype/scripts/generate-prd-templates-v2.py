#!/usr/bin/env python3
"""生成三份需求模板 V2（不覆盖 参考文档/0、模板/ 下原 V1 文件）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "参考文档" / "0、模板"

CN_FONT = "宋体"
CN_HEADING = "黑体"
EN_FONT = "Times New Roman"
HEADER_FILL = "D9E2F3"
HINT_COLOR = RGBColor(0x80, 0x80, 0x80)
HEADING_COLOR = RGBColor(0x1F, 0x29, 0x37)
BORDER_COLOR = "000000"

FIELD_HEADERS = [
    "序号",
    "所在界面/区块",
    "字段中文名称",
    "英文名称",
    "字段类型",
    "必填",
    "校验规则\n（没有明确规则时先限制字符即可，放宽）",
    "备注（如下拉框取值来源、说明等）",
    "是否代码集取值",
    "数据格式样例\n（无特定格式留空）",
]
FIELD_WIDTHS = [1.0, 3.6, 2.6, 2.8, 1.6, 1.0, 3.2, 3.4, 1.4, 2.0]


def set_run_font(run, *, cn=CN_FONT, en=EN_FONT, size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if color is not None:
        run.font.color.rgb = color


def apply_line_spacing(p):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.line_spacing = 1.5


def add_para(doc, text="", *, bold=False, size=11, color=None, align=None, space_after=8):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    apply_line_spacing(p)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_hint(doc, text):
    return add_para(doc, text, size=10, color=HINT_COLOR, space_after=8)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    sizes = {1: 16, 2: 14, 3: 12}
    for run in h.runs:
        set_run_font(run, cn=CN_HEADING, size=sizes.get(level, 12), bold=True, color=HEADING_COLOR)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(8)
    apply_line_spacing(h)
    return h


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    tr_pr.append(flag)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    if tbl_pr is None:
        tbl_pr = OxmlElement("w:tblPr")
        tbl.insert(0, tbl_pr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER_COLOR)
        borders.append(el)
    tbl_pr.append(borders)
    set_table_cell_margins(tbl_pr)


def set_table_cell_margins(tbl_pr, top=80, bottom=80, left=100, right=100):
    """单元格内边距，单位 dxa（20 dxa = 1 pt）。"""
    for old in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def set_col_widths(table, widths_cm):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def fill_cell(cell, text, *, bold=False, size=9, align=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    apply_line_spacing(p)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)


def add_kv_table(doc, rows, col_widths=(4.2, 21.4)):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_borders(table)
    set_col_widths(table, col_widths)
    for i, (k, v) in enumerate(rows):
        fill_cell(table.rows[i].cells[0], k, bold=True, size=10.5)
        set_cell_shading(table.rows[i].cells[0], "F2F2F2")
        fill_cell(table.rows[i].cells[1], v, size=10.5)
    doc.add_paragraph()
    return table


def add_grid_table(doc, headers, rows, col_widths, *, header_fill=HEADER_FILL, empty_hint=False):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        fill_cell(cell, h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, header_fill)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            color = HINT_COLOR if empty_hint and (val in ("", "需补充")) else None
            fill_cell(
                table.rows[ri + 1].cells[ci],
                val,
                size=9,
                color=color,
                align=WD_ALIGN_PARAGRAPH.CENTER if ci in (0, 5, 8) else None,
            )
    set_col_widths(table, col_widths)
    doc.add_paragraph()
    return table


def add_caption(doc, text):
    p = add_para(doc, text, size=10, bold=True, space_after=4)
    return p


def add_page_field(paragraph, kind="PAGE"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {kind} "
    run2._r.append(instr)
    run3 = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run3._r.append(end)
    for r in (run, run2, run3):
        set_run_font(r, size=9, color=HINT_COLOR)


def setup_section(doc, header_text):
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width = Cm(29.7)
        sec.page_height = Cm(21.0)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.0)
        sec.header_distance = Cm(0.8)
        sec.footer_distance = Cm(0.6)

        header = sec.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = hp.add_run(header_text)
        set_run_font(run, size=9, color=HINT_COLOR)

        footer = sec.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = fp.add_run("第 ")
        set_run_font(r1, size=9, color=HINT_COLOR)
        add_page_field(fp, "PAGE")
        r2 = fp.add_run(" 页 / 共 ")
        set_run_font(r2, size=9, color=HINT_COLOR)
        add_page_field(fp, "NUMPAGES")
        r3 = fp.add_run(" 页")
        set_run_font(r3, size=9, color=HINT_COLOR)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(11)
    if normal._element.rPr is None:
        normal._element.get_or_add_rPr()
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.line_spacing = 1.5
    for i, sz in ((1, 16), (2, 14), (3, 12)):
        style = doc.styles[f"Heading {i}"]
        style.font.name = EN_FONT
        style.font.size = Pt(sz)
        style.font.bold = True
        style.font.color.rgb = HEADING_COLOR
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEADING)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.line_spacing = 1.5


def add_title_block(doc, title, subtitle):
    p = add_para(doc, title, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    for run in p.runs:
        set_run_font(run, cn=CN_HEADING, size=18, bold=True, color=HEADING_COLOR)
    add_para(doc, subtitle, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_hint(
        doc,
        "相对原模板的修订版，不覆盖原文件。打开后请更新页码域：全选 → F9。"
        "条款口径见 1.3：用户已讲清的标（已明确），未聊过或由 AI 补写/遗漏的标（未确认）。禁止把未讨论内容标成已明确。",
    )


def add_function_item(doc, num, cn_name, en_name, desc, interaction, notes, drill, *, filled=False):
    add_para(
        doc,
        f"{num}、功能按钮——{cn_name}（英文名称：{en_name}）",
        bold=True,
        size=11,
    )
    add_para(doc, "a. 功能说明")
    color = None if filled else HINT_COLOR
    add_para(doc, f"· 描述：{desc}", size=11, color=color)
    add_para(doc, f"· 业务的事件交互：{interaction}", size=11, color=color)
    add_para(doc, f"· 备注信息（校验规则补充、其他说明等）：{notes}", size=11, color=color)
    add_para(doc, f"b. 下钻页面说明：{drill}", size=11, color=color)
    doc.add_paragraph()


def empty_field_rows(n=8):
    return [[str(i), "需补充", "需补充", "需补充", "需补充", "", "需补充", "需补充", "", ""] for i in range(1, n + 1)]


def add_prd_common_front(doc, info_rows):
    add_heading(doc, "0. 文档信息", 1)
    add_kv_table(doc, info_rows)

    add_heading(doc, "1. 文档概述", 1)
    add_heading(doc, "1.1 文档目的", 2)
    add_para(
        doc,
        "本文档为厦大马来分校本科教务系统提供标准化需求输入，供评审、原型对照和后续开发使用。"
        "一份文档对应一个二级菜单（或同等粒度的独立功能）。",
    )
    add_heading(doc, "1.2 开发背景", 2)
    add_kv_table(
        doc,
        [
            ("开发模式", "边分析边迭代，分模块生成可交互原型"),
            ("目标用户", "学生、教师、各级管理人员（按本菜单实际角色填写）"),
            ("覆盖范围", "本科生（如本菜单另有范围，在此改写；非用户指定的范围标（未确认））"),
        ],
    )
    add_heading(doc, "1.3 填写说明", 2)
    add_para(
        doc,
        "· 条款口径：用户指定且已讲清的规则或内容，在该项后标（已明确）；"
        "没有明确、没有聊过、由 AI 自行补写或可能遗漏的点，标（未确认）。"
        "不得把 AI 发挥的内容标成已明确。",
    )
    add_para(
        doc,
        "· 文档信息中的「确认状态」表示整份文档是否已经过确认人核过（待确认 / 已确认），"
        "与条款上的（已明确）/（未确认）不是同一件事。",
    )
    add_para(doc, "· 标题只用到三级；表须有题注；字段表按「所在界面/区块」分行，一张表可覆盖筛选、列表、弹窗、抽屉。")
    add_para(doc, "· 英文名称与界面原文保持一致；界面为外文时保留原文并加中文说明。")
    add_para(doc, "· 升版时不得覆盖旧版文件夹；须同时产出变更说明（见《需求调整变更说明模板（简版）20260821V2》）。")
    add_para(doc, "· 无实质需求变更时不要升版。跳号升版时，变更说明对照上一有效版本。")


def add_prd_module_skeleton(doc, *, sample=False):
    add_heading(doc, "2. 模块需求（按一个二级菜单填写）", 1)

    add_heading(doc, "2.1 菜单路径与范围", 2)
    if sample:
        add_kv_table(
            doc,
            [
                ("一级菜单", "基础数据（英文名称：Basic Info）"),
                ("二级菜单", "课程信息管理（英文名称：Course Info）"),
                ("三级菜单", "无"),
                ("范围内", "课程基础信息、课程学习成果（CLO）、学生学时（SLT）的维护与查询"),
            ],
        )
    else:
        add_kv_table(
            doc,
            [
                ("一级菜单", "需补充（英文名称：需提供）"),
                ("二级菜单", "需补充（英文名称：需提供）"),
                ("三级菜单", "无 / 需补充"),
                ("范围内", "需补充：本菜单包含的能力"),
            ],
        )

    add_heading(doc, "2.2 建设目标与非目标", 2)
    if sample:
        add_para(
            doc,
            "目标：维护课程主数据，作为开课、排课、选课与学业评估的底层课程数据源。",
        )
        add_para(
            doc,
            "非目标：不在本菜单完成开课安排、排课结果或选课名单；不替代代码集管理本身。",
        )
    else:
        add_para(doc, "目标：需补充（本菜单要解决什么问题、达成什么结果）。", color=HINT_COLOR)
        add_para(doc, "非目标：需补充（明确不做、或由其他菜单承担的能力；如无则填「无」）。", color=HINT_COLOR)

    add_heading(doc, "2.3 角色与权限", 2)
    add_hint(doc, "只写用户已明确的可见/可操作范围；未聊过的权限标（未确认）。")
    if sample:
        role_rows = [
            ("教务", "维护课程库、导入导出、审核相关操作（以实际权限配置为准）", "未确认"),
            ("学院", "维护本院开课课程信息（以实际权限配置为准）", "未确认"),
            ("教师", "本菜单是否可编辑课程（未确认）", "未确认"),
        ]
    else:
        role_rows = [
            ("需补充角色", "需补充：何时用、可看/可操作哪些", "未确认"),
            ("需补充角色", "需补充", "未确认"),
        ]
    add_grid_table(
        doc,
        ["角色", "场景与能力", "口径"],
        role_rows,
        [3.5, 18.5, 3.5],
        empty_hint=not sample,
    )

    add_heading(doc, "2.4 菜单介绍", 2)
    if sample:
        add_para(
            doc,
            "用于管理课程基础信息、课程学习成果、学生学时及变更记录。"
            "数据依赖开课单位与师资等基础数据；支持分步新增/编辑、导入导出等。"
            "是开课、排课、学业评估的课程数据源。",
        )
    else:
        add_para(doc, "菜单内容简介：需补充。", color=HINT_COLOR)

    add_heading(doc, "2.5 界面说明", 2)
    add_para(doc, "筛选区", bold=True)
    if sample:
        add_para(doc, "课程编号（Course Code）、课程名字（Course Name）、开课单位（Offering）、课程分类（Course Classification）；查询 / 重置。")
    else:
        add_para(doc, "需补充：筛选项、查询/重置。", color=HINT_COLOR)
    add_para(doc, "列表列", bold=True)
    if sample:
        add_para(
            doc,
            "序号（No.）、课程编号（Course Code）、课程名字（Course Name）、开课单位（Offering）、"
            "课程分类（Course Classification）、学分（Credit）、操作（Actions）。",
        )
    else:
        add_para(doc, "需补充：列表列（含操作列）。步骤/Tab 差异列请分开写。", color=HINT_COLOR)
    add_para(doc, "关联界面（弹窗 / 抽屉 / 下钻）", bold=True)
    if sample:
        add_para(doc, "新增/编辑下钻页，三个 Tab：General Information（基础信息）、Course Learning Outcome（CLO）（课程学习成果）、Student Learning Time（SLT）（学生学时）。")
    else:
        add_para(doc, "需补充：弹出页内容、功能按钮；无则写「无下钻页面」。", color=HINT_COLOR)

    add_heading(doc, "2.6 数据前后流转", 2)
    if sample:
        add_kv_table(
            doc,
            [
                (
                    "1）关系说明",
                    "用户在列表发起增删改查、导入复制；新增/编辑提交的课程数据经校验后写入课程主表、CLO、学时明细与变更记录；详情按 Tab 读取；删除前校验下游引用；导出由后端汇总文件返回。",
                ),
                ("2）前置条件", "已维护学期、院系、专业、教师、课程库等关键信息；新增课程须绑定开课单位、课程创建人等。"),
                ("3）下游输出", "作为选课申请、开课、成绩、排课、课程变更、培养方案等模块的课程基础数据源。"),
            ],
        )
    else:
        add_kv_table(
            doc,
            [
                ("1）关系说明", "需补充（无则填无）"),
                ("2）前置条件", "需补充：数据从哪里获取（无则填无）"),
                ("3）下游输出", "需补充：输出到哪些模块（无则填无）"),
            ],
        )

    add_heading(doc, "2.7 业务流（操作流程）", 2)
    if sample:
        add_para(
            doc,
            "进入【Course Info】→ 查询列表 → 新增/编辑（依次完善基础信息、CLO、学生学时）→ 保存入库；"
            "删除前校验下游是否引用。",
        )
    else:
        add_para(doc, "需补充，例如：先维护学年学期，再维护本菜单。", color=HINT_COLOR)

    add_heading(doc, "2.8 原型参考", 2)
    if sample:
        add_kv_table(doc, [("原型页面", "需补充页面 ID（示例写法：prototype/index.html → page-xxx）")])
    else:
        add_kv_table(doc, [("原型页面", "需补充：prototype/index.html → page-xxx；无原型则写无")])


def add_prd_fields_and_functions(doc, *, sample=False):
    add_heading(doc, "2.9 字段信息表", 2)
    add_caption(doc, "表 2-1 字段信息表（按所在界面/区块列全）")
    add_hint(doc, "筛选区、列表、各弹窗与抽屉分区均写入本表；用「所在界面/区块」区分，不必为每个弹窗另起一套表头。")
    if sample:
        rows = [
            ["1", "页面·筛选区", "课程编号", "Course Code", "输入框", "否", "—", "模糊匹配", "否", ""],
            ["2", "页面·筛选区", "课程名字", "Course Name", "输入框", "否", "—", "模糊匹配", "否", ""],
            ["3", "页面·筛选区", "开课单位", "Offering", "下拉框", "否", "从学院/单位数据选择", "—", "是", ""],
            ["4", "页面·筛选区", "课程分类", "Course Classification", "下拉框", "否", "—", "—", "是", ""],
            ["5", "页面·列表", "序号", "No.", "序号", "—", "—", "当前页序号", "否", ""],
            ["6", "页面·列表", "课程编号", "Course Code", "文本", "是", "全局唯一", "列表展示", "否", "COS102"],
            ["7", "页面·列表", "课程名字", "Course Name", "文本", "是", "—", "—", "否", ""],
            ["8", "页面·列表", "开课单位", "Offering", "文本", "是", "—", "—", "是", ""],
            ["9", "页面·列表", "课程分类", "Course Classification", "文本", "是", "—", "—", "是", ""],
            ["10", "页面·列表", "学分", "Credit", "数值", "是", "正整数", "—", "否", ""],
            ["11", "页面·列表·操作列", "操作", "Actions", "链接", "—", "—", "查看/编辑等", "否", ""],
            ["12", "下钻页·Tab1 基础信息（General Information）", "课程编号", "Course Code", "输入框", "是", "全局唯一，字符格式限制", "课程唯一编码标识", "否", "COS102"],
            ["13", "下钻页·Tab1 基础信息（General Information）", "课程名字", "Course Name", "输入框", "是", "最多 200 字符", "课程标准名称", "否", ""],
            ["14", "下钻页·Tab1 基础信息（General Information）", "开课单位", "Offering", "下拉框", "是", "从学院数据选择", "归属开课学院", "是", ""],
            ["15", "下钻页·Tab1 基础信息（General Information）", "课程负责人", "Course Owner", "下拉框", "是", "从师资档案选取", "授课责任教师", "是", ""],
            ["16", "下钻页·Tab1 基础信息（General Information）", "课程分类", "Course Classification", "下拉框", "是", "枚举以代码集为准", "区分课程类型", "是", ""],
            ["17", "下钻页·Tab1 基础信息（General Information）", "学分", "Credit", "数值输入框", "是", "正整数", "课程学分", "否", ""],
            ["18", "下钻页·Tab1 基础信息（General Information）", "授课语种", "Medium of Instruction", "下拉框", "是", "预设语种选项", "授课使用语言", "是", ""],
            ["19", "下钻页·Tab1 基础信息（General Information）", "学期类型", "Semester Type", "下拉框", "是", "长学期 / 短学期", "适配校历学期", "是", ""],
            ["20", "下钻页·Tab2 课程学习成果（Course Learning Outcome）", "目标名称", "CLO", "输入框", "是", "编码唯一（CLO1/CLO2）", "成果编号", "否", ""],
            ["21", "下钻页·Tab2 课程学习成果（Course Learning Outcome）", "目标内容", "Outcome", "文本域", "是", "最多 100 字符", "课程学习目标描述", "否", ""],
            ["22", "下钻页·Tab2 课程学习成果（Course Learning Outcome）", "布鲁姆分类层级", "Bloom's Taxonomy Level", "下拉框", "是", "固定枚举", "教学目标层级", "是", ""],
            ["23", "下钻页·Tab2 课程学习成果（Course Learning Outcome）", "教学方法", "Teaching Methods", "下拉多选", "是", "预设项", "授课方式", "是", ""],
            ["24", "下钻页·Tab2 课程学习成果（Course Learning Outcome）", "评估方法", "Assessment Methods", "下拉多选", "是", "预设项", "考核方式", "是", ""],
            ["25", "下钻页·Tab3.1 课程内容大纲（Course Content Outline and Subtopics）", "课程内容", "Course Content", "文本域", "是", "最多 100 字符", "章节主题名称", "否", ""],
            ["26", "下钻页·Tab3.1 课程内容大纲（Course Content Outline and Subtopics）", "关联 CLO", "CLO", "下拉多选", "是", "勾选已创建 CLO", "绑定对应学习成果", "是", ""],
            ["27", "下钻页·Tab3.1 课程内容大纲（Course Content Outline and Subtopics）", "各类学时（实地/在线/自主）", "SLT Hours", "数字输入框", "是", "非负整数", "拆分各类授课学时", "否", ""],
            ["28", "下钻页·Tab3.2 过程性评估（Continuous Assessment）", "过程性评估项", "Continuous Assessment", "下拉选择", "是", "预设项", "考核项目名称", "是", ""],
            ["29", "下钻页·Tab3.2 过程性评估（Continuous Assessment）", "占比", "Percentage", "数值框", "是", "0~100", "考核成绩占比", "否", ""],
            ["30", "下钻页·Tab3.2 过程性评估（Continuous Assessment）", "实地学时", "Physical", "数值框", "是", "非负整数", "面授学时", "否", ""],
            ["31", "下钻页·Tab3.2 过程性评估（Continuous Assessment）", "在线学时", "Online", "数值框", "是", "非负整数", "同步在线学时", "否", ""],
            ["32", "下钻页·Tab3.2 过程性评估（Continuous Assessment）", "自主学时", "NF2F", "数值框", "是", "非负整数", "异步自学学时", "否", ""],
            ["33", "下钻页·Tab3.3 期末评估（Final Assessment）", "期末评估项", "Final Assessment", "下拉选择", "是", "预设项", "期末考核名称", "是", ""],
            ["34", "下钻页·Tab3.3 期末评估（Final Assessment）", "占比", "Percentage", "数值框", "是", "0~100", "期末成绩占比", "否", ""],
            ["35", "下钻页·Tab3.3 期末评估（Final Assessment）", "实地/在线/自主学时", "SLT Hours", "数值框", "是", "非负整数", "拆分学时", "否", ""],
        ]
        add_grid_table(doc, FIELD_HEADERS, rows, FIELD_WIDTHS)
    else:
        add_grid_table(doc, FIELD_HEADERS, empty_field_rows(8), FIELD_WIDTHS, empty_hint=True)

    add_heading(doc, "2.10 功能清单", 2)
    add_hint(doc, "每个按钮/开关单独一条。无下钻则写「无下钻页面」。英文名称须与按钮文案一致。")
    if sample:
        add_function_item(
            doc,
            "1",
            "查询 / 重置",
            "Query / Reset",
            "按筛选条件过滤课程列表；重置清空条件。",
            "筛选区点击查询刷新列表；重置恢复默认筛选。",
            "无额外校验。",
            "无下钻页面。",
            filled=True,
        )
        add_function_item(
            doc,
            "2",
            "新增",
            "Add",
            "分步创建课程全量信息：基础信息 → 课程学习成果 → 学生学时，全字段校验后生成课程档案。",
            "列表页点击新增，打开分步下钻页，分步填写后保存。",
            "课程编号全局不可重复；CLO、学时明细支持多条；保存后生成变更记录。",
            "下钻页三个 Tab：General Information（基础信息）、Course Learning Outcome（CLO）（课程学习成果）、Student Learning Time（SLT）（学生学时）。",
            filled=True,
        )
        add_function_item(
            doc,
            "3",
            "编辑",
            "Edit",
            "需补充：是否与新增共用下钻页、哪些字段保存后不可改。",
            "需补充。",
            "需补充。",
            "需补充。",
            filled=False,
        )
    else:
        add_function_item(doc, "1", "查询 / 重置", "Query / Reset", "需补充", "需补充", "需补充", "无下钻页面。")
        add_function_item(doc, "2", "新增", "Add", "需补充", "需补充", "需补充", "需补充，无则写「无下钻页面」。")
        add_function_item(doc, "3", "编辑", "Edit", "需补充", "需补充", "需补充", "需补充。")

    add_heading(doc, "2.11 未确认事项", 2)
    add_hint(doc, "收录未聊过、未拍板、AI 补写或可能遗漏的点。用户已讲清的不要放在这里。")
    if sample:
        add_para(doc, "· 教师角色在本菜单是否可编辑（未确认）。")
        add_para(doc, "· 课程分类的具体枚举取值（未确认，以代码集为准）。")
        add_para(doc, "· 原型页面 ID（未确认）。")
    else:
        add_para(doc, "· 需补充（无则填无）。", color=HINT_COLOR)


def add_catalog_appendix(doc, *, sample=False):
    add_heading(doc, "附录 A 本模块菜单路径（可选）", 1)
    add_hint(
        doc,
        "只保留本模块相关路径。原模板中的全校「基础数据」目录仅为结构示例；写开课管理等模块时请改成本侧栏路径，不要整表照抄。",
    )
    add_caption(doc, "表 A-1 菜单路径")
    headers = ["一级目录", "一级目录英文名称", "二级目录", "二级目录英文名称", "三级目录", "备注说明"]
    widths = [3.2, 3.6, 3.4, 4.0, 2.4, 9.0]
    if sample:
        rows = [
            ["基础资源", "Basic Info", "学校信息", "University Info", "—", "示例：全局参数与多语言（非本菜单正文范围）"],
            ["基础资源", "Basic Info", "部门信息", "Department Info", "—", "示例：组织架构"],
            ["基础资源", "Basic Info", "代码集管理", "Code Set Management", "—", "示例：下拉字典"],
            ["专业信息", "Programme Info", "专业版本", "Programme Version", "—", "示例"],
            ["场地资源", "Venue Resources", "教学楼管理", "Block Management", "—", "示例（原英文 Lecturer Info 为笔误，已改正）"],
            ["场地资源", "Venue Resources", "教室信息管理", "Classroom Information", "—", "示例"],
            ["课程资源", "Course Info", "课程信息管理", "Course Info", "—", "本参考版正文对应的二级菜单"],
            ["教师资源", "Lecturer Info", "教师信息管理", "Lecturer Information", "—", "示例"],
            ["学年学期校历", "Academic Calendar", "学年学期信息", "Semester Information", "—", "示例"],
        ]
    else:
        rows = [
            ["需补充", "需补充", "需补充", "需补充", "—", "写本模块侧栏路径"],
            ["", "", "需补充", "需补充", "—", ""],
            ["", "", "需补充", "需补充", "—", ""],
        ]
    table = add_grid_table(doc, headers, rows, widths, empty_hint=not sample)
    if not sample:
        return
    # 合并同级目录单元格，便于阅读
    def merge_col(col, start, end):
        a = table.cell(start, col)
        b = table.cell(end, col)
        a.merge(b)

    merge_col(0, 1, 3)
    merge_col(1, 1, 3)
    merge_col(0, 5, 6)
    merge_col(1, 5, 6)


def build_blank_prd():
    doc = Document()
    configure_styles(doc)
    setup_section(doc, "厦大马来分校本科教务系统产品需求文档（空白模板）· V2")
    add_title_block(
        doc,
        "厦大马来分校本科教务系统产品需求文档（空白模板）",
        "模板版本：V2　　创建日期：2026 年 8 月 21 日　　对照原文件：产品需求文档模板（空白模板）.docx",
    )
    add_prd_common_front(
        doc,
        [
            ("需求文档名称", "厦大马来分校本科教务系统产品需求文档 — ＜二级菜单名称＞"),
            ("菜单路径", "＜如：开课管理 → 专业开课 → 开课安排＞"),
            ("英文名称", "需提供"),
            ("文档版本", "V1（按菜单各自计数；升版不得覆盖旧版文件夹）"),
            ("创建/发布日期", "YYYY-MM-DD"),
            ("填写人", ""),
            ("确认人", ""),
            ("确认状态", "□ 待确认　　□ 已确认"),
            ("原型页面", "prototype/index.html → page-xxx"),
            ("对应变更说明", "升版时放在新版本文件夹内，命名见变更说明模板 V2"),
        ],
    )
    add_prd_module_skeleton(doc, sample=False)
    add_prd_fields_and_functions(doc, sample=False)
    add_catalog_appendix(doc, sample=False)
    path = OUT_DIR / "产品需求文档模板（空白模板）20260821V2.docx"
    doc.save(path)
    return path


def build_sample_prd():
    doc = Document()
    configure_styles(doc)
    setup_section(doc, "厦大马来分校本科教务系统产品需求文档（带参考数据版）· V2")
    add_title_block(
        doc,
        "厦大马来分校本科教务系统产品需求文档（带参考数据版）",
        "模板版本：V2　　创建日期：2026 年 8 月 21 日　　对照原文件：产品需求文档模板（带参考数据版).docx",
    )
    add_hint(
        doc,
        "参考数据仅示范写法。未与用户确认的示例规则须标（未确认），不得当成学校已明确政策。"
        "课程分类等枚举以代码集为准。原 V1 中「新增」英文误作 Search、Information 拼写及场地资源英文已在本版改正。",
    )
    add_prd_common_front(
        doc,
        [
            ("需求文档名称", "厦大马来分校本科教务系统产品需求文档 — 课程信息管理"),
            ("菜单路径", "基础数据 → 课程信息管理（示例路径，按侧栏实际填写）"),
            ("英文名称", "Course Info"),
            ("文档版本", "V1（示例）"),
            ("创建/发布日期", "2026-05-25（示例沿用原模板日期）"),
            ("填写人", "示例"),
            ("确认人", "示例"),
            ("确认状态", "□ 待确认　　☑ 已确认（仅作模板示范）"),
            ("原型页面", "需补充页面 ID"),
            ("对应变更说明", "本示例无升版"),
        ],
    )
    add_prd_module_skeleton(doc, sample=True)
    add_prd_fields_and_functions(doc, sample=True)
    add_catalog_appendix(doc, sample=True)
    path = OUT_DIR / "产品需求文档模板（带参考数据版）20260821V2.docx"
    doc.save(path)
    return path


def build_change_note():
    doc = Document()
    configure_styles(doc)
    setup_section(doc, "厦大马来分校本科教务系统——需求调整变更说明（简版）· V2")
    add_title_block(
        doc,
        "厦大马来分校本科教务系统——需求调整变更说明（简版模板）",
        "模板版本：V2　　创建日期：2026 年 8 月 21 日　　对照原文件：需求调整变更说明模板（简版）.docx",
    )
    add_para(doc, "一眼对照：哪个模块 / 哪个功能 / 改了什么", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)

    add_heading(doc, "1 版本信息", 1)
    add_hint(
        doc,
        "对照上一有效版本（若中间版本已废弃，直接写 V2→V4，不必保留空号）。"
        "版本文件夹命名：＜二级菜单名称＞＜YYYYMMDD＞＜Vn＞；变更说明放在新版本文件夹内。无实质需求变更不要升版。",
    )
    add_kv_table(
        doc,
        [
            ("需求文档名称", "厦大马来分校本科教务系统产品需求文档 — ＜二级菜单名称＞"),
            ("上一有效版 → 本版", "V__ → V__"),
            ("上一版文件", "＜菜单＞＜旧日期＞＜旧Vn＞.docx"),
            ("本版文件", "＜菜单＞＜新日期＞＜新Vn＞.docx"),
            ("变更日期", "YYYY-MM-DD"),
            ("填写人", ""),
            ("确认人", ""),
            ("确认状态", "□ 待确认　　□ 已确认"),
        ],
        col_widths=(4.5, 21.1),
    )

    add_heading(doc, "2 本次改了什么（总览）", 1)
    add_para(doc, "按「模块 → 功能/位置 → 变动」填写；一行一事。变更类型只填：新增 / 修改 / 删除 / 澄清。无变动的模块不必出现。")
    add_caption(doc, "表 2-1 变更总览")
    overview_headers = ["序号", "模块（菜单路径）", "功能 / 位置", "变更类型", "调整内容（从什么变成什么）", "状态"]
    overview_widths = [1.4, 5.2, 4.0, 2.2, 9.4, 2.4]
    overview_rows = [
        [
            "1",
            "例：开课管理 → 开课设置 → 开课时间设置",
            "例：单位开课时间",
            "修改",
            "例：允许单位窗口超出全校时间；未配置仍继承全校",
            "已明确",
        ],
    ] + [[str(i), "", "", "", "", "未确认"] for i in range(2, 9)]
    add_grid_table(doc, overview_headers, overview_rows, overview_widths)
    add_para(doc, "填写提示：", bold=True)
    add_para(doc, "· 模块：一级/二级/三级菜单即可，如「开课管理 → 专业开课 → 开课安排」。")
    add_para(doc, "· 功能/位置：按钮、开关、字段、列表列、筛选项、流程步骤、弹窗/抽屉、下游拦截等。")
    add_para(doc, "· 调整内容：用一句话说清「从什么变成什么」；新增写「新增xxx」，删除写「删除xxx」。")
    add_para(
        doc,
        "· 状态：已明确 / 未确认。"
        "已明确=用户指定且已讲清；未确认=未聊过、未拍板、或由 AI 补写/可能遗漏。"
        "第 1 节「确认状态」仍是整份变更说明是否经确认人核过。",
    )

    add_heading(doc, "3 前后对照（可选，复杂变更再填）", 1)
    add_para(doc, "总览表说不清时，把关键条目展开对照；简单修改可跳过本章。")
    add_caption(doc, "表 3-1 前后对照")
    add_grid_table(
        doc,
        ["序号（对应上表）", "变更前（上一有效版）", "变更后（本版）"],
        [["", "", ""], ["", "", ""], ["", "", ""], ["", "", ""]],
        [3.2, 11.2, 11.2],
    )

    add_heading(doc, "4 连带影响（可选）", 1)
    add_para(doc, "仅当本次调整会影响其他模块、原型或测试时填写。")
    add_caption(doc, "表 4-1 连带影响")
    add_grid_table(
        doc,
        ["序号", "影响到哪里", "影响说明", "需同步"],
        [
            ["", "", "", "□ 原型　□ 开发　□ 测试"],
            ["", "", "", "□ 原型　□ 开发　□ 测试"],
            ["", "", "", "□ 原型　□ 开发　□ 测试"],
        ],
        [1.6, 6.0, 12.4, 5.6],
    )

    path = OUT_DIR / "需求调整变更说明模板（简版）20260821V2.docx"
    doc.save(path)
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    written = [build_blank_prd(), build_sample_prd(), build_change_note()]
    for p in written:
        if not p.exists():
            raise SystemExit(f"未生成：{p}")
        print(p.relative_to(ROOT), p.stat().st_size)


if __name__ == "__main__":
    main()
