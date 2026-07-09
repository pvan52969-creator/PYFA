# -*- coding: utf-8 -*-
"""从 PRD 字段表与原型 UI 补充数据生成中英字段对照 Excel。"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = SCRIPT_DIR.parent
OUT = ROOT / "参考文档" / "培养方案管理字段中英对照表.xlsx"

if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from field_i18n_extract import collect_all_rows
from field_i18n_supplement import ENUM_ROWS


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14)
META_FONT = Font(size=10, color="666666")


def _style_header(ws, row: int, ncol: int):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_width(ws, min_w=10, max_w=48):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(length + 2, min_w), max_w)


def _write_fields_sheet(wb: Workbook, rows: list[tuple]):
    ws = wb.active
    ws.title = "字段对照表"
    ws.merge_cells("A1:J1")
    ws["A1"] = "厦大马来分校 · 培养方案与开课管理 · 字段中英对照表"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:J2")
    ws["A2"] = (
        f"对应原型：prototype/index.html、app.js  ·  生成日期：{date.today()}  ·  "
        "范围：培养方案管理 + 开课管理（含全部页面/弹窗/抽屉）  ·  "
        "请客户在「客户确认英文」列填写意见"
    )
    ws["A2"].font = META_FONT

    headers = [
        "序号", "所属模块", "页面/位置", "字段类别", "中文", "英文（原型当前）",
        "操作路径/位置描述", "备注", "客户确认英文", "确认意见"
    ]
    start = 4
    for i, h in enumerate(headers, 1):
        ws.cell(row=start, column=i, value=h)
    _style_header(ws, start, len(headers))

    for idx, row in enumerate(rows, 1):
        r = start + idx
        ws.cell(row=r, column=1, value=idx)
        # row: module, location, category, zh, en, path, note
        for c, val in enumerate(row, 2):
            ws.cell(row=r, column=c, value=val)
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = f"A{start + 1}"
    ws.auto_filter.ref = f"A{start}:J{start + len(rows)}"
    _auto_width(ws, max_w=56)


def _write_enum_sheet(wb: Workbook):
    ws = wb.create_sheet("状态与枚举值")
    headers = ["序号", "类别", "中文", "英文（原型当前）", "备注", "客户确认英文", "确认意见"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))

    for idx, (cat, zh, en, note) in enumerate(ENUM_ROWS, 1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=cat)
        ws.cell(row=r, column=3, value=zh)
        ws.cell(row=r, column=4, value=en)
        ws.cell(row=r, column=5, value=note)
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{1 + len(ENUM_ROWS)}"
    _auto_width(ws)


def _write_index_sheet(wb: Workbook):
    from field_i18n_paths import LOCATION_PATHS

    ws = wb.create_sheet("页面与弹窗索引")
    headers = ["序号", "界面 ID", "类型", "操作路径"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=1, column=i, value=h)
    _style_header(ws, 1, len(headers))

    def _kind(loc_id: str) -> str:
        if loc_id.startswith("page-"):
            return "页面"
        if loc_id.startswith("modal-"):
            return "弹窗"
        if loc_id.startswith("drawer-"):
            return "抽屉"
        if loc_id.startswith("sidebar-"):
            return "导航"
        return "其他"

    for idx, (loc_id, path) in enumerate(sorted(LOCATION_PATHS.items(), key=lambda x: x[0]), 1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx)
        ws.cell(row=r, column=2, value=loc_id)
        ws.cell(row=r, column=3, value=_kind(loc_id))
        ws.cell(row=r, column=4, value=path)
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{1 + len(LOCATION_PATHS)}"
    _auto_width(ws, max_w=72)


def build():
    rows = collect_all_rows()
    wb = Workbook()
    _write_fields_sheet(wb, rows)
    _write_enum_sheet(wb)
    _write_index_sheet(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Generated: {OUT}")
    print(f"  字段对照表：{len(rows)} 条")
    print(f"  状态与枚举值：{len(ENUM_ROWS)} 条")
    # 按模块统计
    from collections import Counter
    c = Counter(r[0] for r in rows)
    for mod, n in sorted(c.items()):
        print(f"    · {mod}: {n}")


if __name__ == "__main__":
    build()
